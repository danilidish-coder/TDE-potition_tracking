"""Standalone personnel/roster import contract, including legacy facade identity."""
from dataclasses import FrozenInstanceError
import hashlib
from pathlib import Path
import sys
import subprocess
import tempfile
import unittest

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
import position_tracking as legacy
import position_tracking_personnel as personnel


class PersonnelImportContract(unittest.TestCase):
    def test_legacy_exports_and_standalone_module_identity(self):
        self.assertIs(legacy.PersonnelRecord, personnel.PersonnelRecord)
        self.assertIs(legacy.PersonnelRepository, personnel.PersonnelRepository)
        self.assertIs(legacy.ShiftRosterRepository, personnel.ShiftRosterRepository)

    def test_module_loads_without_entry_point_or_qt(self):
        program = (
            "import sys; sys.path.insert(0,'src'); import position_tracking_personnel; "
            "assert 'position_tracking' not in sys.modules; "
            "assert not any(n.startswith('PySide6') for n in sys.modules)"
        )
        result = subprocess.run([sys.executable, '-c', program], cwd=ROOT,
                                capture_output=True, text=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_record_normalization_and_immutability(self):
        person = legacy.PersonnelRecord('010001', 'ΑΣΤ', ' ΑΛΦΑ ', '  ΑΝΝΑ  ')
        self.assertEqual(person.full_name, 'ΑΛΦΑ ΑΝΝΑ')
        self.assertEqual(person.display, '010001 - ΑΣΤ ΑΛΦΑ ΑΝΝΑ')
        with self.assertRaises(FrozenInstanceError):
            person.agms = 'bad'

    def test_greek_headers_duplicate_merge_and_read_only_input(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / 'personnel.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(['Τίτλος καταλόγου'])
            sheet.append(['Αριθμός Μητρώου', 'Βαθμός', 'Επώνυμο', 'Όνομα'])
            sheet.append(['010001', 'ΑΣΤ', 'ΑΛΦΑ', 'ΑΝΝΑ'])
            sheet.append([10002.0, None, 'ΒΗΤΑ', None])
            sheet.append([10002, 'ΑΡΧ/ΚΑΣ', None, 'ΒΑΣΙΛΗΣ'])
            sheet.append(['invalid', 'ΑΣΤ', 'INVALID', 'ROW'])
            workbook.save(path)
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            people = legacy.PersonnelRepository.load(path)
            self.assertEqual([p.agms for p in people], ['010001', '10002'])
            self.assertEqual(people[1].display, '10002 - ΑΡΧ/ΚΑΣ ΒΗΤΑ ΒΑΣΙΛΗΣ')
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_headerless_roster_uses_known_ids_and_deduplicates(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / 'roster.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(['notes', '010001', '6912345678'])
            sheet.append(['notes', '10002', '6912345679'])
            sheet.append(['duplicate', '010001', '6912345680'])
            workbook.save(path)
            self.assertEqual(legacy.ShiftRosterRepository.load_agms(path, {'010001', '10002'}), ('010001', '10002'))

    def test_missing_unsupported_corrupt_and_empty_files_fail(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            csv = root / 'not-excel.csv'; csv.write_text('10001,NAME')
            bad = root / 'corrupt.xlsx'; bad.write_bytes(b'not a ZIP workbook')
            empty = root / 'empty.xlsx'; Workbook().save(empty)
            for path in [root / 'missing.xlsx', csv, bad, empty]:
                with self.subTest(path=path.name), self.assertRaises(ValueError):
                    legacy.PersonnelRepository.load(path)


if __name__ == '__main__':
    unittest.main()
