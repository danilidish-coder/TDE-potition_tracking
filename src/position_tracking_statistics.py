"""Personnel statistics workbook and charts, with lazy openpyxl dependency."""
from __future__ import annotations
from pathlib import Path
import os
from position_tracking_common import _clean
from position_tracking_personnel import PersonnelRecord
from position_tracking_report_text import ShiftSessionSource as SharedEventStore

def create_personnel_statistics_xlsx(store: SharedEventStore, people_by_agms: dict[str, PersonnelRecord], target: str | Path | None = None) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import PieChart, Reference
        from openpyxl.chart.label import DataLabelList
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    except Exception as exc:
        raise RuntimeError("Δεν είναι διαθέσιμη η δημιουργία Excel στατιστικών (openpyxl).") from exc

    target_path = Path(target) if target else (store.statistics / "ΣΤΑΤΙΣΤΙΚΑ_ΠΡΟΣΩΠΙΚΟΥ.xlsx")
    target_path.parent.mkdir(parents=True, exist_ok=True)

    # Συγκέντρωση μόνο ολοκληρωμένων START/END συνεδριών.
    stats: dict[str, dict[tuple[str, str], int]] = {}
    for session in store.completed_sessions_all():
        start = session["start"]
        agms = _clean(start.get("agms"))
        if not agms:
            continue
        pos = _clean(start.get("position")) or "—"
        direction = _clean(start.get("direction")).upper() or "—"
        minutes = max(0, int((session["end_dt"] - session["start_dt"]).total_seconds() // 60))
        stats.setdefault(agms, {})[(pos, direction)] = stats.setdefault(agms, {}).get((pos, direction), 0) + minutes

    wb = Workbook()
    ws = wb.active
    ws.title = "ΣΤΑΤΙΣΤΙΚΑ"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "ΣΤΑΤΙΣΤΙΚΑ ΘΕΣΕΩΝ ΠΡΟΣΩΠΙΚΟΥ"
    c.font = Font(bold=True, color="FFFFFF", size=13)
    c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 19
    ws.column_dimensions["D"].width = 3
    for col in ("E", "F", "G", "H"):
        ws.column_dimensions[col].width = 12

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    name_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="EDF4FA")

    people = sorted(people_by_agms.values(), key=lambda p: (p.surname, p.given_name, p.agms))
    row = 4
    for person in people:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        name_cell = ws.cell(row=row, column=1, value=_clean(f"{person.rank} {person.full_name}"))
        name_cell.font = Font(bold=True, size=11)
        name_cell.fill = name_fill
        name_cell.alignment = Alignment(vertical="center")

        hdr = row + 1
        headers = ["ΘΕΣΗ", "ΚΑΤΕΥΘΥΝΣΗ", "ΣΥΝΟΛΙΚΟΣ ΧΡΟΝΟΣ"]
        for col_idx, text in enumerate(headers, 1):
            cell = ws.cell(row=hdr, column=col_idx, value=text)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        entries = sorted((stats.get(person.agms) or {}).items(), key=lambda x: (x[0][0], x[0][1]))
        if not entries:
            entries = [(('—', '—'), 0)]

        data_start = hdr + 1
        pie_by_position: dict[str, int] = {}
        for (pos, direction), minutes in entries:
            ws.cell(row=row+2, column=1, value=pos).font = Font(bold=True)
            ws.cell(row=row+2, column=2, value=direction)
            duration_cell = ws.cell(row=row+2, column=3, value=(minutes / 1440 if minutes else 0))
            duration_cell.number_format = '[h]" ώ "mm" λ"'
            for col_idx in range(1, 4):
                cell = ws.cell(row=row+2, column=col_idx)
                cell.alignment = Alignment(horizontal="center")
                cell.border = border
            pie_by_position[pos] = pie_by_position.get(pos, 0) + minutes
            row += 1

        data_end = row + 1
        # Κρυφά βοηθητικά δεδομένα μόνο για την πίτα.
        helper_start = hdr
        ws.cell(row=helper_start, column=11, value="ΘΕΣΗ")
        ws.cell(row=helper_start, column=12, value="ΛΕΠΤΑ")
        helper_row = helper_start + 1
        for pos, minutes in sorted(pie_by_position.items()):
            ws.cell(row=helper_row, column=11, value=pos)
            ws.cell(row=helper_row, column=12, value=minutes)
            helper_row += 1
        ws.column_dimensions["K"].hidden = True
        ws.column_dimensions["L"].hidden = True

        if any(v > 0 for v in pie_by_position.values()):
            chart = PieChart()
            labels = Reference(ws, min_col=11, min_row=helper_start+1, max_row=helper_row-1)
            data = Reference(ws, min_col=12, min_row=helper_start, max_row=helper_row-1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(labels)
            chart.title = ""
            chart.height = 4.3
            chart.width = 7.0
            chart.legend.position = "b"
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showPercent = True
            ws.add_chart(chart, f"E{hdr-1}")

        row = max(row + 4, hdr + 8)

    temp = target_path.with_suffix(target_path.suffix + ".tmp")
    wb.save(temp)
    os.replace(temp, target_path)
    return target_path

