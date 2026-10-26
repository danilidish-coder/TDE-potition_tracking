"""Shift Word document generation and archive naming; optional docx imported lazily."""
from __future__ import annotations
from datetime import datetime
from pathlib import Path
import re
from position_tracking_common import _clean
from position_tracking_personnel import PersonnelRecord
from position_tracking_report_text import ShiftSessionSource as SharedEventStore, shift_report_title

def create_shift_docx(store: SharedEventStore, shift: dict, people_by_agms: dict[str, PersonnelRecord] | None = None, target: str | Path | None = None) -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except Exception as exc:
        raise RuntimeError("Δεν είναι διαθέσιμη η δημιουργία Word (python-docx).") from exc
    people_by_agms = people_by_agms or {}
    sessions = store.shift_sessions_v2(shift) if int(shift.get("version", 1)) >= 2 else []
    roster = [_clean(x) for x in shift.get("approved", []) if _clean(x)]
    by_agms: dict[str, list[dict]] = {a: [] for a in roster}
    for item in sessions:
        by_agms.setdefault(_clean(item["start"].get("agms")), []).append(item)
    start_at = datetime.fromisoformat(shift["start_at"])
    end_at = datetime.fromisoformat(shift.get("end_at")) if shift.get("end_at") else None
    if target is None:
        target = store.current_docs / "ΤΡΕΧΟΥΣΑ_ΒΑΡΔΙΑ.docx"
    target = Path(target); target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = "Arial"; normal.font.size = Pt(11)
    for text, size, bold in [("ΕΛΛΗΝΙΚΗ ΑΣΤΥΝΟΜΙΑ", 14, True), ("Τ.Δ.Ε. ΚΡΥΣΤΑΛΛΟΠΗΓΗΣ", 13, True), ("ΚΑΤΑΓΡΑΦΗ ΘΕΣΕΩΝ ΒΑΡΔΙΑΣ", 13, True)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{start_at:%d/%m/%Y %H:%M}" + (f" - {end_at:%d/%m/%Y %H:%M}" if end_at else " - ΣΕ ΕΞΕΛΙΞΗ"))
    doc.add_paragraph()
    for agms in roster:
        items = sorted(by_agms.get(agms, []), key=lambda x: x["start_dt"])
        person = people_by_agms.get(agms)
        if person:
            name = _clean(f"{person.rank} {person.surname} {person.given_name}")
        elif items:
            st = items[0]["start"]; name = _clean(f"{st.get('rank','')} {st.get('full_name','')}") or f"ΑΓΜΣ {agms}"
        else:
            name = f"ΑΓΜΣ {agms}"
        p = doc.add_paragraph(); p.add_run(name).bold = True
        if not items:
            doc.add_paragraph("Χωρίς καταχώριση θέσης", style=None)
        for item in items:
            st = item["start"]; sdt = item["start_dt"]; edt = item["end_dt"]
            end_text = edt.strftime("%H:%M") if edt else "ΑΝΟΙΚΤΗ"
            p = doc.add_paragraph()
            r = p.add_run(f"ΘΕΣΗ {st.get('position','')}")
            r.bold = True
            direction = _clean(st.get("direction")).upper()
            if direction:
                p.add_run(f"   {direction}")
            p.add_run(f"   {sdt:%H:%M} - {end_text}")
    doc.add_paragraph()
    officer = shift.get("start_officer") or {}
    officer_name = _clean(f"{officer.get('rank','')} {officer.get('full_name','')}")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Κρυσταλλοπηγή, την {end_at:%d/%m/%Y}" if end_at else "Κρυσταλλοπηγή, την __________")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Ο ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ").bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(officer_name).bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Υπογραφή: __________________________")
    doc.save(target)
    return target



def archive_shift_docx(store: SharedEventStore, shift: dict, people_by_agms: dict[str, PersonnelRecord] | None = None) -> Path:
    stamp = datetime.fromisoformat(_clean(shift.get("end_at")) or _clean(shift.get("start_at")))
    folder = store.archive / f"{stamp:%Y}" / f"{stamp:%m}" / f"{stamp:%Y-%m-%d}"
    folder.mkdir(parents=True, exist_ok=True)
    filename = re.sub(r'[<>:"/\\|?*]+', "_", shift_report_title(shift)).strip() + ".docx"
    target = folder / filename
    return create_shift_docx(store, shift, people_by_agms, target)

