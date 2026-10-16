from __future__ import annotations

import json
import base64
import os
import hashlib
import shutil
import re
import sqlite3
import subprocess
import socket
import sys
import tempfile
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
import zipfile
import unicodedata
from contextlib import contextmanager
from dataclasses import dataclass, asdict
from datetime import datetime, date, timedelta, time as dtime
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from position_tracking_offline import OfflineJournal, OfflineConflict
from position_tracking_windows_session import _windows_session_is_active
from position_tracking_machine_queue import _local_root, _machine_queue_root
from position_tracking_handoff import close_previous_windows_user_positions, reconcile_handoff_before_replay
from position_tracking_bootstrap import cache_windows_client
from position_tracking_common import _clean, _norm_search, _safe_name
from xml.etree import ElementTree as ET
from position_tracking_report_text import (GREEK_MONTH_GENITIVE, _duration_text, build_shift_text, build_roster_shift_text, _format_seconds, shift_report_title, build_roster_shift_html, _escape_html, build_shift_html)

APP_NAME = "TDE Krystallopigis - Katagrafi Theseon"
APP_VERSION = "1.2.5"
SETTINGS_NAME = "position_tracking.settings.json"
DEFAULT_CUTOFF = "10:00"
STATISTICS_PIN_SHA256 = "ca444860f0e1ddca9596007253ea43b3a5133780b0a566d9a096a53568c063fb"

# Exact approved shared visual assets from TDE-Control-Center/SHARED_ASSETS.
POLICE_EMBLEM_PNG_B64 = """iVBORw0KGgoAAAANSUhEUgAAAGIAAABoCAMAAAA94NVaAAABgFBMVEXb5OJam2Erllld3dnp1Cva4t4eq5Tc69hro6ecl13Y3twv16kr49ms5tf06KLx0lcxZWVlcllO0Z6HemKMpZzKti0obTAAG22YxrbCrE+5pz3/qqrAv8L//38mPYg9cZwrijchrcs7wmbRuLn//7QAAP9ZbqJJwWN///+ZmZm7u8O+vsGMn8Cqxqq/wcH/AAD/f38AAAD8/f0cQ4IMOHwjR4QC/f3JycrW1dMRPII2Vo3+/v7n5+dMZ5Z3h6izuciKl7HW2uSao7WxtLwBK3VZcZumq7hCXJBne6Py6a3j4tpkd5zU1tS6w9LT1NR/f3+pqarLy8vV2Nfh3tgu+/z65o6/v7+Zp8T423KBjarm6ObX2NhYc6Fo+v6YlVXe4OmjrMdT+/6pmU9XaG5sdWiypk/W19f++AfOuE3W1tbU19QN2NEpS3pGXHNy2tmKiGOJ2diR9/vt8u8AKosL7OvKysvKysvX4tv9427b5OLS89L66RIaQX3SwkjKy8vk5uXpCAtUAAAAgHRSTlMd/////2T/8///nv////////////////////8D/QL///////8DAf//AgUirf8JeAECAP/+/v7//vz+/gT+/v7+/f/9/f/+/f7+//790P6yAgQFjf7//wT///0qa/////////////9N//8HMv////////8n//9ozfv/UP////+yTmkf+WIAABCHSURBVHjatVoHe9tGtiUpqtuSa9pm29t9vRKYEQYdBKvYJJHqUlQsq9qKbdlxL399z50BQJCWZGdfdr4vCk0COHP7uXeQKfzDV+brL/23pSV7CWt10c4s2vZvDmEvLQ59kbF/MwjbjnZsj745+gVra+vN6Gby228Acaq23O10Wo7Tc2gJRxMfP3Yy8vmri/8/CHsRT9k86gpDNIQwStUwCIKwWvrRED3HMET3E35ftf9+CJs22O2IXk8Y1cB1LYubjDGdW5bbDEJfCOG0pi4KX0DJXPuT3W05PcMPXPVwU9d1jv+Y/Mwtr2QIR3Q27cLCwt8BgX1djGKfeD49EXvH0iUG3/Bc+ghUy6saQnSOyPK/EgI32FMtx6g2LRNP1ytr20wCAIIFmqa59E+ul8u6GxIIXGz1V0GQkVuOqFY47V9n3NH8CIIzCwjaGuOcucXQIlECgExlCvavgFiFEV6JUpMez4NiUC6vaQ6XGFxnRa24oRUhG3OBZXjYhhWKBrR1ubIyl8qw2RJGwBl0pOtFTeiuoWmRpkhN7XqkqSIJ5MBYZqXUE6OFSwXJXOaqGdEruabO9G0uH7lLT/KZVNO2ptZaGZbY09o1X1i4FNqCRS5VVuYSM2wJEehw0FqxyJQ2/H/xNY2Tz+qGVtx2PfyF4jxgW2WmM7e9XWYVgwxy+kWIhcJ/jwqjTjpiPhlVZ4ZmsHJdK1om+S3ijrGya+y6ugNLuCb0xfY0zWPMLTVEprD4BQhkjHHHcJlu1ou8rXll+I0HxZsWIgEPM2V40yeLW22o0FOxsgZJedkKe8BYuhbCXrBbr0oW1FPRtGJRC8gnubNnqYjmvEbL4rryZRfq2+UQDkJqmnCZCYOcDQdIZijgOg4hYMdtaVTHD/CZ8gV3vaBaMuQqhYFXk6Cuvxd77y782uTA2ByyeWYwHraEQTLAEBY07QtNgwTYfiX0DWEMLGQWnfIKaQ3Xedz3cJsZNIS9uHAVhA1vhR04m1h7r5PsvunCtmYtaCOTI606kcM6TgN4SO6eBfE4/ELzyyZcxKszvSq6g+bIDEScEE0mXd/H9bixBgBKD0b/8clyCMVHbJPr7erSs4oIQ8sHxsKlECiSrYZHqacGdzeQi9rbcCFPAmiXLofsAiPrQYVxGZYOpHaNV0dpc/Qh/q/QbVTJT5D1kC+KljRnCQDa1QuiiEDlSp2k2aMHNHstO4WR6cfcJpla1z3648uEwerYpqNdu4BBPki5xXW0olFcM/VQdOzVzyFWMy0BNbGaVnRhOp+Sm3e9CH1tuaRfa1d94cEcTqZvjhhiqdB1qrQViocAzgeFBV+DQIJQxqEkr6HAcxO31kWrn9gjiEXK35SfGeUCrU3hHAwFwnVLVBjbLtbLTCZdXS+JrcQamdidtpyQITrr3hopP2Cmh3xrpZdb6S+3NvCLT7rSqWgxVmsL5B8Swx6CEEbN1M2iEh2Zug4tbTC9vygdJgsXpH7hFPmoGjB6pS2toVcbW3HOzURR1xUhriUlwSmgWtcwdp0NkyfP4WxjAEJPfsJHA45Vwv2Wsp0gaxgDirJtuFPF1CkvwaHgVzoYkuZ4qa3yISlS6AShGfBHyRzgV20IVBJvImtk1J8tp0T74r7yOiRl3NWHoFI0CKGbMa+KIBwD7sJCgxI8jAJTxk4lIX5vd3uUOmCtbYgaoKhRxMUQ+N5ELbo3Pz+/sr6+vrIyPz9zwPCVHrM3Q8Yg7bJcw/0132KWIS6UNTJR/iNrsTLV07DNsRlSagxhPtfPbx/Ozq6vNG7Sevt2fmZ29vvsOZATKUhVTVOSkiLRLBY2Ir/NyNxx1EN24v5avWZRrYAQWgriYOfk9cr8+E/7+8tzci0v7/80PjPz+mTnoC8F/KQU8TjpcZVGlHCVuTuizhgVR23Xr1NYv0oguHmQn3n7bH+uv26o/+0/ezvzeofMriA0aY09KpV1iftKcR7ltEKgtlRUxgvKlhIihni48mxuYI3diT/tzz9IQQgRmtBAYEHjuK/a+I8YAolX+LAUnBb5FUkQcT0IsT8EkY0/LQ9AICFSlgaTdl0KVdHpQ2xR3On6t8iyMIVZikpQAvFT6vmPxgAxfSkExKBcvV2UWQ5JRMROu1T4iGoHa3t7WvAtZIn0BAhy/iGIkeydR7npG9PTj2IInkA0jKpuWmp/e+S2kUchQD4KkAlJVnf3KuV6nMKdDSp8zxOIn7FGbt7O3ZnIZXO3IojnxK/iPRkIiIpkI5RifCHtTX/sl2C+CLrI2kFcqh1PRxrl2Qhi+enIyMjEdyO57MjY2NiLSFF/oFQbQWhIuGAu4KecqnhV5ZAMyNMbx7dIpZbrhe0aqxoJxB71wDOvU4rKZW+8mHD/nNhihq5w/RhCeICgkuxoFgucKQWxVBh1qpSgiK6WkcpKfQhJCefXUxBjf5peHqmM3bgRQazQFRU/qX8bpCgL9BBSeCKBOHJClYpMVgn61gaE/xnEo2/+aeRONpd7FEtBVyRSOEZINLLoo/nnYCKStBFZ32oE0jupqBbNWlKwP4fA3ieyE7nc2KO5tBQpiKrqNlUz0OuQM2UKvwd/CmTVDojGIjyvhsjl5m5Wb/Rl+hwC2Za5a21/m5ofEUGsFn6BdyJidumh/FqIsdyd9r0XyzcevZi71BYSAr1A+dsyQUCK04KUYot4JqgcnFSm2ashHt3613t3pnPf5OauhmBeEay9zpUUxMULv0PJCyQEOhZmXguB3DGR/ebW2IsrIaqKiWHVy4miyGkJok5kbs9l15r7xZ2/jDz6zBYpiFA33cCnLLVdrsQeJSFMIkET9Itl/XgFBJLfH//43Xd3pqfHroJAt44n1SWEW242+qGnhWouwOuO5qIHcQYgxmd/igJ7+s8T2dyfoKikXrwejAuiIa76uAdWFEf3QsFu+ZJ+oFfngaWHQwmkMXtTaQlRNwFjTL9IFPVsvTGYQIxmmYiK73JGBP0TEQSZCzuGpQYcNM5gXj/Tcsknvx9fjlz23o25myOpzD7+vWKcSaY1arI+r2GzRGyjTAujd4RMg6Srtv9tJYHw5CTKvE2plhL5vXsj09lsbvqpXBBi5eA5G0jmSHbMkhQMRKstMvHorDDlVNR4iYKD6X5SkpjkY49P3s7NPR03/iK+y+Zu56bnRgxj3ID23p4oKpVUPVibrwX1thxYwfuT6VxhVMARCGGNODlYf6p2y9L6GgRhZOTpyM8TL9qkqOWbI9Dds9cHZorkqHJB1iaaZ7EmRV6fgfgSISQEikojDYH1w7xiCC/Cm4bRjknOyo6pp6gaFT3ZjUk0nW30UiTHJpIjU2RQRsXgUcVICKf5fvateuxE6eefnyqI/ZkfzIieJySHEhHj2xP/5SN3l5yzBAIUpOGqFLI24dWplgywQVLV+rhy3MSflsdPHsfsXEEojiMb3TJ1Q8a4nUAsoLQGLJlmFTn3pRgpZm4+nB0f4FJphAgCXRUzudN25TCONZ1un9MCrGXQQCvwjSKe7ZWVGKn+gpu3Z6PoUDE3k0JQECSENKfmV6gdCRtHhb8m5qY01ZSNms7BEaC0EjlVuoUhXTUSVvhs5fBxuoMyEkuoQDeI/Mf9vTL3wkWrKvsIyRGUUzlDXZJ5b31ekdv9V69nP+uSBLyS6WDfLvpRuOyG4gapXg+NmIw+on9rRWXxoUZsQnu9/nZ/bv/m/Pr4bqrXkxAyd8AnXWaWrT3q7kVmoCm27TOwatl9Mppp+jTzEY2N4V5vfn3m1fz6PPY8IIVlAIH246gRMY02xBCEFEN29q7lipqPkoVtGFWW6ucUs1hZX4k7VrXolwoNW0w5qXS46s78xllhqO8uZNAocdNCw6q1PRrvgvWiXzfjBTKfGrggeJNfpHOEYPSy4fTKNDaqNzr2/w5BLNod5EIa4QRytyHNlYThuakhQXMjlCjFvcDlqa+BAEauc1crtuWkG2EnNkEKBscstv1G+BYuLFI3qAXkQjVfDIxBStUNl2YpzWQOGQ1ARKCm0lrgEUeDTkXnszELqeqoAX2ymuru1Yy/JAbHUU5oGV5FDA+kYOm4rZDjE9dIjzn7EKu2tDj5E2fSymomKgYxeHPwC2izSemQ1+pBuyinipzmOL+7bDZYOEP0kEpZchqi664/OBoUgT80twstFp/+4F7QSXOj8dJetS+fcMqpl87SExQw3cH5Y0pzNOMsVVgUIdJXk5CwLx+iLqCIB4ynco+cWViBT/PS3YE5qkMCiFKTpzakRiUwxKeBefPgKPjilWiW+7fA6Tn8XrfqJUM5l0gtI3QHAFSAcsO5ehRcKJyejjqimdShWztZy3y4w9XRV+gbRuLE7aqHsmDqfAgBpna69uK15xejcBC1NYR6/jhvZY8nzXMUkzKj88JqBFPacOkcQ530xVmGTgmqESe4fNpcsKfGWy08IME4zx9PHhznzdkPuvleHV0Qjq8U9mPVq1jRMSJ1cW6pSiH58mXr7CpbLBU6jpDzdy/GMPmTu/fz7uzxjpnfoWMRUz5Sd5sbVd8QPVxd9ZrR2aXbo7tpfJ+5CoIGhNUwDOFAgQo+WGHy5GRy5/DdfTN//OT54cljnaaM6qRVnrMaoiFPYJsVy61bLpZnDJ1gDEgxJShzoJz0SpGvHObz+XP+8HDnDzvHH54fHj/QP+ygyyF+SocJiOkK+YGUp0QNkKnmEvbVhwt1CmxmVWF0nQQ5fzJ5y9zJ5588vHv/cDL7Ls/z7w6ef8g/jmd2JBDV+9CnwRzlcfRG9pVHJIUMVXiygY7cFNLxByLj+ZPj/KE1Ofnhwc7tyQf8EAqbhJfxdHSSxa16TZ4wNBv/M3gqNnjQI6pRALme3zOou8QDbh0c6PzDu3eT5u0fHuw8/H7ySfbdiamn0ozKHZYrITznl2sgCq8ktyUIiwTx63KoDFEeZx8cmDvH73YmD+8/eHL7wZNbXDcHIhvlRUFsAOL0KoiFwqghyz4gQCfcquj5niWPuiXMyYcnHNFY44C5e/fubT4AEkGYwVCKGoI4ErUIAqZjvIKSaciDL/ae1GGZ5/cPDx7n3+UfmrWDu5PnKW1FEDqrgqJdeyLWZAkEzKjOy/0QYSwz1Xt61Pnhg3OTRrS37t4dloLamZeDCEMQZw2vDyHLBXODkowuP9xoymPJ98Q6pIVN/aA2DAFm0LLtayAyWmDqfQhdOiRKWVUmJTqWBCUBtX8fz7n1YQhmidbQCeiARy2edn1JtWKIGIVKhkrmgo4l65SVVC4fsAUdI7q97vUHoF1DuqmCUNqI6DrlI4nhyGxXCiAPT79ZYdbkkNZNHfFcVpLQBEgSxpvIdoqlWxUvoJonjzx/GT36dPRLR2EZPrKfZZbj90NqTbrTux5iiSqSWpRqsXdK2aLnCBSBo8Wl6Fb7dPFs/CU4A71E0ybzSGVaEXG79qSYZhVRYfY8+UIMrVZnajS6aXEVK8oO9uhUN7oabrCBZB5dbn/h1N7e3Ly42LQ/iQYdCXa6W2cX8pZV8JPEGel9Juk19sXZVpc02IAe/U9nZxebm/bXvs3y8ePUm4X46tPLXvCwl07/PQI87X78z5eO8ZVvs9hqxf/854VFe+Hq14YW+qfn9tbowl8XcevC10qxuJT54rtVaYmWfpPX4a5fC0ur/2iIq9ffACg1lb6ZgXizAAAAAElFTkSuQmCC"""
GREECE_MAP_SVG_B64 = """PHN2ZyB4bWxucz0iaHR0cDovL3d3dy53My5vcmcvMjAwMC9zdmciIHZpZXdCb3g9IjAgMCAxODI4IDE3MDQiPjxnIGZpbGw9IiM2ZDlmYzEiIGZpbGwtb3BhY2l0eT0iLjEwIiBzdHJva2U9IiM4ZWI4Y2YiIHN0cm9rZS1vcGFjaXR5PSIuNDIiIHN0cm9rZS13aWR0aD0iMyIgc3Ryb2tlLWxpbmVqb2luPSJyb3VuZCI+PHBhdGggZD0iTTExOTYsMTcgMTE0MCwxIDExMzcsOTMgMTAxMCwxMjEgODI2LDQwIDU4Niw5NyA1NzMsMTQ0IDQ1OCwxNDEgMzgwLDIwOSAyODEsMjEyIDI4OSwyODAgMTYyLDQxNSAxMjYsNTE4IDIyMyw2NTcgMjM1LDYxOSAyOTIsNjM5IDI5OCw2ODAgMjIzLDY4OCAyODEsNzk5IDMxOCw3NjQgMzUwLDgxMCA1MTAsNzc1IDY1MCw4NDUgNTg4LDg5OSA0MTMsODA0IDI4MSw5MDEgMzc4LDEwMjQgMzc4LDExNTcgNDA4LDExNzkgNDU4LDExMTEgNTE0LDEyNTYgNTQxLDExNjAgNjQxLDEyNDcgNTU4LDk4NCA2MzQsMTA0NyA2OTksMTAxMiA2MTEsOTAyIDcxMSw4NzAgNzk1LDk1NSA3OTgsODQyIDcwOSw3NzQgNzk0LDc4OSA4NTAsODg4IDg4Niw4NDggODI4LDgzMiA4MTEsNzMxIDcxNCw3MDQgNjY5LDY0MCA1ODUsNjg1IDY0Niw2ODUgNzE4LDc3MyA1MzMsNjg2IDYyMiw2NDAgNTgxLDU5OSA2MDIsNTYxIDY1Miw2MDMgNjI2LDYyNSA2NzEsNjA5IDU0NCw0MDggNTU1LDI3OSA2MDcsMjYxIDYwMCwzMjAgNzI3LDM1OCA3NzYsNDI2IDc0MCwzMjggODUwLDM3MyA3NjAsMzE2IDc0MSwyMzIgOTcwLDE3MyAxMTM5LDIzNyAxMTcxLDExNCAxMjIxLDkzeiIvPjxwYXRoIGQ9Ik02OTgsMTUxNyA3MTIsMTUzMSA3NDYsMTUzMyA3NTUsMTUyNyA3ODMsMTUzMiA3ODgsMTUzOSA4NTUsMTU0MSA4ODgsMTU2MiA5MTcsMTU2MyA5MjMsMTU3NiA5MTksMTYwMyAxMDAxLDE2MDAgMTAzMCwxNTg3IDEwNTksMTU5MSAxMTMxLDE1NzggMTE3MiwxNTg2IDExODcsMTU3NiAxMTk0LDE1NDggMTIwMSwxNTQzIDExOTMsMTUzOSAxMTkzLDE1MjIgMTE4OCwxNTIyIDExNzIsMTU0MiAxMTUxLDE1MzQgMTEwNywxNTYyIDEwOTIsMTU0NyAxMDk0LDE1MjQgMTA5OSwxNTI1IDEwOTQsMTUxOCAxMTAxLDE1MDkgMTA4NCwxNTA3IDEwNDIsMTUyMCAxMDMxLDE1MDkgOTc2LDE1MDkgOTY4LDE0ODggOTU3LDE0OTMgOTA3LDE0ODggODgzLDE1MDEgODMyLDE1MDUgODI4LDE0NzYgODE3LDE0ODMgNzk0LDE0NzMgNzk2LDE0NjcgODE1LDE0NjcgODEzLDE0NTEgNzk5LDE0NTAgNzg4LDE0NjggNzQ4LDE0NjMgNzM4LDE0MjcgNzM2LDE0NjkgNzIxLDE0NzEgNzEyLDE0NTEgNzA5LDE0ODcgNzAwLDE1MDQgNzA0LDE1MTJ6Ii8+PHBhdGggZD0iTTExMDAsNTkyIDExMDMsNTk2IDExMDEsNjA4IDExMDQsNjA3IDExMTAsNjEzIDExMTUsNjE0IDExMjAsNjIxIDExMjgsNjIwIDExNDAsNjI3IDExNDMsNjI0IDExNDIsNjIxIDExNDYsNjE5IDExNDYsNjE1IDExNTAsNjE0IDExNTUsNTk5IDExNjEsNTk3IDExNjQsNTk5IDExNjUsNTk3IDExNzUsNTk4IDExNzksNjAyIDExNzcsNjEwIDExNzEsNjEyIDExNjgsNjE2IDExNjAsNjE2IDExNTcsNjI0IDExNTQsNjIzIDExNTEsNjI4IDExNDUsNjI4IDExNTYsNjQzIDExNjAsNjQ0IDExNjEsNjQxIDExNjcsNjQxIDExNjgsNjQ0IDExNzEsNjQyIDExODcsNjUwIDExODgsNjUzIDExOTEsNjUxIDExOTcsNjUzIDExOTgsNjU2IDExOTksNjU0IDEyMDMsNjU2IDEyMDQsNjUzIDEyMDksNjUxIDEyMjEsNjUxIDEyMjAsNjQ4IDEyMjIsNjQ3IDEyMTcsNjQyIDEyMTgsNjM4IDEyMjIsNjQxIDEyMzIsNjQyIDEyMzAsNjI4IDEyMjIsNjI0IDEyMjEsNjEzIDEyMTcsNjEyIDEyMTgsNjA3IDEyMTQsNjA2IDEyMDcsNTk0IDExOTEsNTgxIDExOTMsNTc0IDExOTcsNTczIDExOTcsNTY3IDExOTEsNTY3IDExODksNTYzIDExODYsNTY0IDExODUsNTYxIDExODksNTU5IDExODUsNTU3IDExODIsNTU5IDExNTcsNTU5IDExNTgsNTY5IDExMzgsNTc5IDExMjUsNTgwIDExMjQsNTgzIDExMTUsNTgzIDExMTAsNTc3IDExMTAsNTgxIDExMDQsNTg2IDExMDQsNTkweiIvPjxwYXRoIGQ9Ik0xNTM3LDEyNDcgMTUzMSwxMjUyIDE1MjUsMTI1MyAxNTI0LDEyNTEgMTUyMywxMjU0IDE1MTYsMTI1NSAxNTAwLDEyNjQgMTQ5NywxMjYzIDE0OTQsMTI2OCAxNDg0LDEyNzQgMTQ4MCwxMjczIDE0NzksMTI3NyAxNDc1LDEyNzggMTQ3NSwxMjgyIDE0NjQsMTI5MiAxNDYwLDEzMDcgMTQ1NiwxMzA4IDE0NTUsMTMxMiAxNDQ1LDEzMTggMTQ0OSwxMzIwIDE0NDgsMTMyNiAxNDUyLDEzMjYgMTQ1OSwxMzM0IDE0NTksMTMzOCAxNDU2LDEzMzkgMTQ1OCwxMzQwIDE0NTcsMTM1MiAxNDUzLDEzNTUgMTQ1NCwxMzYwIDE0NTEsMTM3MCAxNDU0LDEzNzAgMTQ2MywxMzc3IDE0NzEsMTM2OSAxNDc2LDEzNjggMTQ3NSwxMzY1IDE0OTEsMTMzOSAxNTAzLDEzMzIgMTUxMywxMzM1IDE1MTUsMTMyNyAxNTExLDEzMjYgMTUxMSwxMzE3IDE1MTUsMTMwOSAxNTE5LDEzMDkgMTUyMiwxMzA2IDE1MjIsMTMwMiAxNTI1LDEzMDEgMTUyNywxMjg1IDE1MzAsMTI4NCAxNTMwLDEyNzkgMTUzNCwxMjc4IDE1MzMsMTI3NSAxNTM2LDEyNzQgMTUzMywxMjY4IDE1NDAsMTI2MSAxNTM4LDEyNjAgMTUzOCwxMjUyIDE1NDAsMTI1MHoiLz48cGF0aCBkPSJNMTE0MSw3NDEgMTE0MCw3NDQgMTEzMSw3NDEgMTEyMiw3NDUgMTExNiw3NDUgMTExNSw3NDMgMTExNCw3NDYgMTEwNiw3NDUgMTEwNCw3NTMgMTEwMCw3NTQgMTEwNSw3NTggMTExMSw3NzEgMTExNCw3NjkgMTEyMyw3NzUgMTEyMiw3NzggMTEyNiw3ODYgMTEzMSw3OTAgMTEyOSw3OTQgMTEzMSw3OTUgMTEzMiw4MDIgMTEyOCw4MDcgMTEyOCw4MTIgMTEyMCw4MTUgMTExNiw4MTMgMTExMCw4MjQgMTExNCw4MjYgMTExNCw4MjkgMTExOCw4MzAgMTEyMSw4MzggMTEyNCw4MzUgMTEyOSw4NDIgMTEzNSw4NDUgMTEzNSw4NDAgMTEzOSw4MzggMTEzOCw4MzUgMTE0MCw4MzMgMTE0OCw4MzIgMTE0OSw4MzQgMTE0OSw4MTcgMTE1Niw4MTAgMTE2MCw4MTAgMTE1NSw4MDIgMTE1NCw3OTUgMTE1Niw3NzUgMTE1MSw3NjkgMTE1Myw3NjUgMTE1Niw3NjcgMTE1Nyw3NTMgMTE1NCw3NTIgMTE1Myw3NTUgMTE0OSw3NTUgMTE0OCw3NTIgMTE0Myw3NTB6Ii8+PHBhdGggZD0iTTE4Niw3NzEgMTg3LDc4MyAxODQsNzg0IDE4Nyw3OTEgMTgzLDc5MyAxODYsNzk1IDE4NSw4MDEgMTgwLDgwMSAxODEsODA0IDE3NSw4MDkgMTcxLDgwOCAxNzMsODA1IDE2Niw4MDQgMTY2LDc5OCAxNjEsODAxIDE2MSw4MDQgMTU4LDgwNSAxNTgsODE3IDE1Miw4MjEgMTUzLDgyOCAxNTAsODM3IDE1Niw4NDIgMTY0LDg0MSAxNjIsODE1IDE2OCw4MTMgMTcyLDgyNCAxNzEsODI4IDE3NCw4MjkgMTc5LDg0MCAxNzgsODQzIDE3NCw4NDEgMTc1LDg1MiAxNzksODU2IDE4Miw4NTUgMTgzLDg1OSAxOTEsODU0IDE5OCw4NTMgMjAxLDg1OSAyMDYsODU5IDIxMSw4NjUgMjIzLDg2NSAyMjUsODY3IDIyOCw4NjQgMjI2LDg2MyAyMjgsODU2IDIxMyw4MzIgMjE0LDgyOSAyMTEsODI4IDIwOSw4MjIgMjA1LDgyMSAyMDIsODI0IDE5Nyw4MjIgMTkzLDgwOSAxOTcsODA3IDE5Niw3OTAgMTkzLDc4NyAxOTUsNzgzIDE5NCw3NzkgMTkxLDc3OCAxOTMsNzc0eiIvPjxwYXRoIGQ9Ik00OSw0NjIgNDMsNDY5IDQ2LDQ3MiA0Niw0NzYgNTEsNDc4IDQ5LDQ4NiA1Nyw0OTEgNTcsNDk2IDYwLDQ5NyA1OSw1MDEgNjQsNTAxIDY4LDUwOCA3MSw1MDcgNzUsNTExIDgwLDU0MiA4MSw1MzkgODUsNTM5IDg5LDU0NCA5Niw1NDYgMTAxLDU1MiAxMDUsNTUyIDExMyw1NTcgMTE1LDU2MSAxMTYsNTU5IDExOSw1NjAgMTIwLDU1OCAxMTMsNTQxIDEwOCw1NDUgMTAzLDU0NSAxMDEsNTQwIDk3LDU0MCA4OSw1MzUgODYsNTEzIDg5LDUxMiA4Nyw1MTEgODksNTAwIDc4LDQ5NiA3NSw0ODQgNzcsNDc5IDk2LDQ2OSA5Niw0NjQgODYsNDYxIDg0LDQ1NiA4MCw0NTUgNzIsNDYyIDY4LDQ2MiA2Nyw0NTggNTksNDYyeiIvPjxwYXRoIGQ9Ik0xMjM0LDk0MyAxMjM1LDk1MiAxMjM4LDk1NyAxMjQyLDk1NiAxMjQ2LDk1MSAxMjUzLDk1MCAxMjU0LDk1MiAxMjU1LDk0OSAxMjYwLDk0OSAxMjY4LDk1NCAxMjcyLDk2MSAxMjcxLDk2NCAxMjg3LDk2NSAxMjg3LDk2MSAxMjkxLDk1NiAxMzA2LDk1MSAxMzA3LDk0OCAxMzE0LDk0OCAxMzE1LDk1MSAxMzE2LDk0OCAxMzIyLDk0NyAxMzE0LDk0NSAxMzEyLDkzNyAxMzE2LDkzNyAxMzE1LDkzNSAxMzA1LDkzNCAxMzA0LDkzMSAxMzAyLDkzMyAxMzA1LDkzOCAxMjk5LDkzOSAxMjg4LDkzNSAxMjg1LDkyOSAxMjc3LDkyOSAxMjc1LDkyNiAxMjczLDkyOCAxMjU3LDkyOCAxMjU1LDkzMiAxMjUxLDkzMiAxMjQ2LDkzOHoiLz48cGF0aCBkPSJNMTA1NSwxMDcwIDEwNDgsMTA3MCAxMDQ1LDEwODAgMTA0MiwxMDgzIDEwMzksMTA4MiAxMDM2LDEwODYgMTAzMywxMDg0IDEwMjgsMTA5MyAxMDI0LDEwOTMgMTAyNCwxMDk2IDEwMjAsMTA5OSAxMDI3LDExMDMgMTAyNywxMTEyIDEwMzUsMTEyNiAxMDM4LDExMjcgMTA0MCwxMTM0IDEwNDYsMTEzMyAxMDQ4LDExMjggMTA1NywxMTI1IDEwNjAsMTEyMCAxMDY1LDEwODkgMTA2MiwxMDgxIDEwNjMsMTA3NyAxMDYwLDEwNzd6Ii8+PHBhdGggZD0iTTEwMzAsNDEwIDEwMjQsNDEyIDEwMTksNDEwIDEwMTUsNDIxIDEwMDksNDIzIDEwMDgsNDIwIDEwMDYsNDI3IDEwMDEsNDI2IDk5OCw0MjAgOTk5LDQxNyA5OTQsNDEzIDk4MSw0MTUgOTc5LDQxMiA5NzQsNDE1IDk2NSw0MTYgOTY1LDQyNSA5NjksNDI5IDk2NSw0NDUgOTY3LDQ0NyA5NzYsNDQ1IDk3OCw0NDggOTg0LDQ0NCA5ODcsNDQ4IDk4MCw0NTUgOTgyLDQ1NyA5OTAsNDU3IDk4Niw0NDUgOTg3LDQzNyA5OTUsNDMxIDk5OSw0MzEgMTAwMyw0MzYgOTk4LDQ0NiAxMDA1LDQ1MiAxMDA1LDQ1NyAxMDAzLDQ1OCAxMDEzLDQ2MSAxMDE1LDQ1NiAxMDEyLDQ1MiAxMDExLDQ0MSAxMDE0LDQzOCAxMDE0LDQzMyAxMDE5LDQzMSAxMDIwLDQyNCAxMDI4LDQxNiAxMDI4LDQxMXoiLz48cGF0aCBkPSJNMjAwLDkwNiAxOTcsOTExIDE5Nyw5MjEgMjAwLDkyOCAyMDUsOTMwIDIwMiw5MzUgMjA4LDkzNyAyMTAsOTQ1IDIxOSw5NTAgMjIxLDk1NSAyMjcsOTYwIDIyNyw5NjMgMjMzLDk2MyAyMjgsOTU1IDIzNSw5NTIgMjM0LDk0OCAyMzksOTQ0IDI0Niw5NDMgMjU3LDk1MSAyNTcsOTQ4IDI0Miw5MzYgMjQyLDkyOSAyMjksOTIxIDIxOSw5MjAgMjEwLDkwOSAyMDksOTAweiIvPjxwYXRoIGQ9Ik05MTMsODg0IDkxMSw4ODggOTA4LDg4NyA5MDYsODkyIDkwOSw5MDkgOTE3LDkwOSA5MjEsOTEyIDkyNSw5MTkgOTI0LDkyMiA5MjcsOTIyIDkzNCw5MzMgOTQyLDkzOCA5NDQsOTQyIDk0Myw5NDcgOTQ3LDk0OSA5NDYsOTUyIDk1MSw5NTQgOTU1LDk0NCA5NTgsOTQzIDk1Nyw5MzYgOTUxLDkzNiA5NTUsOTIzIDk1Miw5MjQgOTQ4LDkxOSA5NDgsOTE2IDk1MSw5MTUgOTQ5LDkxNCA5NTAsOTA3IDkzMSw5MDMgOTI3LDg5MSA5MjEsODg4IDkyMSw4ODR6Ii8+PHBhdGggZD0iTTg5MiwyMjYgODg4LDIyOSA4ODUsMjI4IDg4NSwyMzIgODgwLDIzNCA4NzIsMjU3IDg3NSwyNjIgODg0LDI2MiA4OTEsMjcxIDg5NiwyNzMgOTAxLDI2OCA5MDcsMjY5IDkxNSwyNTkgOTExLDI1NSA5MTEsMjUxIDkxNCwyNDggOTEwLDI0NCA5MTAsMjQwIDkxMywyMzcgOTA2LDIzMyA5MDcsMjI4IDg5OSwyMzB6Ii8+PHBhdGggZD0iTTY3MiwzNjcgNjcxLDM4NCA2NjcsMzg4IDY3NCwzOTQgNjc3LDQwMCA2NzYsNDA1IDY4MSw0MDkgNjc4LDQxOCA2OTUsNDE2IDcxMyw0MjUgNzE3LDQzMCA3MTgsNDI4IDczNiw0MzEgNzQxLDQyOSA3MzMsNDI5IDczMSw0MjcgNzI5LDQxNSA3MjEsNDE2IDY5Niw0MDMgNjg2LDM4MyA2NzQsMzc1IDY3NiwzNzIgNjczLDM3MXoiLz48cGF0aCBkPSJNMjE5LDY4MSAyMDIsNjkyIDIwMiw2OTcgMTk3LDY5OCAxOTgsNzAzIDE5Niw3MTAgMTkzLDcxMSAxOTMsNzE5IDE4OSw3MjEgMTkxLDcyMiAxOTEsNzM1IDE4OCw3NDQgMTk1LDczMyAxOTksNzM0IDE5Nyw3MzcgMjAwLDc0MCAyMDQsNzQwIDIwNSw3MzggMjEwLDczOCAyMTMsNzMyIDIxNyw3MzMgMjE4LDczMCAyMjMsNzMyIDIyNSw3MjggMjMwLDcyOCAyMjYsNzI0IDIyMiw3MzAgMjE2LDcyOCAyMTgsNzE5IDIxNiw3MTQgMjE5LDcxMSAyMTcsNzAwIDIyMCw2OTYgMjE3LDY5MXoiLz48cGF0aCBkPSJNMTM2MywxMzc4IDEzNTksMTM4MyAxMzYxLDEzOTUgMTM1MSwxNDA3IDEzNTQsMTQwOSAxMzUyLDE0MTkgMTM0OCwxNDIxIDEzNDcsMTQzNSAxMzM1LDE0NDggMTM0NSwxNDU4IDEzNDYsMTQ3NSAxMzQyLDE0ODIgMTM1MCwxNDg2IDEzNTEsMTQ5MSAxMzUzLDE0ODkgMTM1MiwxNDgyIDEzNTYsMTQ3NiAxMzY2LDE0NzIgMTM2NiwxNDY4IDEzNjMsMTQ2OCAxMzU5LDE0NjMgMTM2MSwxNDU2IDEzNTQsMTQ0NSAxMzUwLDE0NDQgMTM1MCwxNDM4IDEzNTMsMTQzNyAxMzUyLDE0MzIgMTM2MCwxNDE0IDEzNjEsMTM5MCAxMzY2LDEzODkgMTM2NCwxMzg4eiIvPjxwYXRoIGQ9Ik01OTcsMTI2MCA1OTUsMTI2NSA1OTAsMTI2OCA1OTAsMTI3NiA1OTQsMTI4MSA1OTIsMTI4NCA1OTUsMTI5MCA1OTEsMTI5OSA1OTcsMTMxMyA2MDYsMTMxOCA2MTIsMTMxOCA2MTMsMTMyMSA2MTYsMTMxNSA2MTMsMTMxNCA2MTMsMTMwNCA2MTgsMTI5NyA2MjMsMTI5NiA2MTQsMTI4NSA2MDMsMTI3N3oiLz48cGF0aCBkPSJNMTM3NSwxMTQzIDEzNjcsMTE0MyAxMzY0LDExMzcgMTM2MiwxMTQyIDEzNTYsMTE0NCAxMzUxLDExNDIgMTM1MCwxMTQ1IDEzMzcsMTE0OSAxMzI1LDExNTggMTMxNiwxMTcxIDEzMTIsMTE3MSAxMzA5LDExNjggMTMwNCwxMTc0IDEzMDMsMTE4MiAxMzAxLDExODMgMTMwOSwxMTkxIDEzMTAsMTE4MCAxMzEyLDExNzkgMTMwOSwxMTc2IDEzMTQsMTE3MiAxMzIwLDExNzAgMTMzNSwxMTcyIDEzNDcsMTE2MCAxMzUxLDExNjEgMTM1MiwxMTU4IDEzNTYsMTE1OSAxMzU3LDExNTYgMTM2OCwxMTUzIDEzNzIsMTE1NHoiLz48cGF0aCBkPSJNMTE5Niw5NTYgMTE5Myw5NTggMTE4Nyw5NTggMTE4Niw5NTYgMTE4NCw5NjEgMTE3NSw5NjggMTE3Miw5NjcgMTE3MCw5NzEgMTE1OCw5NjkgMTE0MSw5NzMgMTE0Miw5NzYgMTEzMSw5OTIgMTEzMSw5OTkgMTEzNSw5OTUgMTE0MCw5OTUgMTE0MSw5OTcgMTE1MSw5OTEgMTE1MCw5ODggMTE1NCw5ODggMTE1NSw5ODUgMTE1OSw5ODcgMTE2MCw5ODQgMTE3MCw5ODMgMTE3Niw5NzggMTE3OSw5NzkgMTE4MCw5NzUgMTE4Niw5NzQgMTE4NSw5NzEgMTE5NCw5NjMgMTE5Miw5NjB6Ii8+PHBhdGggZD0iTTk1Niw5NTkgOTU3LDk2MyA5NzcsOTc5IDk4Miw5ODYgOTkzLDk5MiA5OTgsOTkwIDEwMDMsOTgwIDEwMDMsOTc2IDk5OSw5NzEgOTkwLDk2OCA5ODksOTY1IDk4OCw5NjcgOTcyLDk2NSA5NjcsOTU5eiIvPjxwYXRoIGQ9Ik0xMDEwLDEwODMgMTAwMywxMDg5IDk5OSwxMDg4IDEwMDAsMTA4MyA5OTYsMTA4MyA5ODcsMTA5MyA5ODgsMTA5OCA5ODEsMTEwNCA5ODAsMTExMyA5ODIsMTExNiA5OTAsMTEyMCA5OTksMTEyMCAxMDA1LDExMDkgMTAwOCwxMTA4IDEwMDYsMTA5MiAxMDEyLDEwODV6Ii8+PHBhdGggZD0iTTg2Nyw2NTMgODY3LDY1NiA4NjEsNjYyIDg2NSw2NjMgODY2LDY2OSA4NjIsNjcwIDg2Myw2NzQgODYxLDY3NiA4NTgsNjc2IDg2NSw2NzYgODc1LDY4NSA4ODEsNjg0IDg4NCw2ODcgODg0LDY5MyA4NzgsNjk3IDg4MSw3MDEgODg0LDY5NiA4ODcsNjk2IDg5MSw3MDEgOTAwLDcwMCA4OTksNjkyIDg4Miw2NzkgODgxLDY2OSA4ODQsNjY1IDg3Nyw2NjN6Ii8+PHBhdGggZD0iTTEwMzEsMzAwIDEwNDksMzE3IDEwNTcsMzE4IDEwNTgsMzE1IDEwNjYsMzEzIDEwNjgsMzA2IDEwNzEsMzA1IDEwNjQsMjk4IDEwNTEsMjkzIDEwNDAsMjk0IDEwMzksMjkyeiIvPjxwYXRoIGQ9Ik04NzgsMTE3NSA4NzYsMTE3NCA4NzAsMTE3OSA4NjUsMTE3OSA4NTksMTE3MyA4NTYsMTE3NCA4NTksMTE3NSA4NjIsMTE4MSA4NjgsMTE4MyA4NjgsMTE4OCA4NjQsMTE5MiA4NTksMTE5MSA4NTUsMTE4NyA4NTYsMTE4NCA4NTMsMTE4MyA4NTMsMTE3OSA4NDYsMTE3OSA4NDMsMTE5NyA4NTQsMTE5NiA4NTUsMTE5OCA4NTYsMTE5NSA4NjEsMTE5NiA4NjMsMTE5NCA4NjgsMTE5NCA4NzAsMTE5NyA4NzMsMTE5MyA4NzgsMTE5NCA4NzksMTE4NiA4ODEsMTE4NSA4NzcsMTE4MnoiLz48cGF0aCBkPSJNODQ5LDk1OCA4NDUsOTU5IDg0Miw5NTYgODQyLDk2MiA4MzgsOTYzIDgzNiw5NjggODMzLDk2OSA4MzMsOTc5IDgzNSw5ODAgODM0LDk5NSA4NTIsOTcyIDg1NCw5NjJ6Ii8+PHBhdGggZD0iTTEwMTIsMTE2OCAxMDA5LDExNzEgMTAwOCwxMTgxIDEwMTIsMTE4MiAxMDI0LDExOTcgMTAyOCwxMTk3IDEwMjgsMTE4OCAxMDMyLDExODIgMTAyMiwxMTc4IDEwMTksMTE3NSAxMDE5LDExNjl6Ii8+PHBhdGggZD0iTTExNDgsMTEzOCAxMTM5LDExMzIgMTEzMiwxMTMzIDExMzQsMTE0MCAxMTIzLDExNDggMTExMiwxMTUyIDExMTQsMTE1OCAxMTA3LDExNjMgMTEwMCwxMTYzIDEwOTksMTE2MSAxMDk3LDExNjUgMTA5MywxMTY0IDEwOTQsMTE2NyAxMDkyLDExNjggMTEwNSwxMTcwIDExMDcsMTE2NiAxMTE0LDExNjQgMTEyNiwxMTQ5IDExMjksMTE1MCAxMTMyLDExNDN6Ii8+PHBhdGggZD0iTTg2MCwxMDA2IDg0OCwxMDE4IDg1MywxMDE4IDg1NSwxMDIzIDg1NSwxMDMyIDg1MSwxMDM3IDg1MCwxMDQ1IDg1MywxMDQyIDg1NywxMDQyIDg1NywxMDM4IDg2MywxMDM0IDg2MiwxMDI4IDg2NSwxMDI3IDg2NCwxMDI0IDg2NywxMDIxIDg1OCwxMDEzeiIvPjxwYXRoIGQ9Ik0xMjk4LDExMDAgMTI5NywxMTAzIDEzMDAsMTEwNiAxMzA3LDExMDggMTMwNSwxMTExIDEzMDcsMTExNSAxMzA0LDExMjEgMTMwMSwxMTIyIDEzMDMsMTEyMyAxMzA1LDExMzQgMTMxMiwxMTMzIDEzMTMsMTEyOSAxMzIwLDExMjcgMTMxOSwxMTIwIDEzMjIsMTExOSAxMzE5LDExMTYgMTMxMywxMTE1IDEzMDgsMTExMCAxMzEwLDExMDcgMTI5OSwxMTAzeiIvPjxwYXRoIGQ9Ik03MTYsNTk5IDcxNCw2MDEgNzE3LDYwMiA3MTYsNjA1IDcyNiw2MTUgNzI1LDYyMSA3MjcsNjI2IDczMyw2MjUgNzM2LDYyOSA3NDAsNjMwIDczOSw2MjcgNzQ2LDYyMiA3NDcsNjE4IDczOCw2MjAgNzM2LDYxOCA3MzcsNjEyIDczNCw2MTMgNzI3LDYxMHoiLz48cGF0aCBkPSJNNjkzLDg4MSA2OTQsODg2IDY5Miw4ODggNjg5LDg4NyA2ODcsODkxIDY5Myw4OTIgNjk0LDg4OSA2OTcsODg5IDY5OCw4OTMgNjkyLDkwMCA2ODIsOTAwIDY4Miw5MDcgNjg2LDkwNiA2OTEsOTExIDY5Myw5MDYgNjk5LDkwMyA3MDUsODk0IDcwMSw4OTAgNzAyLDg4NXoiLz48cGF0aCBkPSJNOTQxLDk5OSA5NDQsMTAwMyA5NDMsMTAxMyA5MzUsMTAxOSA5NDAsMTAyNCA5NDAsMTAzMCA5NTMsMTAyMyA5NTIsMTAxNSA5NDksMTAxNCA5NDksMTAwM3oiLz48cGF0aCBkPSJNMTAxNCwxMDE4IDEwMjAsMTAxNyAxMDI2LDEwMjAgMTAzOSwxMDA3IDEwMzcsMTAwNCAxMDM0LDEwMDYgMTAyOCwxMDAyIDEwMjgsMTAwNyAxMDI0LDEwMDggMTAyMSwxMDAwIDEwMTcsMTAwMSAxMDE2LDEwMDUgMTAyMCwxMDEzeiIvPjxwYXRoIGQ9Ik03MDYsOTM3IDY5NSw5MzYgNjg1LDkzOSA2ODksOTQ4IDY5Miw5NDkgNjkxLDk1NCA2OTQsOTU1IDY5OSw5NTIgNzAxLDk0NSA3MDUsOTQzeiIvPjxwYXRoIGQ9Ik04NzEsMTA2NyA4NzAsMTA3MCA4NjcsMTA2OSA4NjEsMTA4MCA4NjEsMTA4NSA4NjYsMTA4NCA4NjgsMTA4OCA4NzMsMTA4OCA4NzcsMTA4MiA4NzYsMTA3MSA4NzIsMTA3MHoiLz48cGF0aCBkPSJNODk5LDExMTAgOTA0LDExMTcgOTAzLDExMjIgOTA4LDExMzIgOTA3LDExMzggOTExLDExMzcgOTE0LDExMzEgOTIwLDExMjkgOTE4LDExMjQgOTEwLDExMTcgOTExLDExMTR6Ii8+PHBhdGggZD0iTTEwMzAsMTIzOCAxMDI4LDEyNDAgMTAzMSwxMjQ1IDEwMzIsMTI0MyAxMDM4LDEyNDQgMTA0MSwxMjYyIDEwMzcsMTI2OCAxMDI5LDEyNjggMTA0NCwxMjcxIDEwNDgsMTI2NiAxMDQ4LDEyNTcgMTAzOSwxMjQ2IDEwNDAsMTI0M3oiLz48cGF0aCBkPSJNMTMyNywxNDg2IDEzMjIsMTQ4NSAxMzIxLDE0ODggMTMxMiwxNDg5IDEzMDcsMTQ5MiAxMzA0LDE0OTcgMTMwMSwxNDk3IDEzMDIsMTUwMyAxMzA2LDE1MDUgMTMxMiwxNDk4IDEzMjUsMTQ5MnoiLz48cGF0aCBkPSJNMTE4OSwxMjEwIDExODgsMTIxNyAxMTkxLDEyMjYgMTE5NCwxMjI5IDEyMDEsMTIzMCAxMjAzLDEyMTggMTIwMCwxMjE1IDExOTYsMTIxNiAxMTkwLDEyMTN6Ii8+PHBhdGggZD0iTTc3OCw1ODIgNzc3LDU4NCA3NzMsNTgyIDc3MSw1OTMgNzY0LDU5OSA3NjUsNjA2IDc1Niw2MTIgNzYwLDYxMyA3NjgsNjA1IDc3MSw2MDYgNzcwLDYwMiA3NzMsNjAxIDc3Myw1OTUgNzc1LDU5MiA3NzgsNTkyeiIvPjxwYXRoIGQ9Ik0xMzcxLDEyNDcgMTM3NCwxMjUwIDEzNzIsMTI1MyAxMzc2LDEyNTQgMTM4MSwxMjQ5IDEzODQsMTI1MyAxMzgzLDEyNTYgMTM4OSwxMjYxIDEzOTQsMTI1OSAxMzg3LDEyNTMgMTM4MywxMjQyeiIvPjxwYXRoIGQ9Ik0xNDYyLDEyMDMgMTQ1OCwxMjA5IDE0NTUsMTIwOSAxNDU1LDEyMTIgMTQ2MCwxMjEzIDE0NjksMTIyMyAxNDcxLDEyMDcgMTQ2MywxMjA3eiIvPjxwYXRoIGQ9Ik03MDIsNjA1IDY5OSw2MDQgNzAwLDYwMSA2OTcsNjAyIDY5Miw1OTkgNjg5LDYwNiA2NzksNjExIDY5MCw2MTIgNjkxLDYxNCA2OTQsNjA4IDcwMiw2MDd6Ii8+PHBhdGggZD0iTTIwNyw3NjQgMjA2LDc3MCAxOTksNzcyIDIwNCw3NzggMjA0LDc4NiAyMDksNzkyIDIwOCw3ODggMjEyLDc3OSAyMDYsNzc0eiIvPjxwYXRoIGQ9Ik0xMzQxLDEyMDggMTM0MywxMjE4IDEzNTAsMTIyMSAxMzUxLDEyMTggMTM1NSwxMjE3IDEzNTEsMTIxMSAxMzUyLDEyMDh6Ii8+PHBhdGggZD0iTTEyNzksMTA3MiAxMjc0LDEwNzYgMTI3NywxMDgwIDEyODIsMTA4MCAxMjgxLDEwODMgMTI4OCwxMDkwIDEyODksMTA5NSAxMjk0LDEwOTAgMTI4NywxMDgzIDEyODUsMTA3NnoiLz48cGF0aCBkPSJNOTkyLDExODUgOTg0LDExODUgOTc5LDExOTUgOTc1LDExOTYgOTc1LDEyMDEgOTg4LDExOTV6Ii8+PHBhdGggZD0iTTg4OSwxMTU4IDg4NCwxMTU1IDg4MSwxMTU4IDg3NywxMTU4IDg3OCwxMTYxIDg3NiwxMTYyIDg3OSwxMTYzIDg3OSwxMTY5IDg4NiwxMTY3eiIvPjxwYXRoIGQ9Ik0xMDUxLDc0MiAxMDU2LDc0OSAxMDU1LDc1MyAxMDYxLDc1NSAxMDYzLDc0NSAxMDU5LDc0MXoiLz48cGF0aCBkPSJNMTA5NCwxMjY1IDExMDAsMTI3MSAxMTAzLDEyNzEgMTEwNCwxMjY4IDExMDksMTI3MCAxMTExLDEyNjUgMTEwOCwxMjY2IDExMDAsMTI1OXoiLz48cGF0aCBkPSJNMjE4LDc5MyAyMTYsNzkzIDIxNiw3OTggMjA4LDc5NyAyMTQsODA5IDIxOSw4MDh6Ii8+PHBhdGggZD0iTTEyMDksMTIwMiAxMjEwLDEyMTQgMTIxMiwxMjE2IDEyMTMsMTIxMyAxMjIxLDEyMTMgMTIxOSwxMjA2eiIvPjxwYXRoIGQ9Ik05NTcsNTE4IDk1Miw1MjEgOTUyLDUzMyA5NjIsNTI0eiIvPjxwYXRoIGQ9Ik05NzQsMTEwMiA5NzQsMTEwNSA5NjYsMTExNSA5NzEsMTEyMyA5NzIsMTExNyA5NzUsMTExNiA5NzMsMTExNSA5NzMsMTEwOCA5NzYsMTEwN3oiLz48cGF0aCBkPSJNOTMzLDExOTcgOTM3LDEyMDEgOTQyLDEyMDIgOTQ5LDEyMTEgOTUzLDEyMTAgOTUzLDEyMDQgOTQ0LDEyMDMgOTQxLDExOTd6Ii8+PHBhdGggZD0iTTcwNywxMDM1IDcwNiwxMDMzIDY5OCwxMDMzIDY3NSwxMDQ1IDY4MCwxMDQ1IDY4NCwxMDQyIDY4NywxMDQ0IDY4OCwxMDQxIDY5NCwxMDQwIDY5NiwxMDM2IDcwMSwxMDM0eiIvPjxwYXRoIGQ9Ik0xNDI5LDEyOTYgMTQxNCwxMjk2IDE0MTQsMTI5OSAxNDExLDEzMDAgMTQyNiwxMzAwIDE0MjcsMTMwMnoiLz48cGF0aCBkPSJNMTIzNSwxMDI5IDEyMjksMTAzNiAxMjMyLDEwMzcgMTIzMiwxMDQxIDEyMjksMTA0MiAxMjMxLDEwNDUgMTIyOSwxMDQ2IDEyMzIsMTA0NyAxMjM0LDEwNTIgMTIzNCwxMDQ3IDEyMzYsMTA0NiAxMjMzLDEwNDUgMTIzMywxMDM2IDEyNDAsMTAzNXoiLz48cGF0aCBkPSJNNzk5LDE2MTYgNzk1LDE2MjEgODA1LDE2MzAgODAzLDE2MjcgODA1LDE2MjF6Ii8+PHBhdGggZD0iTTEyMiw1OTIgMTI4LDYwNSAxMzMsNjAzeiIvPjwvZz48L3N2Zz4="""


from position_tracking_settings import (
    AppConfig as _AppConfig, _app_dir, _open_path, _fallback_settings_path,
    _runtime_root as _settings_runtime_root,
)


def _exe_dir() -> Path:
    return _app_dir()


def _runtime_root() -> Path:
    return _settings_runtime_root(_app_dir())


def _settings_path() -> Path:
    return _local_root() / SETTINGS_NAME


class AppConfig(_AppConfig):
    """Compatibility facade preserving legacy overridable path providers."""

    @classmethod
    def load(cls) -> "AppConfig":
        return super().load(path=_settings_path(), app_directory=_app_dir())

    def save(self) -> Path:
        return super().save(path=_settings_path())




from position_tracking_personnel import PersonnelRecord, PersonnelRepository, ShiftRosterRepository


class BusyError(RuntimeError):
    pass


class SharedEventStore:
    def __init__(self, central_root: str | Path):
        root = Path(central_root)
        self.root = root if root.name == ".tde_runtime" else (root if root.name.upper() == "TDE_DATA" else root / "TDE_DATA")
        self.events = self.root / "events"
        self.active = self.root / "active"
        self.locks = self.root / "locks"
        self.reports = self.root / "reports"
        self.shifts = self.root / "shifts"
        self.audit = self.root / "audit"
        self.backups = self.root / "backups"
        self.current_docs = self.root / "current"
        self.archive = self.root / "archive"
        self.statistics = self.root / "statistics"
        self.communications = self.root / "communications"
        self.comm_clients = self.communications / "clients"
        self.comm_messages = self.communications / "messages"
        self.comm_acks = self.communications / "acks"
        self.comm_commands = self.communications / "commands"
        self.comm_command_acks = self.communications / "command_acks"
        self.shift_db = self.root / "shift_reports.sqlite3"
        self.pending_index = self.root / "pending_sessions"
        self.corrections = self.root / "start_corrections"
        for p in (
            self.events, self.active, self.locks, self.reports, self.shifts, self.audit,
            self.backups, self.current_docs, self.archive, self.statistics,
            self.communications, self.comm_clients, self.comm_messages, self.comm_acks,
            self.comm_commands, self.comm_command_acks, self.pending_index, self.corrections,
        ):
            p.mkdir(parents=True, exist_ok=True)
        self._init_shift_db()
        # Recover END events left unindexed after a process or SQLite interruption.
        try:
            self.reconcile_pending_sessions()
        except Exception:
            pass

    def _active_path(self, agms: str) -> Path:
        return self.active / f"{_safe_name(agms)}.json"

    def _lock_path(self, agms: str) -> Path:
        return self.locks / f"{_safe_name(agms)}.lock"

    @staticmethod
    @contextmanager
    def _file_lock(path: Path, timeout: float, busy_message: str):
        """OS-held byte-range lock. The lock file stays in place on SMB.

        Unlike stale-file mtime locks, Windows releases this lock when the
        owning process dies or its SMB handle is disconnected. Never unlink
        another process's lock based on its age.
        """
        import errno
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            fd = os.open(path, os.O_RDWR | os.O_CREAT, 0o666)
        except OSError as exc:
            raise BusyError(f"Δεν είναι διαθέσιμος ο κοινόχρηστος φάκελος: {exc}") from exc
        locked = False
        try:
            # Never write a bootstrap byte BEFORE acquiring the OS lock.
            # On Windows/SMB a concurrent owner may already lock byte 0;
            # os.write(fd, b"\x00") then fails with EACCES and was treated as
            # a fatal error rather than normal lock contention. Both Windows
            # byte-range locks and POSIX flock support empty lock files.
            deadline = time.monotonic() + timeout
            while True:
                try:
                    os.lseek(fd, 0, os.SEEK_SET)
                    if os.name == "nt":
                        import msvcrt
                        msvcrt.locking(fd, msvcrt.LK_NBLCK, 1)
                    else:
                        import fcntl
                        fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
                    locked = True
                    break
                except OSError as exc:
                    contention = exc.errno in (errno.EACCES, errno.EAGAIN) or getattr(exc, "winerror", None) in (32, 33, 36)
                    if not contention:
                        raise BusyError(f"Αποτυχία κοινόχρηστου κλειδώματος: {exc}") from exc
                    if time.monotonic() >= deadline:
                        raise BusyError(busy_message) from exc
                    time.sleep(0.15)
            yield
        finally:
            if locked:
                try:
                    os.lseek(fd, 0, os.SEEK_SET)
                    if os.name == "nt":
                        import msvcrt
                        msvcrt.locking(fd, msvcrt.LK_UNLCK, 1)
                    else:
                        import fcntl
                        fcntl.flock(fd, fcntl.LOCK_UN)
                finally:
                    os.close(fd)
            else:
                os.close(fd)

    @contextmanager
    def officer_lock(self, agms: str, timeout: float = 6.0):
        with self._file_lock(
            self._lock_path(agms), timeout,
            "Η καταχώριση χρησιμοποιείται από άλλο σταθμό. Δοκιμάστε ξανά."
        ):
            yield

    @staticmethod
    def _atomic_json(path: Path, data: dict):
        path.parent.mkdir(parents=True, exist_ok=True)
        temp = path.with_name(path.name + f".{uuid.uuid4().hex}.tmp")
        temp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(temp, path)


    @staticmethod
    def local_client_id() -> str:
        computer = _safe_name(socket.gethostname())
        user = _safe_name(os.environ.get("USERNAME") or os.environ.get("USER") or "USER")
        # Distinguish two Windows sessions/processes sharing the same user.
        return f"{computer}__{user}__{os.getpid()}"

    def write_client_heartbeat(self, client_id: str, extra: dict | None = None) -> dict:
        payload = {
            "client_id": _clean(client_id),
            "computer": socket.gethostname(),
            "user": os.environ.get("USERNAME") or os.environ.get("USER") or "",
            "pid": os.getpid(),
            "last_seen": datetime.now().isoformat(timespec="seconds"),
        }
        if extra:
            payload.update(extra)
        self._atomic_json(self.comm_clients / f"{_safe_name(client_id)}.json", payload)
        return payload

    def remove_client_heartbeat(self, client_id: str):
        (self.comm_clients / f"{_safe_name(client_id)}.json").unlink(missing_ok=True)

    def online_clients(self, max_age_seconds: int = 15) -> list[dict]:
        now = datetime.now()
        result = []
        for path in self.comm_clients.glob("*.json"):
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                seen = datetime.fromisoformat(_clean(data.get("last_seen")))
                if (now - seen).total_seconds() <= max(3, int(max_age_seconds)):
                    result.append(data)
            except Exception:
                continue
        return sorted(result, key=lambda x: (_clean(x.get("computer")), _clean(x.get("user"))))

    def send_message(
        self,
        sender: dict,
        target_client_ids: list[str],
        text: str,
        priority: str = "ΟΔΗΓΙΑ",
        target_label: str = "",
    ) -> dict:
        targets = sorted({_clean(x) for x in target_client_ids if _clean(x)})
        text = _clean(text)
        if not targets:
            raise ValueError("Δεν βρέθηκε ενεργός παραλήπτης.")
        if not text:
            raise ValueError("Γράψτε το μήνυμα / οδηγία.")
        message_id = uuid.uuid4().hex
        payload = {
            "message_id": message_id,
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "sender": sender or {},
            "targets": targets,
            "target_label": _clean(target_label),
            "priority": _clean(priority).upper() or "ΟΔΗΓΙΑ",
            "text": text,
        }
        self._atomic_json(self.comm_messages / f"{message_id}.json", payload)
        self.log_admin_action("MESSAGE_SENT", {
            "message_id": message_id, "targets": targets, "priority": payload["priority"],
            "target_label": payload["target_label"],
        })
        return payload

    def acknowledge_message(self, message_id: str, client_id: str):
        folder = self.comm_acks / _safe_name(message_id)
        payload = {
            "message_id": _clean(message_id),
            "client_id": _clean(client_id),
            "ack_at": datetime.now().isoformat(timespec="seconds"),
            "computer": socket.gethostname(),
            "user": os.environ.get("USERNAME") or os.environ.get("USER") or "",
        }
        self._atomic_json(folder / f"{_safe_name(client_id)}.json", payload)

    def pending_messages(self, client_id: str, max_age_days: int = 7) -> list[dict]:
        client_id = _clean(client_id)
        cutoff = datetime.now() - timedelta(days=max(1, int(max_age_days)))
        result = []
        for path in sorted(self.comm_messages.glob("*.json")):
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                if client_id not in [_clean(x) for x in data.get("targets", [])]:
                    continue
                created = datetime.fromisoformat(_clean(data.get("created_at")))
                if created < cutoff:
                    continue
                ack = self.comm_acks / _safe_name(data.get("message_id")) / f"{_safe_name(client_id)}.json"
                if not ack.exists():
                    result.append(data)
            except Exception:
                continue
        return sorted(result, key=lambda x: _clean(x.get("created_at")))

    def message_acknowledgements(self, message_id: str) -> list[dict]:
        folder = self.comm_acks / _safe_name(message_id)
        result = []
        if folder.exists():
            for path in folder.glob("*.json"):
                try:
                    result.append(json.loads(path.read_text(encoding="utf-8")))
                except Exception:
                    pass
        return result

    def recent_sent_messages(self, limit: int = 20) -> list[dict]:
        result = []
        for path in sorted(self.comm_messages.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                data["acknowledgements"] = self.message_acknowledgements(_clean(data.get("message_id")))
                result.append(data)
            except Exception:
                continue
            if len(result) >= max(1, int(limit)):
                break
        return result

    def send_close_command(self, sender: dict, target_client_id: str, reason: str = "") -> dict:
        target_client_id = _clean(target_client_id)
        if not target_client_id:
            raise ValueError("Δεν επιλέχθηκε τερματικό.")
        command_id = uuid.uuid4().hex
        payload = {
            "command_id": command_id,
            "type": "CLOSE_APP",
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "sender": sender or {},
            "target_client_id": target_client_id,
            "reason": _clean(reason),
        }
        self._atomic_json(self.comm_commands / f"{command_id}.json", payload)
        self.log_admin_action("REMOTE_CLOSE_SENT", {
            "command_id": command_id, "target_client_id": target_client_id, "reason": _clean(reason)
        })
        return payload

    def pending_commands(self, client_id: str) -> list[dict]:
        client_id = _clean(client_id)
        result = []
        for path in sorted(self.comm_commands.glob("*.json")):
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                if _clean(data.get("target_client_id")) != client_id:
                    continue
                ack = self.comm_command_acks / f"{_safe_name(data.get('command_id'))}.json"
                if not ack.exists():
                    result.append(data)
            except Exception:
                continue
        return sorted(result, key=lambda x: _clean(x.get("created_at")))

    def acknowledge_command(self, command_id: str, client_id: str):
        payload = {
            "command_id": _clean(command_id),
            "client_id": _clean(client_id),
            "ack_at": datetime.now().isoformat(timespec="seconds"),
            "computer": socket.gethostname(),
            "user": os.environ.get("USERNAME") or os.environ.get("USER") or "",
        }
        self._atomic_json(self.comm_command_acks / f"{_safe_name(command_id)}.json", payload)

    @contextmanager
    def _shift_db_write_lock(self, timeout: float = 10.0):
        with self._file_lock(
            self.locks / "shift_reports.sqlite.lock", timeout,
            "Το αρχείο βαρδιών χρησιμοποιείται από άλλο σταθμό. Δοκιμάστε ξανά."
        ):
            yield

    @contextmanager
    def _shift_db_connect(self):
        con = sqlite3.connect(str(self.shift_db), timeout=10.0)
        try:
            con.row_factory = sqlite3.Row
            con.execute("PRAGMA journal_mode=DELETE")
            con.execute("PRAGMA synchronous=FULL")
            con.execute("PRAGMA busy_timeout=10000")
            yield con
            con.commit()
        except Exception:
            con.rollback()
            raise
        finally:
            con.close()

    @staticmethod
    def _upsert_shift_con(con, shift: dict):
        shift_id = _clean(shift.get("shift_id"))
        if not shift_id:
            raise ValueError("Μη έγκυρη βάρδια.")
        status = _clean(shift.get("status")).upper() or "DRAFT"
        start_at = _clean(shift.get("start_at"))
        end_at = _clean(shift.get("end_at")) or None
        payload = json.dumps(shift, ensure_ascii=False, separators=(",", ":"))
        con.execute(
            """
            INSERT INTO shifts (shift_id, status, start_at, end_at, data_json, updated_at)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(shift_id) DO UPDATE SET
                status=excluded.status,
                start_at=excluded.start_at,
                end_at=excluded.end_at,
                data_json=excluded.data_json,
                updated_at=excluded.updated_at
            """,
            (shift_id, status, start_at, end_at, payload, datetime.now().isoformat(timespec="seconds")),
        )

    @staticmethod
    def _upsert_personnel_session_con(con, start: dict, end: dict):
        start_event_id = _clean(start.get("event_id"))
        end_event_id = _clean(end.get("event_id"))
        if not start_event_id or not end_event_id:
            return
        try:
            start_dt = datetime.fromisoformat(_clean(start.get("timestamp")))
            end_dt = datetime.fromisoformat(_clean(end.get("timestamp")))
        except Exception:
            return
        if end_dt < start_dt:
            return
        agms = _clean(start.get("agms") or end.get("agms"))
        rank = _clean(start.get("rank") or end.get("rank"))
        surname = _clean(start.get("surname") or end.get("surname"))
        given_name = _clean(start.get("given_name") or end.get("given_name"))
        full_name = _clean(start.get("full_name") or end.get("full_name") or f"{surname} {given_name}")
        person_search = _norm_search(f"{agms} {rank} {surname} {given_name} {full_name}")
        duration_seconds = max(0, int((end_dt - start_dt).total_seconds()))
        con.execute(
            """
            INSERT INTO personnel_sessions (
                session_id, start_event_id, end_event_id, agms, rank, surname, given_name,
                full_name, person_search, position, direction, start_at, end_at,
                duration_seconds, start_computer, start_user, end_computer, end_user
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(start_event_id) DO UPDATE SET
                end_event_id=excluded.end_event_id,
                agms=excluded.agms,
                rank=excluded.rank,
                surname=excluded.surname,
                given_name=excluded.given_name,
                full_name=excluded.full_name,
                person_search=excluded.person_search,
                position=excluded.position,
                direction=excluded.direction,
                start_at=excluded.start_at,
                end_at=excluded.end_at,
                duration_seconds=excluded.duration_seconds,
                start_computer=excluded.start_computer,
                start_user=excluded.start_user,
                end_computer=excluded.end_computer,
                end_user=excluded.end_user
            """,
            (
                start_event_id, start_event_id, end_event_id, agms, rank, surname, given_name,
                full_name, person_search, _clean(start.get("position")),
                _clean(start.get("direction")).upper(),
                start_dt.isoformat(timespec="seconds"), end_dt.isoformat(timespec="seconds"),
                duration_seconds, _clean(start.get("computer")), _clean(start.get("user")),
                _clean(end.get("computer")), _clean(end.get("user")),
            ),
        )

    def _store_personnel_session(self, start: dict, end: dict):
        with self._shift_db_write_lock():
            with self._shift_db_connect() as con:
                self._upsert_personnel_session_con(con, start, end)

    def _has_end_receipt(self, active: dict) -> bool:
        start_id = _safe_name(active.get("event_id"))
        receipt = self.pending_index / f"{start_id}.json"
        if not receipt.is_file():
            return False
        try:
            data = json.loads(receipt.read_text(encoding="utf-8"))
            return _clean(data.get("start", {}).get("event_id")) == _clean(active.get("event_id"))
        except Exception as exc:
            raise BusyError(f"Αδύνατος έλεγχος προσωρινής λήξης {start_id}: {exc}") from exc

    def reconcile_pending_sessions(self) -> int:
        """Recover committed END events after a crash, without dropping history."""
        done = 0
        for path in sorted(self.pending_index.glob("*.json")):
            data = json.loads(path.read_text(encoding="utf-8"))
            start, end = data["start"], data["end"]
            agms = _clean(start.get("agms"))
            with self.officer_lock(agms):
                active_path = self._active_path(agms)
                if active_path.is_file():
                    current = json.loads(active_path.read_text(encoding="utf-8"))
                    if _clean(current.get("event_id")) == _clean(start.get("event_id")):
                        active_path.unlink()
                self._store_personnel_session(start, end)
                path.unlink(missing_ok=True)
                done += 1
        return done

    @staticmethod
    def _session_row_dict(row) -> dict:
        return {
            "session_id": _clean(row["session_id"]),
            "start_event_id": _clean(row["start_event_id"]),
            "end_event_id": _clean(row["end_event_id"]),
            "agms": _clean(row["agms"]),
            "rank": _clean(row["rank"]),
            "surname": _clean(row["surname"]),
            "given_name": _clean(row["given_name"]),
            "full_name": _clean(row["full_name"]),
            "position": _clean(row["position"]),
            "direction": _clean(row["direction"]).upper(),
            "start_at": _clean(row["start_at"]),
            "end_at": _clean(row["end_at"]),
            "duration_seconds": int(row["duration_seconds"] or 0),
        }

    @staticmethod
    def _session_overlaps_clock_window(start_dt: datetime, end_dt: datetime, time_from: dtime, time_to: dtime) -> bool:
        day = start_dt.date() - timedelta(days=1)
        last = end_dt.date()
        while day <= last:
            window_start = datetime.combine(day, time_from)
            window_end = datetime.combine(day, time_to)
            if time_to <= time_from:
                window_end += timedelta(days=1)
            if start_dt < window_end and end_dt > window_start:
                return True
            day += timedelta(days=1)
        return False

    def search_personnel_sessions(
        self,
        person_query: str = "",
        date_from: date | None = None,
        date_to: date | None = None,
        position: str = "",
        direction: str = "",
        time_from: dtime | None = None,
        time_to: dtime | None = None,
        moment: datetime | None = None,
        min_duration_minutes: int = 0,
        max_duration_minutes: int = 0,
        limit: int | None = None,
    ) -> list[dict]:
        clauses = ["1=1"]
        params: list[object] = []
        query = _norm_search(person_query)
        if query:
            clauses.append("person_search LIKE ?")
            params.append(f"%{query}%")
        if date_from is not None:
            clauses.append("end_at>=?")
            params.append(f"{date_from:%Y-%m-%d}T00:00:00")
        if date_to is not None:
            clauses.append("start_at<=?")
            params.append(f"{date_to:%Y-%m-%d}T23:59:59")
        position = _clean(position)
        if position and position != "ΟΛΕΣ":
            clauses.append("position=?")
            params.append(position)
        direction = _clean(direction).upper()
        if direction and direction != "ΟΛΑ":
            clauses.append("direction=?")
            params.append(direction)
        if moment is not None:
            stamp = moment.isoformat(timespec="seconds")
            clauses.append("start_at<=?")
            clauses.append("end_at>?")
            params.extend([stamp, stamp])
        if int(min_duration_minutes or 0) > 0:
            clauses.append("duration_seconds>=?")
            params.append(int(min_duration_minutes) * 60)
        if int(max_duration_minutes or 0) > 0:
            clauses.append("duration_seconds<=?")
            params.append(int(max_duration_minutes) * 60)

        sql = "SELECT * FROM personnel_sessions WHERE " + " AND ".join(clauses) + " ORDER BY start_at DESC"
        if limit is not None and int(limit) > 0 and time_from is None and time_to is None:
            sql += " LIMIT ?"
            params.append(int(limit))

        rows: list[dict] = []
        try:
            with self._shift_db_connect() as con:
                db_rows = con.execute(sql, params).fetchall()
            for row in db_rows:
                item = self._session_row_dict(row)
                if time_from is not None and time_to is not None:
                    try:
                        sdt = datetime.fromisoformat(item["start_at"])
                        edt = datetime.fromisoformat(item["end_at"])
                    except Exception:
                        continue
                    if not self._session_overlaps_clock_window(sdt, edt, time_from, time_to):
                        continue
                rows.append(item)
                if limit is not None and int(limit) > 0 and len(rows) >= int(limit):
                    break
        except Exception:
            return []
        return rows

    def list_personnel_workers(self) -> list[dict]:
        result: list[dict] = []
        seen: set[str] = set()
        try:
            with self._shift_db_connect() as con:
                rows = con.execute(
                    """
                    SELECT agms, rank, surname, given_name, full_name
                    FROM personnel_sessions
                    ORDER BY start_at DESC
                    """
                ).fetchall()
            for row in rows:
                agms = _clean(row["agms"])
                if not agms or agms in seen:
                    continue
                seen.add(agms)
                result.append({
                    "agms": agms,
                    "rank": _clean(row["rank"]),
                    "surname": _clean(row["surname"]),
                    "given_name": _clean(row["given_name"]),
                    "full_name": _clean(row["full_name"]),
                })
        except Exception:
            return []
        return sorted(result, key=lambda x: (_norm_search(x.get("surname")), _norm_search(x.get("given_name")), x.get("agms", "")))

    def list_session_positions(self) -> list[str]:
        try:
            with self._shift_db_connect() as con:
                rows = con.execute(
                    "SELECT DISTINCT position FROM personnel_sessions WHERE position<>'' ORDER BY position"
                ).fetchall()
            return [_clean(row["position"]) for row in rows if _clean(row["position"])]
        except Exception:
            return []

    def _init_shift_db(self):
        self.shift_db.parent.mkdir(parents=True, exist_ok=True)
        with self._shift_db_write_lock():
            with self._shift_db_connect() as con:
                con.execute(
                    """
                    CREATE TABLE IF NOT EXISTS shifts (
                        shift_id TEXT PRIMARY KEY,
                        status TEXT NOT NULL,
                        start_at TEXT NOT NULL,
                        end_at TEXT,
                        data_json TEXT NOT NULL,
                        updated_at TEXT NOT NULL
                    )
                    """
                )
                con.execute("CREATE INDEX IF NOT EXISTS idx_shifts_start_at ON shifts(start_at DESC)")
                con.execute("CREATE INDEX IF NOT EXISTS idx_shifts_status ON shifts(status)")
                con.execute(
                    """
                    CREATE TABLE IF NOT EXISTS meta (
                        key TEXT PRIMARY KEY,
                        value TEXT
                    )
                    """
                )
                con.execute(
                    """
                    CREATE TABLE IF NOT EXISTS personnel_sessions (
                        session_id TEXT PRIMARY KEY,
                        start_event_id TEXT UNIQUE NOT NULL,
                        end_event_id TEXT NOT NULL,
                        agms TEXT NOT NULL,
                        rank TEXT,
                        surname TEXT,
                        given_name TEXT,
                        full_name TEXT,
                        person_search TEXT NOT NULL,
                        position TEXT NOT NULL,
                        direction TEXT NOT NULL,
                        start_at TEXT NOT NULL,
                        end_at TEXT NOT NULL,
                        duration_seconds INTEGER NOT NULL,
                        start_computer TEXT,
                        start_user TEXT,
                        end_computer TEXT,
                        end_user TEXT
                    )
                    """
                )
                con.execute("CREATE INDEX IF NOT EXISTS idx_personnel_sessions_start ON personnel_sessions(start_at DESC)")
                con.execute("CREATE INDEX IF NOT EXISTS idx_personnel_sessions_end ON personnel_sessions(end_at DESC)")
                con.execute("CREATE INDEX IF NOT EXISTS idx_personnel_sessions_person ON personnel_sessions(person_search)")
                con.execute("CREATE INDEX IF NOT EXISTS idx_personnel_sessions_position ON personnel_sessions(position, direction)")
                migrated = con.execute(
                    "SELECT value FROM meta WHERE key='legacy_shift_json_migrated_v1'"
                ).fetchone()
                if not migrated:
                    newest = None
                    for path in sorted(self.shifts.glob("SHIFT_*.json")):
                        try:
                            data = json.loads(path.read_text(encoding="utf-8"))
                            if _clean(data.get("shift_id")):
                                self._upsert_shift_con(con, data)
                                newest = data
                        except Exception:
                            continue
                    current_path = self._current_shift_path()
                    try:
                        current = json.loads(current_path.read_text(encoding="utf-8"))
                        if _clean(current.get("shift_id")):
                            self._upsert_shift_con(con, current)
                            newest = current
                    except Exception:
                        pass
                    if newest and _clean(newest.get("shift_id")):
                        con.execute(
                            "INSERT OR REPLACE INTO meta(key,value) VALUES('current_report_shift_id',?)",
                            (_clean(newest.get("shift_id")),),
                        )
                    con.execute(
                        "INSERT OR REPLACE INTO meta(key,value) VALUES('legacy_shift_json_migrated_v1','1')"
                    )

                personnel_migrated = con.execute(
                    "SELECT value FROM meta WHERE key='legacy_personnel_sessions_migrated_v1'"
                ).fetchone()
                if not personnel_migrated:
                    starts: dict[str, dict] = {}
                    for event in self.iter_all_events():
                        kind = _clean(event.get("kind")).upper()
                        if kind == "START":
                            starts[_clean(event.get("event_id"))] = event
                        elif kind == "END":
                            start = starts.get(_clean(event.get("start_event_id")))
                            if start:
                                self._upsert_personnel_session_con(con, start, event)
                    con.execute(
                        "INSERT OR REPLACE INTO meta(key,value) VALUES('legacy_personnel_sessions_migrated_v1','1')"
                    )

    def _store_shift_snapshot(self, shift: dict, make_current_report: bool = False):
        with self._shift_db_write_lock():
            with self._shift_db_connect() as con:
                self._upsert_shift_con(con, shift)
                if make_current_report:
                    con.execute(
                        "INSERT OR REPLACE INTO meta(key,value) VALUES('current_report_shift_id',?)",
                        (_clean(shift.get("shift_id")),),
                    )

    def get_shift(self, shift_id: str) -> dict | None:
        shift_id = _clean(shift_id)
        if not shift_id:
            return None
        try:
            with self._shift_db_connect() as con:
                row = con.execute("SELECT data_json FROM shifts WHERE shift_id=?", (shift_id,)).fetchone()
            return json.loads(row["data_json"]) if row else None
        except Exception:
            return None

    def _current_report_shift_id(self) -> str:
        try:
            with self._shift_db_connect() as con:
                row = con.execute("SELECT value FROM meta WHERE key='current_report_shift_id'").fetchone()
            return _clean(row["value"]) if row else ""
        except Exception:
            return ""

    def _effective_start_timestamp(self, event_id: str, timestamp: str) -> str:
        path = self.corrections / f"{_safe_name(event_id)}.json"
        if not path.is_file():
            return timestamp
        data = json.loads(path.read_text(encoding="utf-8"))
        if _clean(data.get("start_event_id")) != _clean(event_id):
            raise BusyError("Ασυμφωνία στα στοιχεία διόρθωσης ώρας.")
        return _clean(data["corrected_start_at"])

    def _apply_start_correction(self, event: dict) -> dict:
        if _clean(event.get("kind")).upper() != "START":
            return event
        revised = self._effective_start_timestamp(_clean(event.get("event_id")), _clean(event.get("timestamp")))
        if revised == _clean(event.get("timestamp")):
            return event
        result = dict(event)
        result["original_timestamp"] = event["timestamp"]
        result["timestamp"] = revised
        return result

    def correct_active_start(self, agms: str, new_when: datetime, officer: PersonnelRecord, reason: str) -> dict:
        """Audited logical correction; immutable raw START event remains intact."""
        reason = _clean(reason)
        if len(reason) < 5:
            raise ValueError("Συμπληρώστε αιτιολογία τουλάχιστον 5 χαρακτήρων.")
        if new_when > datetime.now():
            raise ValueError("Δεν επιτρέπεται μελλοντική ώρα έναρξης.")
        with self.officer_lock(agms):
            active = self.get_active(agms)
            if not active:
                raise ValueError("Η θέση έχει ήδη κλείσει ή δεν υπάρχει.")
            old_when = datetime.fromisoformat(active["timestamp"])
            if new_when.date() != old_when.date():
                raise ValueError("Για ασφάλεια, η διόρθωση περιορίζεται στην ίδια ημερομηνία.")
            event_id = _clean(active.get("event_id"))
            if not event_id:
                raise ValueError("Δεν βρέθηκε αναγνωριστικό έναρξης.")
            correction = {
                "start_event_id": event_id, "agms": agms,
                "previous_start_at": active["timestamp"],
                "corrected_start_at": new_when.isoformat(timespec="seconds"),
                "corrected_by": self._person_payload(officer),
                "reason": reason,
                "recorded_at": datetime.now().isoformat(timespec="seconds"),
            }
            path = self.corrections / f"{_safe_name(event_id)}.json"
            if path.exists():
                previous = json.loads(path.read_text(encoding="utf-8"))
                correction["original_start_at"] = previous.get("original_start_at", previous.get("previous_start_at"))
            else:
                correction["original_start_at"] = active["timestamp"]
            self._atomic_json(path, correction)
            self.log_admin_action("ACTIVE_START_TIME_CORRECTED", correction)
            return correction

    def get_active(self, agms: str) -> dict | None:
        path = self._active_path(agms)
        if not self.active.is_dir():
            raise BusyError("Ο κοινός φάκελος ενεργών θέσεων δεν είναι διαθέσιμος.")
        try:
            item = json.loads(path.read_text(encoding="utf-8"))
            return None if self._has_end_receipt(item) else self._apply_start_correction(item)
        except FileNotFoundError:
            return None
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            # Fail closed: an unreadable shared active file is NOT a free booth.
            raise BusyError(f"Δεν είναι ασφαλής η ανάγνωση ενεργής θέσης {agms}: {exc}") from exc

    def _existing_operation_event(self, op_id: str, when: datetime) -> dict | None:
        if not op_id:
            return None
        # A retry after a server acknowledgement was lost must never record
        # another START/END. Shared lock for AGMS protects this lookup.
        folder = self.events / f"{when:%Y}" / f"{when:%m}" / f"{when:%Y-%m-%d}"
        prefix = f"{when:%Y%m%d_%H%M%S}_"
        for path in folder.glob(f"{prefix}*_{_safe_name(op_id)}.json"):
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                if data.get("event_id") == op_id:
                    return data
            except (OSError, ValueError):
                raise BusyError("Δεν διαβάζεται η προηγούμενη καταχώριση συγχρονισμού.")
        return None

    def _write_event(self, event: dict) -> Path:
        dt = datetime.fromisoformat(event["timestamp"])
        folder = self.events / f"{dt:%Y}" / f"{dt:%m}" / f"{dt:%Y-%m-%d}"
        filename = f"{dt:%Y%m%d_%H%M%S_%f}_{event['event_id']}.json"
        path = folder / filename
        self._atomic_json(path, event)
        return path

    def start(self, person: PersonnelRecord, position: str, direction: str, when: datetime | None = None, client_id: str = "", event_id: str = "", origin_computer: str = "", origin_user: str = "") -> dict:
        when = when or datetime.now()
        position = _clean(position)
        direction = _clean(direction).upper()
        if not position:
            raise ValueError("Επιλέξτε θέση.")
        if direction not in {"ΕΙΣΟΔΟΣ", "ΕΞΟΔΟΣ"}:
            raise ValueError("Επιλέξτε κατεύθυνση ΕΙΣΟΔΟΣ ή ΕΞΟΔΟΣ.")
        with self.officer_lock(person.agms):
            previous = self._existing_operation_event(event_id, when)
            if previous:
                if previous.get("kind") != "START" or previous.get("agms") != person.agms:
                    raise BusyError("Ασυμφωνία αναγνωριστικού START.")
                current = self.get_active(person.agms)
                if current and _clean(current.get("event_id")) == event_id:
                    return previous
                if current:
                    raise BusyError("Η παλιά έναρξη αντικαταστάθηκε. Δεν επιτρέπεται επανακατάληψη.")
                # If START was recorded but its active file was interrupted,
                # recover only if the original shift was never ended.
                if self._has_end_receipt(previous):
                    return previous
                with self._shift_db_connect() as db:
                    closed = db.execute(
                        "SELECT 1 FROM personnel_sessions WHERE start_event_id=? LIMIT 1", (event_id,)
                    ).fetchone()
                if closed:
                    return previous
                with self.officer_lock("__POSITION_START__"):
                    if any(_clean(x.get("position")) == position for x in self.active_all()):
                        raise BusyError("Η θέση έχει ήδη καταληφθεί στο μεταξύ.")
                    self._atomic_json(self._active_path(person.agms), previous)
                return previous
            # Serialize START operations across all stations so two PCs cannot
            # claim the same physical position at the same moment.
            with self.officer_lock("__POSITION_START__"):
                current = self.get_active(person.agms)
                if current and event_id and current.get("event_id") == event_id:
                    return current
                if current:
                    raise ValueError(
                        f"Υπάρχει ήδη ανοικτή θέση για {person.full_name}: "
                        f"{current.get('position','')} {current.get('direction','')} από {current.get('timestamp','')}. "
                        "Πρέπει πρώτα να γίνει ΛΗΞΗ."
                    )
                occupied = next(
                    (
                        item for item in self.active_all()
                        if _clean(item.get("position")) == position
                        and _clean(item.get("agms")) != _clean(person.agms)
                    ),
                    None,
                )
                if occupied:
                    occupant = _clean(
                        f"{occupied.get('rank','')} {occupied.get('surname','')} {occupied.get('given_name','')}"
                    ) or _clean(occupied.get("full_name")) or _clean(occupied.get("agms"))
                    raise ValueError(
                        f"Η ΘΕΣΗ {position} χρησιμοποιείται ήδη από {occupant}. "
                        "Πρέπει πρώτα να γίνει ΛΗΞΗ της ενεργής θέσης."
                    )
                event = {
                    "version": 2,
                    "event_id": event_id or uuid.uuid4().hex,
                    "kind": "START",
                    "timestamp": when.isoformat(timespec="seconds"),
                    "agms": person.agms,
                    "rank": person.rank,
                    "surname": person.surname,
                    "given_name": person.given_name,
                    "full_name": person.full_name,
                    "position": position,
                    "direction": direction,
                    "computer": origin_computer or socket.gethostname(),
                    "user": origin_user or os.environ.get("USERNAME") or os.environ.get("USER") or "",
                    "client_id": _clean(client_id),
                }
                path = self._write_event(event)
                event["event_file"] = str(path)
                try:
                    self._atomic_json(self._active_path(person.agms), event)
                except Exception:
                    # SMB can commit replace() and lose its acknowledgement.
                    # Never erase the durable START when a marker exists, or
                    # when its state cannot be read safely. The original ID
                    # lets a pending offline operation reconcile on retry.
                    try:
                        self._active_path(person.agms).read_text(encoding="utf-8")
                    except FileNotFoundError:
                        try:
                            if self.active.is_dir():
                                # Confirmed pre-commit failure: retain the old
                                # guarantee that no phantom START is reported.
                                path.unlink(missing_ok=True)
                        except OSError:
                            pass  # uncertain share state: preserve evidence
                    except OSError:
                        pass
                    raise
        try:
            self.sync_active_shift()
        except Exception:
            # START was durably recorded. Do not report its failure merely
            # because synchronizing the shift membership is delayed.
            pass
        return event

    def end(self, person: PersonnelRecord, when: datetime | None = None, expected_start_event_id: str = "", end_reason: str = "", end_officer: PersonnelRecord | None = None, event_id: str = "", origin_computer: str = "", origin_user: str = "") -> dict:
        when = when or datetime.now()
        with self.officer_lock(person.agms):
            previous = self._existing_operation_event(event_id, when)
            if previous:
                if previous.get("kind") != "END" or previous.get("agms") != person.agms:
                    raise BusyError("Ασυμφωνία αναγνωριστικού END.")
                current = self.get_active(person.agms)
                if current and _clean(current.get("event_id")) == _clean(previous.get("start_event_id")):
                    # Repair interruption between END event and active release.
                    receipt = self.pending_index / f"{_safe_name(current.get('event_id'))}.json"
                    self._atomic_json(receipt, {"start": current, "end": previous})
                    self._active_path(person.agms).unlink(missing_ok=True)
                    try:
                        self._store_personnel_session(current, previous)
                    except Exception:
                        pass  # durable receipt is reconciled on subsequent startup
                    else:
                        receipt.unlink(missing_ok=True)
                elif current:
                    raise BusyError("Η θέση άλλαξε μετά την προηγούμενη ΛΗΞΗ.")
                return previous
            current = self.get_active(person.agms)
            if not current:
                raise ValueError(f"Δεν υπάρχει ανοικτή θέση για {person.full_name}.")
            if expected_start_event_id and _clean(current.get("event_id")) != _clean(expected_start_event_id):
                raise BusyError("Η θέση άλλαξε χρήστη πριν επιβεβαιωθεί η λήξη. Δεν τροποποιήθηκε η νέα καταχώριση.")
            event = {
                "version": 2,
                "event_id": event_id or uuid.uuid4().hex,
                "kind": "END",
                "timestamp": when.isoformat(timespec="seconds"),
                "start_event_id": current.get("event_id", ""),
                "agms": person.agms,
                "rank": current.get("rank", person.rank),
                "surname": current.get("surname", person.surname),
                "given_name": current.get("given_name", person.given_name),
                "full_name": current.get("full_name", person.full_name),
                "position": current.get("position", ""),
                "direction": current.get("direction", ""),
                "computer": origin_computer or socket.gethostname(),
                "user": origin_user or os.environ.get("USERNAME") or os.environ.get("USER") or "",
            }
            if end_reason:
                event["end_reason"] = _clean(end_reason)
            if end_officer:
                event["end_officer"] = self._person_payload(end_officer)
            self._write_event(event)
            # Durable receipt before releasing a booth; history indexing can be
            # retried later if SQLite is locked or temporarily unavailable.
            receipt = self.pending_index / f"{_safe_name(current.get('event_id'))}.json"
            self._atomic_json(receipt, {"start": current, "end": event})
            self._active_path(person.agms).unlink(missing_ok=True)
            try:
                self._store_personnel_session(current, event)
            except Exception:
                event["session_index_pending"] = True
            else:
                receipt.unlink(missing_ok=True)
            return event

    def active_all(self) -> list[dict]:
        if not self.active.is_dir():
            raise BusyError("Δεν είναι διαθέσιμος ο κοινός φάκελος ενεργών θέσεων.")
        result: list[dict] = []
        for path in sorted(self.active.glob("*.json")):
            try:
                item = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, ValueError, json.JSONDecodeError) as exc:
                raise BusyError(f"Δεν διαβάζεται η δέσμευση {path.name}: {exc}") from exc
            if not self._has_end_receipt(item):
                result.append(self._apply_start_correction(item))
        return sorted(result, key=lambda x: (_clean(x.get("direction")), _clean(x.get("position")), _clean(x.get("full_name"))))

    def close_previous_windows_user_positions(
        self, computer: str, windows_user: str, *, protected_start_ids=()
    ) -> list[dict]:
        return close_previous_windows_user_positions(
            self, computer, windows_user, PersonnelRecord, _clean, BusyError,
            protected_start_ids=protected_start_ids,
        )

    def iter_events(self, start_dt: datetime, end_dt: datetime) -> list[dict]:
        result: list[dict] = []
        current = start_dt.date()
        last = end_dt.date()
        while current <= last:
            folder = self.events / f"{current:%Y}" / f"{current:%m}" / f"{current:%Y-%m-%d}"
            if folder.exists():
                for path in sorted(folder.glob("*.json")):
                    try:
                        e = self._apply_start_correction(json.loads(path.read_text(encoding="utf-8")))
                        ts = datetime.fromisoformat(e.get("timestamp", ""))
                        if start_dt <= ts < end_dt:
                            result.append(e)
                    except Exception:
                        pass
            current += timedelta(days=1)
        return sorted(result, key=lambda e: e.get("timestamp", ""))

    def iter_all_events(self) -> list[dict]:
        result: list[dict] = []
        if not self.events.exists():
            return result
        for path in sorted(self.events.rglob("*.json")):
            try:
                event = self._apply_start_correction(json.loads(path.read_text(encoding="utf-8")))
                if _clean(event.get("timestamp")):
                    result.append(event)
            except Exception:
                pass
        return sorted(result, key=lambda e: _clean(e.get("timestamp")))

    def completed_sessions_all(self) -> list[dict]:
        sessions: list[dict] = []
        for row in reversed(self.search_personnel_sessions()):
            try:
                start_dt = datetime.fromisoformat(row["start_at"])
                end_dt = datetime.fromisoformat(row["end_at"])
            except Exception:
                continue
            start = {
                "event_id": row["start_event_id"], "kind": "START", "timestamp": row["start_at"],
                "agms": row["agms"], "rank": row["rank"], "surname": row["surname"],
                "given_name": row["given_name"], "full_name": row["full_name"],
                "position": row["position"], "direction": row["direction"],
            }
            end = {
                "event_id": row["end_event_id"], "kind": "END", "timestamp": row["end_at"],
                "start_event_id": row["start_event_id"], "agms": row["agms"],
                "rank": row["rank"], "surname": row["surname"], "given_name": row["given_name"],
                "full_name": row["full_name"], "position": row["position"], "direction": row["direction"],
            }
            sessions.append({"start": start, "end": end, "start_dt": start_dt, "end_dt": end_dt})
        return sessions

    def save_shift(self, start_at: datetime, agms_list: list[str] | tuple[str, ...], lookback_minutes: int = 30, source_file: str = "") -> dict:
        unique = []
        seen = set()
        for value in agms_list:
            agms = _clean(value)
            if agms and agms not in seen:
                seen.add(agms); unique.append(agms)
        if not unique:
            raise ValueError("Η βάρδια δεν έχει προσωπικό.")
        lookback = max(0, min(int(lookback_minutes), 180))
        shift_id = f"SHIFT_{start_at:%Y%m%d_%H%M%S}"
        existing = self.get_shift(shift_id)
        if existing and _clean(existing.get("status")).upper() == "FINALIZED":
            raise ValueError("Η βάρδια έχει οριστικοποιηθεί και δεν μπορεί να αντικατασταθεί.")
        data = {
            "version": 1,
            "shift_id": shift_id,
            "start_at": start_at.isoformat(timespec="seconds"),
            "lookback_minutes": lookback,
            "max_window_hours": 12,
            "source_file": _clean(source_file),
            "agms": unique,
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "created_by": os.environ.get("USERNAME") or os.environ.get("USER") or "",
            "computer": socket.gethostname(),
            "status": "DRAFT",
        }
        self._store_shift_snapshot(data)
        self.log_admin_action("SHIFT_SAVED", {"shift_id": shift_id, "people": len(unique), "lookback_minutes": lookback})
        return data

    def _current_shift_path(self) -> Path:
        return self.shifts / "CURRENT_SHIFT.json"

    def current_shift(self) -> dict | None:
        path = self._current_shift_path()
        if not self.shifts.is_dir():
            raise BusyError("Ο κοινός φάκελος βαρδιών δεν είναι διαθέσιμος.")
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            return data if _clean(data.get("status")).upper() == "ACTIVE" else None
        except FileNotFoundError:
            return None
        except Exception as exc:
            # A read failure is not proof that no shift exists.
            raise BusyError(f"Δεν είναι ασφαλής η ανάγνωση τρέχουσας βάρδιας: {exc}") from exc

    def current_report_shift(self) -> dict | None:
        """Return the active/last shift; finalized history itself lives in SQLite."""
        path = self._current_shift_path()
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            if _clean(data.get("shift_id")) and _clean(data.get("status")).upper() in {"ACTIVE", "FINALIZED"}:
                return data
        except Exception:
            pass
        shift_id = self._current_report_shift_id()
        if shift_id:
            data = self.get_shift(shift_id)
            if data and _clean(data.get("status")).upper() == "FINALIZED":
                return data
        return None

    @staticmethod
    def _person_payload(person: PersonnelRecord) -> dict:
        return {
            "agms": person.agms, "rank": person.rank, "surname": person.surname,
            "given_name": person.given_name, "full_name": person.full_name,
        }

    def start_officer_shift(
        self,
        officer: PersonnelRecord,
        when: datetime | None = None,
        initial_approved: list[str] | tuple[str, ...] | None = None,
        initial_excluded: list[str] | tuple[str, ...] | None = None,
    ) -> dict:
        with self.officer_lock("__SHIFT_CONTROL__", timeout=20.0):
            if self.current_shift():
                raise ValueError("Υπάρχει ήδη ενεργή βάρδια. Πρέπει πρώτα να γίνει ΛΗΞΗ ΒΑΡΔΙΑΣ.")
            when = when or datetime.now()
            active_now = self.active_all()
            active_by_agms = {}
            for item in active_now:
                agms = _clean(item.get("agms"))
                if agms and agms not in active_by_agms:
                    active_by_agms[agms] = item

            if initial_approved is None and initial_excluded is None:
                approved = list(active_by_agms)
                excluded = []
            else:
                approved = []
                seen = set()
                for value in initial_approved or []:
                    agms = _clean(value)
                    if agms and agms in active_by_agms and agms not in seen:
                        approved.append(agms); seen.add(agms)
                excluded = []
                for value in initial_excluded or []:
                    agms = _clean(value)
                    if agms and agms in active_by_agms and agms not in seen:
                        excluded.append(agms); seen.add(agms)
                for agms in active_by_agms:
                    if agms not in seen:
                        excluded.append(agms)

            # Κρατάμε την πραγματική ώρα έναρξης ΟΛΩΝ όσων είναι ήδη ενεργοί
            # κατά την έναρξη της βάρδιας, ακόμη κι αν αρχικά εξαιρεθούν.
            # Έτσι μια μεταγενέστερη διόρθωση ΕΞΑΙΡΕΣΗ -> ΑΠΟΔΟΧΗ
            # επαναφέρει πλήρως την πραγματική αρχική ώρα εργασίας.
            initial_active = {}
            for agms, item in active_by_agms.items():
                initial_active[agms] = {
                    "timestamp": _clean(item.get("timestamp")),
                    "event_id": _clean(item.get("event_id")),
                }

            shift = {
                "version": 2,
                "shift_id": f"SHIFT_{when:%Y%m%d_%H%M%S_%f}",
                "status": "ACTIVE",
                "start_at": when.isoformat(timespec="seconds"),
                "start_officer": self._person_payload(officer),
                "approved": approved,
                "excluded": excluded,
                "pending": [],
                "initial_active": initial_active,
                "created_at": datetime.now().isoformat(timespec="seconds"),
            }
            self._atomic_json(self._current_shift_path(), shift)
            self.log_admin_action("SHIFT_STARTED", {
                "shift_id": shift["shift_id"],
                "officer_agms": officer.agms,
                "initial_people": len(approved),
                "initial_excluded": len(excluded),
            })
            return shift

    def _sync_active_shift_locked(self) -> dict | None:
        shift = self.current_shift()
        if not shift:
            return None
        approved = [_clean(x) for x in shift.get("approved", []) if _clean(x)]
        excluded = [_clean(x) for x in shift.get("excluded", []) if _clean(x)]
        pending = [_clean(x) for x in shift.get("pending", []) if _clean(x)]
        known = set(approved) | set(excluded) | set(pending)
        changed = False
        for item in self.active_all():
            agms = _clean(item.get("agms"))
            if agms and agms not in known:
                pending.append(agms); known.add(agms); changed = True
        if changed:
            shift["pending"] = pending
            shift["updated_at"] = datetime.now().isoformat(timespec="seconds")
            self._atomic_json(self._current_shift_path(), shift)
        return shift

    def sync_active_shift(self) -> dict | None:
        with self.officer_lock("__SHIFT_CONTROL__", timeout=20.0):
            return self._sync_active_shift_locked()

    def set_shift_member_state(self, agms: str, state: str) -> dict:
        with self.officer_lock("__SHIFT_CONTROL__", timeout=20.0):
            shift = self.current_shift()
            if not shift:
                raise ValueError("Δεν υπάρχει ενεργή βάρδια.")
            agms = _clean(agms)
            state = _clean(state).upper()
            if state not in {"APPROVED", "EXCLUDED"}:
                raise ValueError("Μη έγκυρη κατάσταση μέλους βάρδιας.")
            for key in ("approved", "excluded", "pending"):
                shift[key] = [x for x in shift.get(key, []) if _clean(x) != agms]
            shift["approved" if state == "APPROVED" else "excluded"].append(agms)
            shift["updated_at"] = datetime.now().isoformat(timespec="seconds")
            self._atomic_json(self._current_shift_path(), shift)
            self.log_admin_action("SHIFT_MEMBER_STATE", {"shift_id": shift["shift_id"], "agms": agms, "state": state})
            return shift

    def end_officer_shift_member(
        self, officer: PersonnelRecord, agms: str, expected_shift_id: str,
        expected_start_event_id: str,
    ) -> dict:
        """Individually end a confirmed, approved member without finalizing shift."""
        agms = _clean(agms)
        if not agms or not expected_start_event_id:
            raise BusyError("Η επιλεγμένη θέση δεν έχει έγκυρο αναγνωριστικό.")
        with self.officer_lock("__SHIFT_CONTROL__", timeout=20.0):
            shift = self.current_shift()
            if not shift or _clean(shift.get("shift_id")) != _clean(expected_shift_id):
                raise BusyError("Η βάρδια άλλαξε. Ελέγξτε ξανά την τρέχουσα λίστα.")
            if agms not in {_clean(x) for x in shift.get("approved", [])}:
                raise BusyError("Ο αστυνομικός δεν ανήκει πλέον στους αποδεκτούς της βάρδιας.")
            active = self.get_active(agms)
            if not active or _clean(active.get("event_id")) != _clean(expected_start_event_id):
                raise BusyError("Η ενεργή θέση άλλαξε. Δεν καταχωρίστηκε λανθασμένη λήξη.")
            member = PersonnelRecord(
                agms=agms, rank=_clean(active.get("rank")),
                surname=_clean(active.get("surname")), given_name=_clean(active.get("given_name")),
            )
            event = self.end(
                member, expected_start_event_id=expected_start_event_id,
                end_reason="SHIFT_SINGLE_BY_OFFICER", end_officer=officer,
            )
            try:
                self.log_admin_action("SHIFT_SINGLE_END_POSITION", {
                    "shift_id": expected_shift_id, "officer_agms": officer.agms,
                    "member_agms": agms, "position": _clean(active.get("position")),
                    "start_event_id": expected_start_event_id,
                    "end_event_id": _clean(event.get("event_id")),
                })
            except Exception:
                # Audit actor and reason are also embedded in durable END event.
                pass
            return event

    def end_officer_shift(
        self, officer: PersonnelRecord, when: datetime | None = None,
        expected_active_start_ids: dict[str, str] | None = None,
        expected_shift_id: str = "",
    ) -> dict:
        with self.officer_lock("__SHIFT_CONTROL__", timeout=20.0):
            shift = self._sync_active_shift_locked()
            if not shift:
                raise ValueError("Δεν υπάρχει ενεργή βάρδια.")
            if expected_shift_id and _clean(shift.get("shift_id")) != _clean(expected_shift_id):
                raise BusyError("Η βάρδια άλλαξε. Ανανεώστε τη λίστα πριν συνεχίσετε.")
            if shift.get("pending"):
                raise ValueError("Υπάρχουν άτομα σε ΑΝΑΜΟΝΗ. Κάντε πρώτα ΑΠΟΔΟΧΗ ή ΕΞΑΙΡΕΣΗ.")
            if expected_active_start_ids is not None:
                approved = set(map(_clean, shift.get("approved", [])))
                actual = {
                    _clean(item.get("agms")): _clean(item.get("event_id"))
                    for item in self.active_all()
                    if _clean(item.get("agms")) in approved
                }
                if actual != expected_active_start_ids:
                    raise BusyError(
                        "Οι ενεργές θέσεις άλλαξαν μετά την επιβεβαίωση. "
                        "Ελέγξτε ξανά τον κατάλογο πριν από τη λήξη βάρδιας."
                    )
            when = when or datetime.now()
            # A finalized roster must not leave its selected officers blocking the
            # physical desks. Close only APPROVED active records, through the same
            # audited END path used by the normal operator action. This operation
            # is deliberately retryable: already closed records are skipped.
            closed_agms = []
            for agms in sorted(set(_clean(a) for a in shift.get("approved", []))):
                if not agms:
                    continue
                current = self.get_active(agms)
                if not current:
                    continue
                started = datetime.fromisoformat(current["timestamp"])
                if started > when:
                    raise ValueError(
                        f"Η έναρξη του ΑΓΜΣ {agms} είναι μεταγενέστερη της λήξης βάρδιας."
                    )
                member = PersonnelRecord(
                    agms=agms,
                    rank=_clean(current.get("rank")),
                    surname=_clean(current.get("surname")),
                    given_name=_clean(current.get("given_name")),
                )
                expected_start = (
                    expected_active_start_ids.get(agms, "")
                    if expected_active_start_ids is not None
                    else _clean(current.get("event_id"))
                )
                if not expected_start or _clean(current.get("event_id")) != _clean(expected_start):
                    raise BusyError(
                        "Η ενεργή καταχώριση άλλαξε μετά την επιβεβαίωση. "
                        "Δεν πραγματοποιήθηκε λήξη νεότερης θέσης."
                    )
                self.end(
                    member, when=when, expected_start_event_id=expected_start,
                    end_reason="SHIFT_FINALIZED_BY_OFFICER", end_officer=officer,
                )
                self.log_admin_action("SHIFT_AUTO_END_POSITION", {
                    "shift_id": shift["shift_id"],
                    "officer_agms": officer.agms,
                    "member_agms": agms,
                    "end_at": when.isoformat(timespec="seconds"),
                    "reason": "FINALIZE_SELECTED_SHIFT_MEMBER",
                })
                closed_agms.append(agms)
            shift["closed_active_agms"] = closed_agms
            shift["status"] = "FINALIZED"
            shift["end_at"] = when.isoformat(timespec="seconds")
            shift["end_officer"] = self._person_payload(officer)
            shift["finalized_at"] = datetime.now().isoformat(timespec="seconds")
            # Το ιστορικό βαρδιών αποθηκεύεται σε SQLite. Το CURRENT_SHIFT.json
            # μένει μόνο ως ελαφρύ ζωντανό αρχείο συγχρονισμού μεταξύ σταθμών.
            # Persist the finalized report before replacing CURRENT_SHIFT.
            # If SQLite is unavailable, the shift remains ACTIVE and retryable.
            self._store_shift_snapshot(shift, make_current_report=True)
            self._atomic_json(self._current_shift_path(), shift)
            self.log_admin_action("SHIFT_ENDED", {"shift_id": shift["shift_id"], "officer_agms": officer.agms, "approved": len(shift.get("approved", [])), "excluded": len(shift.get("excluded", []))})
            return shift

    def shift_sessions_v2(self, shift: dict) -> list[dict]:
        approved = {_clean(x) for x in shift.get("approved", []) if _clean(x)}
        if not approved:
            return []
        start_at = datetime.fromisoformat(shift["start_at"])
        end_at = datetime.fromisoformat(shift.get("end_at") or datetime.now().isoformat(timespec="seconds"))
        initial = {a: dict(item) for a, item in (shift.get("initial_active", {}) or {}).items()}
        for info in initial.values():
            if _clean(info.get("event_id")):
                info["timestamp"] = self._effective_start_timestamp(
                    _clean(info["event_id"]), _clean(info.get("timestamp"))
                )
        search_start = start_at
        for agms in approved:
            try:
                ts = datetime.fromisoformat(_clean((initial.get(agms) or {}).get("timestamp")))
                if ts < search_start:
                    search_start = ts
            except Exception:
                pass
        events = self.iter_events(search_start, end_at + timedelta(seconds=1))
        starts: dict[str, dict] = {}
        sessions: list[dict] = []
        for e in events:
            agms = _clean(e.get("agms"))
            if agms not in approved:
                continue
            if e.get("kind") == "START":
                starts[_clean(e.get("event_id"))] = e
            elif e.get("kind") == "END":
                st = starts.get(_clean(e.get("start_event_id")))
                if not st:
                    continue
                sdt = datetime.fromisoformat(st["timestamp"]); edt = datetime.fromisoformat(e["timestamp"])
                threshold = start_at
                try:
                    threshold = datetime.fromisoformat(_clean((initial.get(agms) or {}).get("timestamp")))
                except Exception:
                    pass
                if sdt >= threshold and sdt < end_at:
                    sessions.append({"start": st, "end": e, "start_dt": sdt, "end_dt": min(edt, end_at)})
        for st in self.active_all():
            agms = _clean(st.get("agms"))
            if agms not in approved:
                continue
            try:
                sdt = datetime.fromisoformat(st["timestamp"])
            except Exception:
                continue
            threshold = start_at
            try:
                threshold = datetime.fromisoformat(_clean((initial.get(agms) or {}).get("timestamp")))
            except Exception:
                pass
            if threshold <= sdt < end_at:
                sessions.append({"start": st, "end": None, "start_dt": sdt, "end_dt": end_at if shift.get("end_at") else None})
        return sorted(sessions, key=lambda x: (_clean(x["start"].get("full_name")), x["start_dt"]))

    def list_shifts(self) -> list[dict]:
        result: list[dict] = []
        try:
            with self._shift_db_connect() as con:
                rows = con.execute("SELECT data_json FROM shifts ORDER BY start_at DESC").fetchall()
            for row in rows:
                try:
                    result.append(json.loads(row["data_json"]))
                except Exception:
                    continue
        except Exception:
            pass
        return result


    def search_finalized_shifts(
        self,
        date_from: date | None = None,
        date_to: date | None = None,
        officer_query: str = "",
        limit: int | None = None,
    ) -> list[dict]:
        clauses = ["status='FINALIZED'", "end_at IS NOT NULL", "end_at<>''"]
        params: list[object] = []
        if date_from is not None:
            clauses.append("start_at>=?")
            params.append(f"{date_from:%Y-%m-%d}T00:00:00")
        if date_to is not None:
            clauses.append("start_at<=?")
            params.append(f"{date_to:%Y-%m-%d}T23:59:59")

        sql = "SELECT data_json FROM shifts WHERE " + " AND ".join(clauses) + " ORDER BY start_at DESC"
        if limit is not None and int(limit) > 0:
            sql += " LIMIT ?"
            params.append(int(limit))

        result: list[dict] = []
        try:
            with self._shift_db_connect() as con:
                rows = con.execute(sql, params).fetchall()
            query = _clean(officer_query).upper()
            if query:
                query = "".join(ch for ch in unicodedata.normalize("NFD", query) if unicodedata.category(ch) != "Mn")
            for row in rows:
                try:
                    shift = json.loads(row["data_json"])
                except Exception:
                    continue
                if query:
                    officer = shift.get("start_officer") or shift.get("end_officer") or {}
                    haystack = _clean(
                        f"{officer.get('rank','')} {officer.get('surname','')} "
                        f"{officer.get('given_name','')} {officer.get('full_name','')} {officer.get('agms','')}"
                    ).upper()
                    haystack = "".join(ch for ch in unicodedata.normalize("NFD", haystack) if unicodedata.category(ch) != "Mn")
                    if query not in haystack:
                        continue
                result.append(shift)
        except Exception:
            return []
        return result

    def roster_sessions(self, start_at: datetime, agms_list: list[str] | tuple[str, ...], lookback_minutes: int = 30, horizon_hours: int = 12) -> list[dict]:
        allowed = {_clean(x) for x in agms_list if _clean(x)}
        search_start = start_at - timedelta(minutes=max(0, int(lookback_minutes)))
        search_end = start_at + timedelta(hours=max(1, int(horizon_hours)))
        events = self.iter_events(search_start, search_end)
        starts: dict[str, dict] = {}
        sessions: list[dict] = []
        for e in events:
            if _clean(e.get("agms")) not in allowed:
                continue
            if e.get("kind") == "START":
                starts[e.get("event_id", "")] = e
            elif e.get("kind") == "END":
                sid = e.get("start_event_id", "")
                st = starts.get(sid)
                if not st:
                    continue
                sdt = datetime.fromisoformat(st["timestamp"]); edt = datetime.fromisoformat(e["timestamp"])
                sessions.append({"start": st, "end": e, "start_dt": sdt, "end_dt": edt})
        for st in self.active_all():
            if _clean(st.get("agms")) not in allowed:
                continue
            try:
                sdt = datetime.fromisoformat(st["timestamp"])
            except Exception:
                continue
            if search_start <= sdt < search_end:
                sessions.append({"start": st, "end": None, "start_dt": sdt, "end_dt": None})
        return sorted(sessions, key=lambda x: (_clean(x["start"].get("full_name")), x["start_dt"]))

    def log_admin_action(self, action: str, details: dict | None = None) -> Path:
        record = {
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "action": _clean(action),
            "details": details or {},
            "user": os.environ.get("USERNAME") or os.environ.get("USER") or "",
            "computer": socket.gethostname(),
        }
        path = self.audit / "admin_audit.jsonl"
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8") as fh:
            fh.write(json.dumps(record, ensure_ascii=False) + "\n")
        return path

    def backup_data(self) -> Path:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        target = self.backups / f"TDE_DATA_BACKUP_{stamp}.zip"
        temp = target.with_suffix(target.suffix + ".tmp")
        with zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for path in self.root.rglob("*"):
                if not path.is_file():
                    continue
                try:
                    rel = path.relative_to(self.root)
                except Exception:
                    continue
                if rel.parts and rel.parts[0].lower() == "backups":
                    continue
                zf.write(path, arcname=str(rel))
        os.replace(temp, target)
        self.log_admin_action("BACKUP_CREATED", {"file": str(target)})
        return target

    def finalize_shift(self, shift: dict, report_text: str = "") -> dict:
        shift_id = _clean(shift.get("shift_id"))
        if not shift_id:
            raise ValueError("Μη έγκυρη βάρδια.")
        current = self.get_shift(shift_id)
        if not current:
            raise ValueError("Δεν βρέθηκε η βάρδια για οριστικοποίηση.")
        if _clean(current.get("status")).upper() == "FINALIZED":
            return current
        current["status"] = "FINALIZED"
        current["finalized_at"] = datetime.now().isoformat(timespec="seconds")
        current["finalized_by"] = os.environ.get("USERNAME") or os.environ.get("USER") or ""
        current["finalized_computer"] = socket.gethostname()
        if report_text:
            current["final_report_text"] = report_text
            current["final_report_sha256"] = hashlib.sha256(report_text.encode("utf-8")).hexdigest()
        self._store_shift_snapshot(current)
        self.log_admin_action("SHIFT_FINALIZED", {"shift_id": shift_id, "people": len(current.get("agms", []))})
        return current

    def shift_sessions(self, day: date, cutoff: str = DEFAULT_CUTOFF) -> tuple[datetime, datetime, list[dict]]:
        hh, mm = [int(x) for x in cutoff.split(":", 1)]
        start_dt = datetime.combine(day, dtime(hh, mm))
        end_dt = start_dt + timedelta(days=1)
        # Read a wider window so sessions that started shortly before the cutoff can still be paired if needed.
        events = self.iter_events(start_dt - timedelta(days=1), end_dt)
        starts: dict[str, dict] = {}
        sessions: list[dict] = []
        for e in events:
            if e.get("kind") == "START":
                starts[e.get("event_id", "")] = e
            elif e.get("kind") == "END":
                sid = e.get("start_event_id", "")
                s = starts.get(sid)
                if not s:
                    continue
                sdt = datetime.fromisoformat(s["timestamp"])
                edt = datetime.fromisoformat(e["timestamp"])
                if edt <= start_dt or sdt >= end_dt:
                    continue
                sessions.append({"start": s, "end": e, "start_dt": sdt, "end_dt": edt})
        # Open sessions that fall in the selected shift.
        for s in self.active_all():
            try:
                sdt = datetime.fromisoformat(s["timestamp"])
            except Exception:
                continue
            if sdt < end_dt:
                sessions.append({"start": s, "end": None, "start_dt": sdt, "end_dt": None})
        sessions.sort(key=lambda x: (x["start"].get("full_name", ""), x["start_dt"]))
        return start_dt, end_dt, sessions








def create_personnel_statistics_xlsx(store: SharedEventStore, people_by_agms: dict[str, PersonnelRecord], target: str | Path | None = None) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import PieChart, Reference
        from openpyxl.chart.label import DataLabelList
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    except Exception as exc:
        raise RuntimeError("Δεν είναι διαθέσιμη η δημιουργία Excel στατιστικών (openpyxl).") from exc

    target_path = Path(target) if target else (store.statistics / "ΣΤΑΤΙΣΤΙΚΑ_ΠΡΟΣΩΠΙΚΟΥ.xlsx")
    target_path.parent.mkdir(parents=True, exist_ok=True)

    # Συγκέντρωση μόνο ολοκληρωμένων START/END συνεδριών.
    stats: dict[str, dict[tuple[str, str], int]] = {}
    for session in store.completed_sessions_all():
        start = session["start"]
        agms = _clean(start.get("agms"))
        if not agms:
            continue
        pos = _clean(start.get("position")) or "—"
        direction = _clean(start.get("direction")).upper() or "—"
        minutes = max(0, int((session["end_dt"] - session["start_dt"]).total_seconds() // 60))
        stats.setdefault(agms, {})[(pos, direction)] = stats.setdefault(agms, {}).get((pos, direction), 0) + minutes

    wb = Workbook()
    ws = wb.active
    ws.title = "ΣΤΑΤΙΣΤΙΚΑ"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "ΣΤΑΤΙΣΤΙΚΑ ΘΕΣΕΩΝ ΠΡΟΣΩΠΙΚΟΥ"
    c.font = Font(bold=True, color="FFFFFF", size=13)
    c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 19
    ws.column_dimensions["D"].width = 3
    for col in ("E", "F", "G", "H"):
        ws.column_dimensions[col].width = 12

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    name_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="EDF4FA")

    people = sorted(people_by_agms.values(), key=lambda p: (p.surname, p.given_name, p.agms))
    row = 4
    for person in people:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        name_cell = ws.cell(row=row, column=1, value=_clean(f"{person.rank} {person.full_name}"))
        name_cell.font = Font(bold=True, size=11)
        name_cell.fill = name_fill
        name_cell.alignment = Alignment(vertical="center")

        hdr = row + 1
        headers = ["ΘΕΣΗ", "ΚΑΤΕΥΘΥΝΣΗ", "ΣΥΝΟΛΙΚΟΣ ΧΡΟΝΟΣ"]
        for col_idx, text in enumerate(headers, 1):
            cell = ws.cell(row=hdr, column=col_idx, value=text)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        entries = sorted((stats.get(person.agms) or {}).items(), key=lambda x: (x[0][0], x[0][1]))
        if not entries:
            entries = [(('—', '—'), 0)]

        data_start = hdr + 1
        pie_by_position: dict[str, int] = {}
        for (pos, direction), minutes in entries:
            ws.cell(row=row+2, column=1, value=pos).font = Font(bold=True)
            ws.cell(row=row+2, column=2, value=direction)
            duration_cell = ws.cell(row=row+2, column=3, value=(minutes / 1440 if minutes else 0))
            duration_cell.number_format = '[h]" ώ "mm" λ"'
            for col_idx in range(1, 4):
                cell = ws.cell(row=row+2, column=col_idx)
                cell.alignment = Alignment(horizontal="center")
                cell.border = border
            pie_by_position[pos] = pie_by_position.get(pos, 0) + minutes
            row += 1

        data_end = row + 1
        # Κρυφά βοηθητικά δεδομένα μόνο για την πίτα.
        helper_start = hdr
        ws.cell(row=helper_start, column=11, value="ΘΕΣΗ")
        ws.cell(row=helper_start, column=12, value="ΛΕΠΤΑ")
        helper_row = helper_start + 1
        for pos, minutes in sorted(pie_by_position.items()):
            ws.cell(row=helper_row, column=11, value=pos)
            ws.cell(row=helper_row, column=12, value=minutes)
            helper_row += 1
        ws.column_dimensions["K"].hidden = True
        ws.column_dimensions["L"].hidden = True

        if any(v > 0 for v in pie_by_position.values()):
            chart = PieChart()
            labels = Reference(ws, min_col=11, min_row=helper_start+1, max_row=helper_row-1)
            data = Reference(ws, min_col=12, min_row=helper_start, max_row=helper_row-1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(labels)
            chart.title = ""
            chart.height = 4.3
            chart.width = 7.0
            chart.legend.position = "b"
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showPercent = True
            ws.add_chart(chart, f"E{hdr-1}")

        row = max(row + 4, hdr + 8)

    temp = target_path.with_suffix(target_path.suffix + ".tmp")
    wb.save(temp)
    os.replace(temp, target_path)
    return target_path


def create_shift_docx(store: SharedEventStore, shift: dict, people_by_agms: dict[str, PersonnelRecord] | None = None, target: str | Path | None = None) -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except Exception as exc:
        raise RuntimeError("Δεν είναι διαθέσιμη η δημιουργία Word (python-docx).") from exc
    people_by_agms = people_by_agms or {}
    sessions = store.shift_sessions_v2(shift) if int(shift.get("version", 1)) >= 2 else []
    roster = [_clean(x) for x in shift.get("approved", []) if _clean(x)]
    by_agms: dict[str, list[dict]] = {a: [] for a in roster}
    for item in sessions:
        by_agms.setdefault(_clean(item["start"].get("agms")), []).append(item)
    start_at = datetime.fromisoformat(shift["start_at"])
    end_at = datetime.fromisoformat(shift.get("end_at")) if shift.get("end_at") else None
    if target is None:
        target = store.current_docs / "ΤΡΕΧΟΥΣΑ_ΒΑΡΔΙΑ.docx"
    target = Path(target); target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = "Arial"; normal.font.size = Pt(11)
    for text, size, bold in [("ΕΛΛΗΝΙΚΗ ΑΣΤΥΝΟΜΙΑ", 14, True), ("Τ.Δ.Ε. ΚΡΥΣΤΑΛΛΟΠΗΓΗΣ", 13, True), ("ΚΑΤΑΓΡΑΦΗ ΘΕΣΕΩΝ ΒΑΡΔΙΑΣ", 13, True)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{start_at:%d/%m/%Y %H:%M}" + (f" - {end_at:%d/%m/%Y %H:%M}" if end_at else " - ΣΕ ΕΞΕΛΙΞΗ"))
    doc.add_paragraph()
    for agms in roster:
        items = sorted(by_agms.get(agms, []), key=lambda x: x["start_dt"])
        person = people_by_agms.get(agms)
        if person:
            name = _clean(f"{person.rank} {person.surname} {person.given_name}")
        elif items:
            st = items[0]["start"]; name = _clean(f"{st.get('rank','')} {st.get('full_name','')}") or f"ΑΓΜΣ {agms}"
        else:
            name = f"ΑΓΜΣ {agms}"
        p = doc.add_paragraph(); p.add_run(name).bold = True
        if not items:
            doc.add_paragraph("Χωρίς καταχώριση θέσης", style=None)
        for item in items:
            st = item["start"]; sdt = item["start_dt"]; edt = item["end_dt"]
            end_text = edt.strftime("%H:%M") if edt else "ΑΝΟΙΚΤΗ"
            p = doc.add_paragraph()
            r = p.add_run(f"ΘΕΣΗ {st.get('position','')}")
            r.bold = True
            direction = _clean(st.get("direction")).upper()
            if direction:
                p.add_run(f"   {direction}")
            p.add_run(f"   {sdt:%H:%M} - {end_text}")
    doc.add_paragraph()
    officer = shift.get("start_officer") or {}
    officer_name = _clean(f"{officer.get('rank','')} {officer.get('full_name','')}")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Κρυσταλλοπηγή, την {end_at:%d/%m/%Y}" if end_at else "Κρυσταλλοπηγή, την __________")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Ο ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ").bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(officer_name).bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Υπογραφή: __________________________")
    doc.save(target)
    return target


def create_shift_pdf(store: SharedEventStore, shift: dict, people_by_agms: dict[str, PersonnelRecord] | None = None, target: str | Path | None = None) -> Path:
    """Create the shift report directly as an A4 PDF, without Word."""
    try:
        from PySide6.QtCore import QMarginsF
        from PySide6.QtGui import QPageLayout, QPageSize, QTextDocument
        from PySide6.QtPrintSupport import QPrinter
    except Exception as exc:
        raise RuntimeError("Δεν είναι διαθέσιμη η δημιουργία PDF.") from exc

    people_by_agms = people_by_agms or {}
    sessions = store.shift_sessions_v2(shift) if int(shift.get("version", 1)) >= 2 else store.roster_sessions(
        datetime.fromisoformat(shift["start_at"]),
        shift.get("agms", []),
        int(shift.get("lookback_minutes", 30)),
        int(shift.get("max_window_hours", 12)),
    )
    roster = [_clean(x) for x in (shift.get("approved", []) if int(shift.get("version", 1)) >= 2 else shift.get("agms", [])) if _clean(x)]
    by_agms: dict[str, list[dict]] = {a: [] for a in roster}
    for item in sessions:
        by_agms.setdefault(_clean(item["start"].get("agms")), []).append(item)

    start_at = datetime.fromisoformat(shift["start_at"])
    end_at = datetime.fromisoformat(shift.get("end_at")) if shift.get("end_at") else None
    if target is None:
        preview_dir = Path(tempfile.gettempdir()) / "TDE_Krystallopigis" / "pdf_preview"
        preview_dir.mkdir(parents=True, exist_ok=True)
        target = preview_dir / "ΑΝΑΦΟΡΑ_ΒΑΡΔΙΑΣ.pdf"
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)

    body_parts = [
        "<div class='head'>ΕΛΛΗΝΙΚΗ ΑΣΤΥΝΟΜΙΑ</div>",
        "<div class='subhead'>Τ.Δ.Ε. ΚΡΥΣΤΑΛΛΟΠΗΓΗΣ</div>",
        "<div class='subhead'>ΚΑΤΑΓΡΑΦΗ ΘΕΣΕΩΝ ΒΑΡΔΙΑΣ</div>",
        "<div class='period'>" + _escape_html(
            f"{start_at:%d/%m/%Y %H:%M}" + (f" - {end_at:%d/%m/%Y %H:%M}" if end_at else " - ΣΕ ΕΞΕΛΙΞΗ")
        ) + "</div>",
    ]

    if not roster:
        body_parts.append("<p class='empty'>Δεν υπάρχουν καταχωρίσεις στη βάρδια.</p>")

    for agms in roster:
        items = sorted(by_agms.get(agms, []), key=lambda x: x["start_dt"])
        person = people_by_agms.get(agms)
        if person:
            name = _clean(f"{person.rank} {person.surname} {person.given_name}")
        elif items:
            st = items[0]["start"]
            name = _clean(f"{st.get('rank','')} {st.get('full_name','')}") or f"ΑΓΜΣ {agms}"
        else:
            name = f"ΑΓΜΣ {agms}"

        body_parts.append(f"<div class='person'>{_escape_html(name)}</div>")
        if not items:
            body_parts.append("<div class='row muted'>Χωρίς καταχώριση θέσης</div>")
        for item in items:
            st = item["start"]
            sdt = item["start_dt"]
            edt = item["end_dt"]
            end_text = edt.strftime("%H:%M") if edt else "ΑΝΟΙΚΤΗ"
            direction = _clean(st.get("direction")).upper()
            details = f"ΘΕΣΗ {st.get('position','')}"
            if direction:
                details += f"   {direction}"
            details += f"   {sdt:%H:%M} - {end_text}"
            body_parts.append(f"<div class='row'>{_escape_html(details)}</div>")

    officer = shift.get("start_officer") or shift.get("end_officer") or {}
    officer_name = _clean(f"{officer.get('rank','')} {officer.get('full_name','')}")
    place_date = f"Κρυσταλλοπηγή, την {end_at:%d/%m/%Y}" if end_at else "Κρυσταλλοπηγή, την __________"
    body_parts.extend([
        f"<div class='place'>{_escape_html(place_date)}</div>",
        "<div class='signature'><b>Ο ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ</b><br><br>"
        + _escape_html(officer_name)
        + "<br><br>Υπογραφή: __________________________</div>",
    ])

    html = """<!doctype html><html><head><meta charset='utf-8'><style>
        body { font-family: Arial, 'Segoe UI', sans-serif; font-size: 10.5pt; color: #000; }
        .head { text-align:center; font-size:14pt; font-weight:700; margin-bottom:4px; }
        .subhead { text-align:center; font-size:12.5pt; font-weight:700; margin-bottom:3px; }
        .period { text-align:center; margin:10px 0 18px 0; }
        .person { font-weight:700; margin-top:11px; margin-bottom:3px; }
        .row { margin-left:14px; margin-bottom:2px; }
        .muted { font-style:italic; }
        .empty { margin-top:20px; }
        .place { text-align:center; margin-top:28px; }
        .signature { text-align:center; margin-top:12px; }
    </style></head><body>""" + "".join(body_parts) + "</body></html>"

    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(str(target))
    printer.setPageSize(QPageSize(QPageSize.A4))
    printer.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)

    document = QTextDocument()
    document.setHtml(html)
    document.print_(printer)

    if not target.is_file() or target.stat().st_size == 0:
        raise RuntimeError("Δεν δημιουργήθηκε το PDF της αναφοράς.")
    return target




def summarize_personnel_history(rows: list[dict]) -> dict:
    total = 0
    by_direction: dict[str, int] = {}
    by_position: dict[str, int] = {}
    by_direction_position: dict[tuple[str, str], int] = {}
    people: set[str] = set()
    for row in rows:
        sec = max(0, int(row.get("duration_seconds") or 0))
        total += sec
        direction = _clean(row.get("direction")).upper() or "—"
        position = _clean(row.get("position")) or "—"
        agms = _clean(row.get("agms"))
        if agms:
            people.add(agms)
        by_direction[direction] = by_direction.get(direction, 0) + sec
        by_position[position] = by_position.get(position, 0) + sec
        key = (direction, position)
        by_direction_position[key] = by_direction_position.get(key, 0) + sec
    return {
        "total_seconds": total,
        "people": len(people),
        "by_direction": by_direction,
        "by_position": by_position,
        "by_direction_position": by_direction_position,
    }


def create_personnel_history_pdf(
    rows: list[dict],
    criteria: str,
    summary: dict | None = None,
    target: str | Path | None = None,
) -> Path:
    """Render the current personnel-history search as an on-demand PDF."""
    try:
        from PySide6.QtCore import QMarginsF
        from PySide6.QtGui import QPageLayout, QPageSize, QTextDocument
        from PySide6.QtPrintSupport import QPrinter
    except Exception as exc:
        raise RuntimeError("Δεν είναι διαθέσιμη η δημιουργία PDF.") from exc

    if not rows:
        raise ValueError("Δεν υπάρχουν αποτελέσματα για εξαγωγή σε PDF.")

    summary = summary or summarize_personnel_history(rows)
    if target is None:
        preview_dir = Path(tempfile.gettempdir()) / "TDE_Krystallopigis" / "pdf_preview"
        preview_dir.mkdir(parents=True, exist_ok=True)
        target = preview_dir / "ΙΣΤΟΡΙΚΟ_ΠΡΟΣΩΠΙΚΟΥ.pdf"
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)

    direction_lines = []
    for direction in ("ΕΙΣΟΔΟΣ", "ΕΞΟΔΟΣ"):
        seconds = int(summary.get("by_direction", {}).get(direction, 0))
        if not seconds:
            continue
        positions = []
        for (item_direction, position), sec in sorted(
            summary.get("by_direction_position", {}).items(),
            key=lambda x: (x[0][0], x[0][1]),
        ):
            if item_direction == direction and sec:
                positions.append(f"Θέση {position}: {_format_seconds(sec)}")
        detail = " · ".join(positions)
        direction_lines.append(
            f"<div><b>{_escape_html(direction)}:</b> {_format_seconds(seconds)}"
            + (f" &nbsp; ({_escape_html(detail)})" if detail else "")
            + "</div>"
        )

    body_rows = []
    for row in rows:
        try:
            sdt = datetime.fromisoformat(_clean(row.get("start_at")))
            edt = datetime.fromisoformat(_clean(row.get("end_at")))
            date_text = sdt.strftime("%d/%m/%Y")
            start_text = sdt.strftime("%H:%M:%S")
            end_text = edt.strftime("%H:%M:%S")
        except Exception:
            date_text = _clean(row.get("start_at"))[:10]
            start_text = _clean(row.get("start_at"))[11:19]
            end_text = _clean(row.get("end_at"))[11:19]
        name = _clean(f"{row.get('rank','')} {row.get('surname','')} {row.get('given_name','')}") or _clean(row.get("full_name"))
        cells = [
            date_text,
            _clean(row.get("agms")),
            name,
            _clean(row.get("position")),
            _clean(row.get("direction")),
            start_text,
            end_text,
            _format_seconds(int(row.get("duration_seconds") or 0)),
        ]
        body_rows.append("<tr>" + "".join(f"<td>{_escape_html(str(v))}</td>" for v in cells) + "</tr>")

    html = """<!doctype html><html><head><meta charset='utf-8'><style>
        body { font-family: Arial, 'Segoe UI', sans-serif; font-size: 9pt; color:#000; }
        h1 { text-align:center; font-size:14pt; margin:0 0 3px 0; }
        h2 { text-align:center; font-size:12pt; margin:0 0 12px 0; }
        .criteria { margin:8px 0 12px 0; padding:7px; border:1px solid #999; }
        .summary { margin:8px 0 12px 0; line-height:1.45; }
        table { width:100%; border-collapse:collapse; }
        th, td { border:1px solid #777; padding:4px; }
        th { font-weight:bold; background:#eee; }
        .right { text-align:right; }
    </style></head><body>
    <h1>ΕΛΛΗΝΙΚΗ ΑΣΤΥΝΟΜΙΑ</h1>
    <h2>Τ.Δ.Ε. ΚΡΥΣΤΑΛΛΟΠΗΓΗΣ — ΙΣΤΟΡΙΚΟ ΠΡΟΣΩΠΙΚΟΥ</h2>
    <div class='criteria'><b>Κριτήρια:</b> """ + _escape_html(criteria or "Χωρίς πρόσθετα φίλτρα") + """</div>
    <div class='summary'>
      <div><b>Καταχωρίσεις:</b> """ + str(len(rows)) + """ &nbsp; <b>Άτομα:</b> """ + str(summary.get("people", 0)) + """</div>
      <div><b>Συνολικός χρόνος:</b> """ + _format_seconds(int(summary.get("total_seconds", 0))) + """</div>
      """ + "".join(direction_lines) + """
    </div>
    <table>
      <tr><th>ΗΜΕΡΟΜΗΝΙΑ</th><th>ΑΓΜΣ</th><th>ΟΝΟΜΑ</th><th>ΘΕΣΗ</th><th>ΚΑΤΕΥΘΥΝΣΗ</th><th>ΕΝΑΡΞΗ</th><th>ΛΗΞΗ</th><th>ΔΙΑΡΚΕΙΑ</th></tr>
      """ + "".join(body_rows) + """
    </table></body></html>"""

    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(str(target))
    printer.setPageSize(QPageSize(QPageSize.A4))
    printer.setPageOrientation(QPageLayout.Landscape)
    printer.setPageMargins(QMarginsF(10, 10, 10, 10), QPageLayout.Millimeter)
    document = QTextDocument()
    document.setHtml(html)
    document.print_(printer)
    if not target.is_file() or target.stat().st_size == 0:
        raise RuntimeError("Δεν δημιουργήθηκε το PDF του ιστορικού προσωπικού.")
    return target


def archive_shift_pdf(store: SharedEventStore, shift: dict, people_by_agms: dict[str, PersonnelRecord] | None = None) -> Path:
    """Compatibility alias: reports are rendered on demand, not stored in the archive."""
    return create_shift_pdf(store, shift, people_by_agms)






def archive_shift_docx(store: SharedEventStore, shift: dict, people_by_agms: dict[str, PersonnelRecord] | None = None) -> Path:
    stamp = datetime.fromisoformat(_clean(shift.get("end_at")) or _clean(shift.get("start_at")))
    folder = store.archive / f"{stamp:%Y}" / f"{stamp:%m}" / f"{stamp:%Y-%m-%d}"
    folder.mkdir(parents=True, exist_ok=True)
    filename = re.sub(r'[<>:"/\\|?*]+', "_", shift_report_title(shift)).strip() + ".docx"
    target = folder / filename
    return create_shift_docx(store, shift, people_by_agms, target)








def run(args=None) -> int:
    from position_tracking_theme import THEME_DARK, THEME_LIGHT, _theme_stylesheet
    from position_tracking_rendering import coherent_widget_update
    from PySide6.QtCore import Qt, QDate, QDateTime, QTime, QTimer, QRectF, QPointF, QByteArray
    from PySide6.QtGui import QColor, QFont, QFontMetrics, QPainter, QPen, QIcon, QPixmap, QPalette, QPainterPath
    from PySide6.QtSvg import QSvgRenderer
    from PySide6.QtWidgets import (
        QApplication, QComboBox, QDateEdit, QFormLayout, QFrame, QHBoxLayout,
        QLabel, QLineEdit, QMainWindow, QMessageBox, QPushButton, QTextBrowser, QVBoxLayout, QWidget,
        QInputDialog, QDialog, QDateTimeEdit, QSpinBox, QListWidget, QListWidgetItem, QFileDialog, QToolButton,
        QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView, QCheckBox, QTimeEdit, QPlainTextEdit
    )


    def _police_emblem_pixmap(size: int = 82) -> QPixmap:
        """Render the exact approved Hellenic Police emblem shared by TDE-Control-Center."""
        pixmap = QPixmap()
        try:
            pixmap.loadFromData(base64.b64decode(POLICE_EMBLEM_PNG_B64), "PNG")
        except Exception:
            return QPixmap()
        if pixmap.isNull():
            return QPixmap()
        return pixmap.scaled(size, size, Qt.KeepAspectRatio, Qt.SmoothTransformation)


    class GreeceMapWidget(QWidget):
        """Render the exact approved shared Greece map without redrawing or recolouring it."""
        def __init__(self, parent=None):
            super().__init__(parent)
            self.setFixedSize(86, 86)
            self.setToolTip("ΕΛΛΑΔΑ")
            try:
                svg_bytes = base64.b64decode(GREECE_MAP_SVG_B64)
                self._renderer = QSvgRenderer(QByteArray(svg_bytes), self)
            except Exception:
                self._renderer = QSvgRenderer(self)

        def set_dark(self, dark: bool):
            # Compatibility with the existing theme code. The approved SVG stays unchanged.
            self.update()

        def paintEvent(self, event):
            painter = QPainter(self)
            painter.setRenderHint(QPainter.Antialiasing)
            if self._renderer.isValid():
                bounds = QRectF(self.rect().adjusted(4, 4, -4, -4))
                natural = self._renderer.defaultSize()
                if natural.isValid() and natural.width() > 0 and natural.height() > 0:
                    source_ratio = natural.width() / natural.height()
                    target_ratio = bounds.width() / bounds.height() if bounds.height() else source_ratio
                    if target_ratio > source_ratio:
                        width = bounds.height() * source_ratio
                        bounds.setX(bounds.x() + (bounds.width() - width) / 2.0)
                        bounds.setWidth(width)
                    else:
                        height = bounds.width() / source_ratio
                        bounds.setY(bounds.y() + (bounds.height() - height) / 2.0)
                        bounds.setHeight(height)
                self._renderer.render(painter, bounds)
            painter.end()


    class ArrowComboBox(QComboBox):
        """QComboBox with a permanently visible dropdown arrow."""
        def paintEvent(self, event):
            super().paintEvent(event)
            painter = QPainter(self)
            painter.setPen(self.palette().color(QPalette.Text))
            arrow_rect = self.rect().adjusted(self.width() - 28, 0, -2, 0)
            painter.drawText(arrow_rect, Qt.AlignCenter, "▼")
            painter.end()


    def _network_icon() -> QIcon:
        pixmap = QPixmap(256, 256)
        pixmap.fill(Qt.transparent)
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.Antialiasing)
        painter.setPen(Qt.NoPen)
        painter.setBrush(QColor("#0a4f8c"))
        painter.drawRoundedRect(18, 18, 220, 220, 48, 48)
        nodes = [(128, 128, 28), (66, 72, 22), (190, 72, 22), (66, 184, 22), (190, 184, 22)]
        painter.setPen(QPen(QColor("#8fd0ff"), 12, Qt.SolidLine, Qt.RoundCap))
        for x, y, _ in nodes[1:]:
            painter.drawLine(128, 128, x, y)
        for x, y, radius in nodes:
            painter.setPen(QPen(QColor("#dff3ff"), 5))
            painter.setBrush(QColor("#1976c9") if (x, y) != (128, 128) else QColor("#f5f9ff"))
            painter.drawEllipse(x - radius, y - radius, radius * 2, radius * 2)
        painter.end()
        return QIcon(pixmap)

    class BoothMapWidget(QWidget):
        BOOTH_POS = {
            "7": (0.46, 0.10), "6": (0.52, 0.10),
            "5": (0.52, 0.24), "4": (0.52, 0.32),
            "3": (0.40, 0.40), "2": (0.46, 0.40), "1": (0.52, 0.40),
            "10": (0.40, 0.64), "9": (0.46, 0.64), "8": (0.52, 0.64),
            "11": (0.52, 0.72), "12": (0.52, 0.80),
            "14": (0.46, 0.90), "13": (0.52, 0.90),
        }

        def __init__(self, parent=None):
            super().__init__(parent)
            self.active: dict[str, list[dict]] = {}
            self._worker_text_width = 84
            self.setMinimumSize(760, 560)

        @staticmethod
        def _worker_display_name(item: dict) -> str:
            surname = _clean(item.get("surname"))
            given_name = _clean(item.get("given_name"))
            full_name = _clean(item.get("full_name"))
            return _clean(f"{surname} {given_name}") or full_name or "—"

        def set_active(self, items: list[dict]):
            grouped: dict[str, list[dict]] = {}
            for item in items:
                grouped.setdefault(_clean(item.get("position")), []).append(item)
            self.active = grouped
            font = QFont("Segoe UI", 8, QFont.Bold)
            metrics = QFontMetrics(font)
            self._worker_text_width = max(
                [84] + [metrics.horizontalAdvance(self._worker_display_name(item)) for item in items]
            )
            self.update()

        def recommended_dialog_width(self) -> int:
            # The closest booths are 6% of the map width apart. Size the map so
            # the longest worker name fits without requiring manual stretching.
            booth_target = min(max(self._worker_text_width + 20, 92), 190)
            return max(1100, int(booth_target / 0.06) + 90)

        def paintEvent(self, event):
            painter = QPainter(self)
            painter.setRenderHint(QPainter.Antialiasing)
            w, h = self.width(), self.height()
            painter.fillRect(self.rect(), QColor("#06101c"))

            # ground / lanes
            painter.setPen(QPen(QColor("#17364e"), 1))
            step = 42
            for x in range(0, w, step): painter.drawLine(x, 0, x, h)
            for y in range(0, h, step): painter.drawLine(0, y, w, y)
            painter.setPen(QPen(QColor("#69849a"), 2, Qt.DashLine))
            for frac in (0.18, 0.34, 0.68, 0.84):
                y = int(h * frac); painter.drawLine(30, y, w - 30, y)

            # buildings
            painter.setPen(QPen(QColor("#38bdf8"), 2))
            painter.setBrush(QColor("#162638"))
            buildings = [(0.10,0.49,0.22,0.10,"ΤΕΛΩΝΕΙΟ"),(0.34,0.49,0.22,0.10,"ΑΣΤΥΝΟΜΙΑ"),(0.58,0.49,0.15,0.10,"ΑΞ/ΚΟΣ ΥΠΗΡΕΣΙΑΣ")]
            painter.setFont(QFont("Segoe UI", 10, QFont.Bold))
            for x,y,bw,bh,label in buildings:
                r=QRectF(w*x,h*y,w*bw,h*bh); painter.drawRoundedRect(r,8,8); painter.drawText(r,Qt.AlignCenter,label)

            # labels
            painter.setFont(QFont("Segoe UI", 12, QFont.Bold))
            painter.setPen(QColor("#4ade80")); painter.drawText(QRectF(w*0.76,h*0.18,w*0.20,h*0.08),Qt.AlignCenter,"ΕΙΣΟΔΟΣ ↑")
            painter.setPen(QColor("#f87171")); painter.drawText(QRectF(w*0.76,h*0.70,w*0.20,h*0.08),Qt.AlignCenter,"ΕΞΟΔΟΣ ↓")

            # booths
            desired_name_width = min(max(self._worker_text_width + 20, 92), 190)
            bw = max(54, min(int(w*0.075), desired_name_width))
            bh = max(52, int(h*0.075))
            for pos,(xf,yf) in self.BOOTH_POS.items():
                x=int(w*xf); y=int(h*yf)
                r=QRectF(x,y,bw,bh)
                people=self.active.get(pos,[])
                occupied=bool(people)
                # subtle pseudo-3D depth: shadow + lower/right faces, then main face
                shadow_r = QRectF(r.x()+7, r.y()+8, r.width(), r.height())
                painter.setPen(Qt.NoPen)
                painter.setBrush(QColor(0, 0, 0, 95))
                painter.drawRoundedRect(shadow_r, 9, 9)
                painter.setBrush(QColor("#0f5f3f" if occupied else "#064865"))
                painter.drawRoundedRect(QRectF(r.x()+4, r.y()+5, r.width(), r.height()), 9, 9)
                painter.setPen(QPen(QColor("#a7f3d0" if occupied else "#bae6fd"),2))
                painter.setBrush(QColor("#1f9d68" if occupied else "#0878a8"))
                painter.drawRoundedRect(r,9,9)
                # highlight on upper edge for a raised/card-like effect
                painter.setPen(QPen(QColor(255,255,255,70),1))
                painter.drawLine(int(r.x()+7), int(r.y()+4), int(r.right()-7), int(r.y()+4))
                painter.setPen(QColor("#ffffff"))
                painter.setFont(QFont("Segoe UI", 11, QFont.Bold))
                painter.drawText(QRectF(r.x(),r.y()+2,r.width(),18),Qt.AlignCenter,f"ΘΕΣΗ {pos}")
                if people:
                    names=[]
                    for item in people[:2]:
                        worker_name = self._worker_display_name(item)
                        ts=_clean(item.get("timestamp")).replace("T"," ")
                        hhmm=ts[11:16] if len(ts)>=16 else ""
                        direction=_clean(item.get("direction"))
                        names.append(f"{worker_name}\n{direction} {hhmm}".strip())
                    if len(people)>2: names.append(f"+{len(people)-2} ακόμη")
                    painter.setFont(QFont("Segoe UI", 8, QFont.Bold))
                    painter.drawText(QRectF(r.x()+3,r.y()+19,r.width()-6,r.height()-21),Qt.AlignCenter|Qt.TextWordWrap,"\n".join(names))

    class LiveMapDialog(QDialog):
        def __init__(self, store: SharedEventStore, people_by_agms: dict[str, PersonnelRecord], officer_getter, parent=None):
            super().__init__(parent)
            self.store = store
            self.people_by_agms = people_by_agms
            self.officer_getter = officer_getter
            self._live_pool = ThreadPoolExecutor(max_workers=1, thread_name_prefix="TDE-live-map")
            self._live_future = None
            self.setWindowTitle("ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ - Ζωντανή εικόνα θέσεων")
            self.resize(1100, 820)
            self.setStyleSheet("""
                QDialog{background:qlineargradient(x1:0,y1:0,x2:1,y2:1,stop:0 #eef5f9,stop:1 #dfeaf1);}
                QLabel{color:#193b50;}
            """)
            lay=QVBoxLayout(self); lay.setContentsMargins(22,22,22,22); lay.setSpacing(14)
            top=QHBoxLayout()
            title=QLabel("ΖΩΝΤΑΝΗ ΕΙΚΟΝΑ ΘΕΣΕΩΝ"); title.setStyleSheet("font-size:22px;font-weight:900;color:#123f5a;letter-spacing:1px")
            self.updated=QLabel("")
            top.addWidget(title); top.addStretch(); top.addWidget(self.updated)
            lay.addLayout(top)
            self.map=BoothMapWidget(); lay.addWidget(self.map,1)
            self.other=QLabel(""); self.other.setWordWrap(True); self.other.setStyleSheet("background:#ffffff;border:1px solid #cfdae2;padding:10px 12px;border-radius:12px;color:#36596f;font-weight:600")
            lay.addWidget(self.other)

            action_row=QHBoxLayout()
            self.active_person=QComboBox()
            self.active_person.setMinimumHeight(40)
            self.active_person.setMinimumWidth(350)
            self.end_active_btn=QPushButton("ΛΗΞΗ ΕΝΕΡΓΗΣ ΘΕΣΗΣ")
            self.end_active_btn.setMinimumHeight(40)
            self.end_active_btn.setStyleSheet("font-weight:900")
            self.end_active_btn.clicked.connect(self._end_selected_active)
            self.edit_start_btn=QPushButton("ΔΙΟΡΘΩΣΗ ΩΡΑΣ ΕΝΑΡΞΗΣ")
            self.edit_start_btn.setMinimumHeight(40)
            self.edit_start_btn.clicked.connect(self._correct_selected_start)
            action_row.addWidget(QLabel("Ενεργός αστυνομικός:"))
            action_row.addWidget(self.active_person,1)
            action_row.addWidget(self.end_active_btn)
            action_row.addWidget(self.edit_start_btn)
            lay.addLayout(action_row)

            self.timer=QTimer(self); self.timer.timeout.connect(self.refresh); self.timer.start(2000)
            self.refresh()

        def _selected_active(self) -> dict | None:
            agms=_clean(self.active_person.currentData())
            if not agms:
                return None
            return self.store.get_active(agms)

        def _correct_selected_start(self):
            officer=self.officer_getter() if callable(self.officer_getter) else None
            if not officer:
                QMessageBox.warning(self,"ΑΞΥΠ","Εισαγάγετε έγκυρο Α.Γ.Μ.Σ. ΑΞΥΠ.")
                return
            active=self._selected_active()
            if not active:
                QMessageBox.warning(self,"Διόρθωση","Επιλέξτε ενεργό αστυνομικό.")
                return
            started=datetime.fromisoformat(active["timestamp"])
            proposed,ok=QInputDialog.getText(
                self,"ΔΙΟΡΘΩΣΗ ΩΡΑΣ ΕΝΑΡΞΗΣ",
                "Πραγματική ώρα έναρξης (ΗΗ/ΜΜ/ΕΕΕΕ ΩΩ:ΛΛ):",
                QLineEdit.Normal,started.strftime("%d/%m/%Y %H:%M"),
            )
            if not ok: return
            try:
                revised=datetime.strptime(_clean(proposed),"%d/%m/%Y %H:%M")
            except ValueError:
                QMessageBox.warning(self,"Διόρθωση","Χρησιμοποιήστε ΗΗ/ΜΜ/ΕΕΕΕ ΩΩ:ΛΛ.")
                return
            reason,ok=QInputDialog.getText(
                self,"Αιτιολογία διόρθωσης",
                "Αιτιολογία (π.χ. προσωρινή διακοπή εφαρμογής):",
            )
            if not ok: return
            if QMessageBox.question(
                self,"Επιβεβαίωση",
                f"ΑΓΜΣ: {active.get('agms')}\nΠροηγούμενη: {started:%d/%m/%Y %H:%M}\n"
                f"Νέα: {revised:%d/%m/%Y %H:%M}\nΑιτία: {_clean(reason)}\n\nΝα αποθηκευτεί η διόρθωση;",
                QMessageBox.Yes | QMessageBox.No,QMessageBox.No,
            )!=QMessageBox.Yes: return
            try:
                self.store.correct_active_start(_clean(active.get("agms")),revised,officer,reason)
                self.refresh()
                QMessageBox.information(self,"Διόρθωση","Η νέα ώρα καταγράφηκε με ιστορικό αλλαγής.")
            except Exception as exc:
                QMessageBox.critical(self,"Διόρθωση",str(exc))

        def _end_selected_active(self):
            officer=self.officer_getter() if callable(self.officer_getter) else None
            if not officer:
                QMessageBox.warning(
                    self,
                    "Αξιωματικός Υπηρεσίας",
                    "Πληκτρολογήστε πρώτα έγκυρο Α.Γ.Μ.Σ. Αξιωματικού Υπηρεσίας στο κύριο παράθυρο."
                )
                return
            active=self._selected_active()
            if not active:
                QMessageBox.information(self,"Λήξη ενεργής θέσης","Η επιλεγμένη καταγραφή δεν είναι πλέον ενεργή.")
                self.refresh()
                return
            agms=_clean(active.get("agms"))
            person=self.people_by_agms.get(agms)
            if not person:
                person=PersonnelRecord(
                    agms=agms,
                    rank=_clean(active.get("rank")),
                    surname=_clean(active.get("surname")),
                    given_name=_clean(active.get("given_name")),
                )
            name=_clean(f"{person.rank} {person.surname} {person.given_name}") or f"ΑΓΜΣ {agms}"
            pos=_clean(active.get("position"))
            direction=_clean(active.get("direction"))
            started=_clean(active.get("timestamp")).replace("T"," ")
            confirm=QMessageBox.question(
                self,
                "Επιβεβαίωση λήξης ενεργής θέσης",
                f"{name}\nΘέση {pos} — {direction}\nΈναρξη: {started}\n\n"
                "Να καταχωριστεί ΛΗΞΗ από τον Αξιωματικό Υπηρεσίας;",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No
            )
            if confirm != QMessageBox.Yes:
                return
            try:
                event=self.store.end(person)
                self.store.log_admin_action("ACTIVE_POSITION_ENDED_BY_OFFICER", {
                    "officer_agms": officer.agms,
                    "officer_name": officer.full_name,
                    "target_agms": agms,
                    "target_name": person.full_name,
                    "position": pos,
                    "direction": direction,
                    "start_at": _clean(active.get("timestamp")),
                    "end_at": _clean(event.get("timestamp")),
                })
                QMessageBox.information(
                    self,
                    "Λήξη ενεργής θέσης",
                    f"Η θέση του {name} έκλεισε στις {_clean(event.get('timestamp')).replace('T',' ')}."
                )
                self.refresh()
            except Exception as exc:
                QMessageBox.critical(self,"Δεν καταχωρίστηκε",str(exc))

        def closeEvent(self, event):
            self.timer.stop()
            self._live_pool.shutdown(wait=False, cancel_futures=True)
            super().closeEvent(event)

        def refresh(self):
            future = self._live_future
            if future is None:
                self._live_future = self._live_pool.submit(self.store.active_all)
                return
            if not future.done():
                return
            self._live_future = None
            try:
                items = future.result()
            except Exception:
                self.updated.setText("ΑΠΩΛΕΙΑ ΣΥΝΔΕΣΗΣ — αναμονή επαναφοράς")
                return
            self.map.set_active(items)
            current_agms=_clean(self.active_person.currentData())
            self.active_person.blockSignals(True)
            self.active_person.clear()
            for item in items:
                agms=_clean(item.get("agms"))
                label=_clean(
                    f"{item.get('rank','')} {item.get('surname','')} {item.get('given_name','')}"
                ) or _clean(item.get("full_name")) or f"ΑΓΜΣ {agms}"
                pos=_clean(item.get("position"))
                direction=_clean(item.get("direction"))
                self.active_person.addItem(f"{label} — Θέση {pos} {direction}",agms)
            if current_agms:
                idx=self.active_person.findData(current_agms)
                if idx >= 0:
                    self.active_person.setCurrentIndex(idx)
            self.active_person.blockSignals(False)
            self.end_active_btn.setEnabled(bool(items))
            self.edit_start_btn.setEnabled(bool(items))

            recommended = self.map.recommended_dialog_width()
            screen = QApplication.primaryScreen()
            if screen:
                available = screen.availableGeometry()
                target_width = min(recommended, max(1100, available.width() - 60))
                target_height = min(820, max(700, available.height() - 80))
            else:
                target_width = min(recommended, 1600)
                target_height = 820
            if self.width() < target_width:
                self.resize(target_width, target_height)
            other=[]
            for x in items:
                pos=_clean(x.get("position"))
                if pos not in BoothMapWidget.BOOTH_POS:
                    other.append(f"{x.get('rank','')} {x.get('full_name','')} — {pos} — {x.get('direction','')} — από {_clean(x.get('timestamp')).replace('T',' ')}")
            self.other.setText("Άλλες ενεργές θέσεις: " + (" | ".join(other) if other else "καμία"))
            self.updated.setText("Ενημέρωση: " + datetime.now().strftime("%H:%M:%S"))

    class MainWindow(QMainWindow):
        def __init__(self):
            super().__init__()
            self.config = AppConfig.load()
            self.people: tuple[PersonnelRecord, ...] = ()
            self._people_by_agms: dict[str, PersonnelRecord] = {}
            self.store: SharedEventStore | None = None
            self._pending_prompt_active = False
            self._last_shift_signature = ""
            self._force_exit = False
            self._message_dialog_active = False
            self._app_started_at = datetime.now().isoformat(timespec="seconds")
            self.client_id = SharedEventStore.local_client_id()
            self._network_pool = ThreadPoolExecutor(max_workers=2, thread_name_prefix="TDE-network")
            # Network heartbeat/shift polling must never starve a booth action
            # when the SMB host disconnects or an old worker blocks in I/O.
            self._position_pool = ThreadPoolExecutor(max_workers=1, thread_name_prefix="TDE-position")
            self._state_pool = ThreadPoolExecutor(max_workers=2, thread_name_prefix="TDE-state")
            self._position_operation_future = None
            self._last_bootstrap_check = 0.0
            self._offline_future = None
            self._offline_polling = False
            self._pending_central_store = None
            self._handoff_initial_error = ""
            self.offline = OfflineJournal(_machine_queue_root())
            self._agms_query_serial = 0
            self._shift_status_future = None
            self._communication_future = None
            self._shift_future = None
            self.setWindowTitle(f"Τ.Δ.Ε. Κρυσταλλοπηγής - Καταγραφή Θέσεων v{APP_VERSION}")
            emblem_icon = _police_emblem_pixmap(96)
            if not emblem_icon.isNull():
                self.setWindowIcon(QIcon(emblem_icon))
            self.resize(1160, 720)
            self.setMinimumSize(1000, 650)
            self._build()
            self._apply_config()
            self._apply_theme()
            self.shift_timer = QTimer(self)
            self.shift_timer.timeout.connect(self._check_shift_updates)
            self.shift_timer.start(3000)
            self.communication_timer = QTimer(self)
            self.communication_timer.timeout.connect(self._communication_tick)
            self.communication_timer.start(2500)
            QApplication.instance().aboutToQuit.connect(self._cleanup_client)
            self.offline_timer = QTimer(self)
            self.offline_timer.timeout.connect(self._offline_tick)
            self.offline_timer.start(3000)
            self._offline_tick()
            self._communication_tick()

        def _build(self):
            self.root = QWidget(); self.root.setObjectName("rootPanel"); self.setCentralWidget(self.root)
            outer = QVBoxLayout(self.root); outer.setContentsMargins(28,22,28,20); outer.setSpacing(14)

            self.header_card = QFrame()
            self.header_card.setObjectName("serviceHeader")
            header = QHBoxLayout(self.header_card)
            header.setContentsMargins(20, 12, 20, 12)
            header.setSpacing(16)

            self.police_logo = QLabel()
            self.police_logo.setFixedSize(86, 86)
            self.police_logo.setAlignment(Qt.AlignCenter)
            emblem = _police_emblem_pixmap(78)
            if not emblem.isNull():
                self.police_logo.setPixmap(emblem)
            header.addWidget(self.police_logo)

            titles = QVBoxLayout()
            titles.setSpacing(1)
            self.police_title_label = QLabel("ΕΛΛΗΝΙΚΗ ΑΣΤΥΝΟΜΙΑ")
            self.police_title_label.setObjectName("serviceKicker")
            self.title_label = QLabel("ΔΙΕΥΘΥΝΣΗ ΑΣΤΥΝΟΜΙΑΣ ΦΛΩΡΙΝΑΣ")
            self.title_label.setObjectName("serviceTitle")
            self.department_label = QLabel("ΤΜΗΜΑ ΔΙΑΒΑΤΗΡΙΑΚΟΥ ΕΛΕΓΧΟΥ ΚΡΥΣΤΑΛΛΟΠΗΓΗΣ")
            self.department_label.setObjectName("serviceDepartment")
            self.subtitle_label = QLabel("ΚΑΤΑΓΡΑΦΗ ΘΕΣΕΩΝ")
            self.subtitle_label.setObjectName("appSubtitle")
            for label in (self.police_title_label, self.title_label, self.department_label, self.subtitle_label):
                label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
                titles.addWidget(label)
            header.addLayout(titles, 1)

            self.theme_btn = QToolButton()
            self.theme_btn.setMinimumWidth(108)
            self.theme_btn.setMinimumHeight(34)
            self.theme_btn.clicked.connect(self._toggle_theme)
            header.addWidget(self.theme_btn, 0, Qt.AlignTop)

            self.greece_map = GreeceMapWidget()
            header.addWidget(self.greece_map, 0, Qt.AlignRight | Qt.AlignVCenter)
            outer.addWidget(self.header_card)

            columns = QHBoxLayout(); columns.setSpacing(16)

            self.operator_card = QFrame(); self.operator_card.setObjectName("card")
            left = QVBoxLayout(self.operator_card); left.setContentsMargins(20,18,20,18); left.setSpacing(12)
            self.operator_title = QLabel("ΠΡΟΣΩΠΙΚΟ")
            self.operator_title.setAlignment(Qt.AlignCenter); self.operator_title.setStyleSheet("font-size:18px;font-weight:900")
            left.addWidget(self.operator_title)
            operator_form = QFormLayout(); operator_form.setSpacing(12)

            self.agms_input = QLineEdit(); self.agms_input.setPlaceholderText("Πληκτρολογήστε Α.Γ.Μ.Σ.")
            self.agms_input.setMinimumHeight(42); self.agms_input.textChanged.connect(self._agms_changed)
            self.person_name = QLabel("—"); self.person_name.setWordWrap(True); self.person_name.setStyleSheet("font-size:16px;font-weight:800")
            self.direction = ArrowComboBox(); self.direction.addItems(["ΕΙΣΟΔΟΣ","ΕΞΟΔΟΣ"]); self.direction.setMinimumHeight(42)
            self.position = ArrowComboBox(); self.position.setMinimumHeight(42); self.position.currentTextChanged.connect(self._position_changed)
            operator_form.addRow("Α.Γ.Μ.Σ.", self.agms_input)
            operator_form.addRow("Αστυνομικός", self.person_name)
            operator_form.addRow("ΚΑΤΕΥΘΥΝΣΗ", self.direction)
            operator_form.addRow("ΘΕΣΗ", self.position)
            left.addLayout(operator_form)

            self.active_status = QLabel("Πληκτρολογήστε Α.Γ.Μ.Σ.")
            self.active_status.setObjectName("statusBox")
            self.active_status.setAlignment(Qt.AlignCenter); self.active_status.setWordWrap(True)
            left.addWidget(self.active_status)
            left.addStretch()

            op_buttons = QHBoxLayout(); op_buttons.setSpacing(12)
            self.start_btn = QPushButton("ΕΝΑΡΞΗ"); self.start_btn.setObjectName("startButton"); self.start_btn.setMinimumHeight(54)
            self.end_btn = QPushButton("ΛΗΞΗ"); self.end_btn.setObjectName("endButton"); self.end_btn.setMinimumHeight(54)
            self.start_btn.clicked.connect(self._start); self.end_btn.clicked.connect(self._end)
            op_buttons.addWidget(self.start_btn); op_buttons.addWidget(self.end_btn)
            left.addLayout(op_buttons)

            self.officer_card = QFrame(); self.officer_card.setObjectName("card")
            right = QVBoxLayout(self.officer_card); right.setContentsMargins(22,18,22,18); right.setSpacing(13)
            self.officer_title = QLabel("ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ")
            self.officer_title.setAlignment(Qt.AlignCenter); self.officer_title.setStyleSheet("font-size:18px;font-weight:900")
            right.addWidget(self.officer_title)

            officer_form = QFormLayout(); officer_form.setSpacing(12)
            self.officer_agms = QLineEdit(); self.officer_agms.setPlaceholderText("Πληκτρολογήστε Α.Γ.Μ.Σ.")
            self.officer_agms.setMinimumHeight(42); self.officer_agms.textChanged.connect(self._officer_changed)
            self.officer_name = QLabel("—"); self.officer_name.setWordWrap(True); self.officer_name.setStyleSheet("font-size:16px;font-weight:800")
            officer_form.addRow("Α.Γ.Μ.Σ.", self.officer_agms)
            officer_form.addRow("Αξιωματικός", self.officer_name)
            right.addLayout(officer_form)

            self.shift_status = QLabel("Δεν υπάρχει ενεργή βάρδια.")
            self.shift_status.setAlignment(Qt.AlignCenter); self.shift_status.setWordWrap(True); self.shift_status.setObjectName("statusBox")
            right.addWidget(self.shift_status)
            right.addStretch()

            shift_buttons = QHBoxLayout(); shift_buttons.setSpacing(12)
            self.start_shift_btn = QPushButton("ΕΝΑΡΞΗ ΒΑΡΔΙΑΣ"); self.start_shift_btn.setObjectName("startButton"); self.start_shift_btn.setMinimumHeight(54)
            self.end_shift_btn = QPushButton("ΛΗΞΗ ΒΑΡΔΙΑΣ"); self.end_shift_btn.setObjectName("endButton"); self.end_shift_btn.setMinimumHeight(54)
            self.start_shift_btn.clicked.connect(self._start_officer_shift_ui)
            self.end_shift_btn.clicked.connect(self._end_officer_shift_ui)
            shift_buttons.addWidget(self.start_shift_btn); shift_buttons.addWidget(self.end_shift_btn)
            right.addLayout(shift_buttons)

            self.comm_status = QLabel("ΜΗΝΥΜΑΤΑ / ΟΔΗΓΙΕΣ — αναμονή σύνδεσης")
            self.comm_status.setObjectName("infoBox")
            self.comm_status.setAlignment(Qt.AlignCenter)
            self.comm_status.setWordWrap(True)
            self.messages_btn = QPushButton("ΜΗΝΥΜΑΤΑ / ΟΔΗΓΙΕΣ")
            self.messages_btn.setObjectName("metalPrimary")
            self.messages_btn.setMinimumHeight(42)
            self.messages_btn.clicked.connect(self._open_messages)
            right.addWidget(self.comm_status)
            right.addWidget(self.messages_btn)
            self.offline_conflicts_btn = QPushButton("ΕΛΕΓΧΟΣ OFFLINE ΑΣΥΜΦΩΝΙΩΝ")
            self.offline_conflicts_btn.setMinimumHeight(38)
            self.offline_conflicts_btn.clicked.connect(self._offline_conflicts_dialog)
            right.addWidget(self.offline_conflicts_btn)

            columns.addWidget(self.operator_card, 1); columns.addWidget(self.officer_card, 1)
            outer.addLayout(columns, 1)

            bottom = QHBoxLayout(); bottom.setSpacing(10)
            self.print_btn = QPushButton("ΤΡΕΧΟΥΣΑ ΑΝΑΦΟΡΑ")
            self.archive_btn = QPushButton("ΑΡΧΕΙΟ ΑΝΑΦΟΡΩΝ")
            self.history_btn = QPushButton("ΙΣΤΟΡΙΚΟ ΠΡΟΣΩΠΙΚΟΥ")
            for b in (self.print_btn, self.archive_btn, self.history_btn):
                b.setObjectName("metalPrimary")
                b.setMinimumHeight(44)
                bottom.addWidget(b)
            self.print_btn.clicked.connect(self._report)
            self.archive_btn.clicked.connect(self._report_archive)
            self.history_btn.clicked.connect(self._personnel_history)
            outer.addLayout(bottom)

            bottom2 = QHBoxLayout(); bottom2.setSpacing(10)
            self.active_positions_btn = QPushButton("ΕΝΕΡΓΕΣ ΘΕΣΕΙΣ ΕΡΓΑΣΙΑΣ")
            self.personnel_list_btn = QPushButton("ΕΠΙΛΟΓΗ ΛΙΣΤΑΣ ΠΡΟΣΩΠΙΚΟΥ")
            self.exit_app_btn = QPushButton("ΕΞΟΔΟΣ ΑΠΟ ΤΟ ΠΡΟΓΡΑΜΜΑ")
            self.active_positions_btn.setObjectName("metalPrimary")
            self.personnel_list_btn.setObjectName("metalPrimary")
            self.exit_app_btn.setObjectName("exitButton")
            for b in (self.active_positions_btn, self.personnel_list_btn, self.exit_app_btn):
                b.setMinimumHeight(44)
                bottom2.addWidget(b)
            self.active_positions_btn.clicked.connect(self._live_map)
            self.personnel_list_btn.clicked.connect(self._choose_personnel_list)
            self.exit_app_btn.clicked.connect(self._request_real_exit)
            outer.addLayout(bottom2)

            self.info = QLabel(""); self.info.setObjectName("infoBox"); self.info.setAlignment(Qt.AlignCenter); self.info.setWordWrap(True)
            outer.addWidget(self.info)
        def closeEvent(self, event):
            if self._force_exit:
                event.accept()
                return
            event.ignore()
            self.showMinimized()

        def _cleanup_client(self):
            try:
                self._network_pool.shutdown(wait=False, cancel_futures=True)
                self._position_pool.shutdown(wait=False, cancel_futures=True)
                self._state_pool.shutdown(wait=False, cancel_futures=True)
            except Exception:
                pass
            if self.store:
                try:
                    self.store.remove_client_heartbeat(self.client_id)
                except Exception:
                    pass

        def _end_own_active_positions(self) -> int:
            """Close only START records created by this exact app process."""
            if not self.store:
                return 0
            mine = [
                item for item in self.store.active_all()
                if _clean(item.get("client_id")) == self.client_id
            ]
            for item in mine:
                member = PersonnelRecord(
                    agms=_clean(item.get("agms")),
                    rank=_clean(item.get("rank")),
                    surname=_clean(item.get("surname")),
                    given_name=_clean(item.get("given_name")),
                )
                self.store.end(member, expected_start_event_id=_clean(item.get("event_id")))
            return len(mine)

        def _request_real_exit(self):
            answer = QMessageBox.question(
                self,
                "Έξοδος από το πρόγραμμα",
                "Να κλείσει πραγματικά η εφαρμογή?\\n\\n"
                "Οι ενεργές θέσεις που ανοίχθηκαν από αυτό το παράθυρο "
                "θα λάβουν ΛΗΞΗ πριν κλείσει το πρόγραμμα.\\n"
                "Αν δεν είναι διαθέσιμος ο κεντρικός φάκελος, η έξοδος "
                "δεν θα ολοκληρωθεί με ψευδή επιβεβαίωση.",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No,
            )
            if answer != QMessageBox.Yes:
                return
            try:
                self._end_own_active_positions()
            except Exception as exc:
                QMessageBox.critical(
                    self, "Η λήξη δεν καταχωρίστηκε",
                    f"Δεν επιβεβαιώθηκε η αποδέσμευση της θέσης.\\n{exc}\\n"
                    "Ελέγξτε τον κεντρικό υπολογιστή ή ζητήστε από τον ΑΞΥΠ "
                    "να ολοκληρώσει τη θέση από τη ζωντανή εικόνα."
                )
                return
            self._force_exit = True
            self._cleanup_client()
            QApplication.quit()

        def _sender_payload(self) -> dict:
            officer = self._officer_selected() if hasattr(self, "officer_agms") else None
            if officer:
                return SharedEventStore._person_payload(officer)
            return {
                "computer": socket.gethostname(),
                "user": os.environ.get("USERNAME") or os.environ.get("USER") or "",
            }

        def _client_worker_labels(self, clients: list[dict]) -> dict[str, str]:
            active_by_computer: dict[str, list[dict]] = {}
            if self.store:
                for item in self.store.active_all():
                    active_by_computer.setdefault(_clean(item.get("computer")).upper(), []).append(item)
            labels = {}
            for client in clients:
                cid = _clean(client.get("client_id"))
                computer = _clean(client.get("computer"))
                user = _clean(client.get("user"))
                people = active_by_computer.get(computer.upper(), [])
                worker_text = ""
                if people:
                    names = []
                    for item in people:
                        name = _clean(
                            f"{item.get('rank','')} {item.get('surname','')} {item.get('given_name','')}"
                        ) or _clean(item.get("full_name"))
                        pos = _clean(item.get("position"))
                        direction = _clean(item.get("direction"))
                        names.append(f"{name} — Θέση {pos} {direction}".strip())
                    worker_text = " | " + " / ".join(names)
                labels[cid] = f"{computer} — {user}{worker_text}"
            return labels

        def _communication_io(self, window_state: str):
            """Shared-folder IO must not run on the Qt GUI thread."""
            session_active = _windows_session_is_active()
            # Central closures are owned only by the offline-sync worker.
            # Running an independent heartbeat-side close here could overtake
            # the original unsent END and discard its timestamp/event ID.
            self.store.write_client_heartbeat(self.client_id, {
                "app_started_at": self._app_started_at,
                "window_state": window_state,
                "windows_session_active": session_active,
            })
            online = self.store.online_clients()
            commands = self.store.pending_commands(self.client_id)
            pending = self.store.pending_messages(self.client_id)
            return online, commands, pending

        def _communication_tick(self):
            if not self.store:
                return
            future = self._communication_future
            if future is not None:
                if not future.done():
                    return
                self._communication_future = None
                try:
                    online, commands, pending = future.result()
                    if hasattr(self, "comm_status"):
                        self.comm_status.setText(
                            f"ΜΗΝΥΜΑΤΑ / ΟΔΗΓΙΕΣ — ενεργοί σταθμοί: {len(online)}"
                        )
                except Exception:
                    if hasattr(self, "comm_status"):
                        self.comm_status.setText("ΜΗΝΥΜΑΤΑ / ΟΔΗΓΙΕΣ — ΕΛΕΓΧΟΣ ΣΥΝΔΕΣΗΣ")
                    return
                if commands:
                    command = commands[0]
                    if _clean(command.get("type")).upper() == "CLOSE_APP":
                        try:
                            self._end_own_active_positions()
                            self.store.acknowledge_command(
                                _clean(command.get("command_id")), self.client_id
                            )
                        except Exception:
                            if hasattr(self, "comm_status"):
                                self.comm_status.setText(
                                    "ΕΚΚΡΕΜΕΙ ΛΗΞΗ ΘΕΣΗΣ — δεν εκτελέστηκε απομακρυσμένο κλείσιμο"
                                )
                            return
                        self._force_exit = True
                        self._cleanup_client()
                        QApplication.quit()
                        return
                if pending and not self._message_dialog_active:
                    self._show_message_alarm(pending[0])
            if self._communication_future is None:
                state = "ΕΛΑΧΙΣΤΟΠΟΙΗΜΕΝΟ" if self.isMinimized() else "ΑΝΟΙΧΤΟ"
                self._communication_future = self._network_pool.submit(
                    self._communication_io, state
                )

        def _show_message_alarm(self, message: dict):
            if not self.store:
                return
            self._message_dialog_active = True
            try:
                priority = _clean(message.get("priority")).upper() or "ΟΔΗΓΙΑ"
                sender = message.get("sender") or {}
                sender_name = _clean(
                    f"{sender.get('rank','')} {sender.get('surname','')} {sender.get('given_name','')}"
                ) or _clean(sender.get("full_name")) or _clean(sender.get("user")) or "Αξιωματικός Υπηρεσίας"
                created = _clean(message.get("created_at")).replace("T", " ")
                text_value = _clean(message.get("text"))

                self.showNormal()
                self.raise_()
                self.activateWindow()
                if priority in {"ΟΔΗΓΙΑ", "ΕΠΕΙΓΟΝ"}:
                    QApplication.beep()

                box = QMessageBox(self)
                box.setWindowTitle(f"{priority} — ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ")
                box.setIcon(QMessageBox.Warning if priority == "ΕΠΕΙΓΟΝ" else QMessageBox.Information)
                box.setText(
                    f"<b>{priority}</b><br><br>"
                    f"<b>Από:</b> {_escape_html(sender_name)}<br>"
                    f"<b>Ώρα:</b> {_escape_html(created)}<br><br>"
                    f"{_escape_html(text_value)}"
                )
                box.setStandardButtons(QMessageBox.Ok)
                ack_button = box.button(QMessageBox.Ok)
                if ack_button:
                    ack_button.setText("ΕΛΑΒΑ ΓΝΩΣΗ")
                    ack_button.setMinimumWidth(180)
                    ack_button.setMinimumHeight(44)
                box.setWindowFlag(Qt.WindowStaysOnTopHint, True)
                box.exec()
                self.store.acknowledge_message(_clean(message.get("message_id")), self.client_id)
            finally:
                self._message_dialog_active = False

        def _resolve_message_targets(self, selector: str, online: list[dict]) -> list[str]:
            selector = _clean(selector)
            own = self.client_id
            if selector == "ALL":
                return [_clean(x.get("client_id")) for x in online if _clean(x.get("client_id")) != own]
            if selector.startswith("CLIENT:"):
                cid = selector.split(":", 1)[1]
                return [cid] if cid and cid != own else []

            wanted_computers: set[str] = set()
            if self.store:
                for item in self.store.active_all():
                    match = False
                    if selector.startswith("DIR:"):
                        match = _clean(item.get("direction")).upper() == selector.split(":", 1)[1].upper()
                    elif selector.startswith("POS:"):
                        match = _clean(item.get("position")) == selector.split(":", 1)[1]
                    if match:
                        wanted_computers.add(_clean(item.get("computer")).upper())
            return [
                _clean(client.get("client_id"))
                for client in online
                if _clean(client.get("client_id")) != own
                and _clean(client.get("computer")).upper() in wanted_computers
            ]

        def _open_messages(self):
            if not self.store:
                QMessageBox.warning(self, "Μηνύματα", "Δεν είναι διαθέσιμος ο κοινός φάκελος δεδομένων.")
                return
            officer = self._officer_selected()
            if not officer:
                QMessageBox.warning(self, "Μηνύματα", "Πληκτρολογήστε πρώτα έγκυρο Α.Γ.Μ.Σ. Αξιωματικού Υπηρεσίας.")
                return

            dlg = QDialog(self)
            dlg.setWindowTitle("ΜΗΝΥΜΑΤΑ / ΟΔΗΓΙΕΣ ΑΞΙΩΜΑΤΙΚΟΥ ΥΠΗΡΕΣΙΑΣ")
            dlg.resize(900, 680)
            lay = QVBoxLayout(dlg)
            lay.setContentsMargins(20, 18, 20, 18)
            lay.setSpacing(10)

            online = self.store.online_clients()
            labels = self._client_worker_labels(online)
            recipient = ArrowComboBox()
            recipient.addItem("ΟΛΟΙ ΟΙ ΕΝΕΡΓΟΙ ΣΤΑΘΜΟΙ", "ALL")
            recipient.addItem("ΟΛΗ Η ΕΙΣΟΔΟΣ", "DIR:ΕΙΣΟΔΟΣ")
            recipient.addItem("ΟΛΗ Η ΕΞΟΔΟΣ", "DIR:ΕΞΟΔΟΣ")
            active_positions = sorted({_clean(x.get("position")) for x in self.store.active_all() if _clean(x.get("position"))})
            for pos in active_positions:
                recipient.addItem(f"ΘΕΣΗ {pos}", f"POS:{pos}")
            for client in online:
                cid = _clean(client.get("client_id"))
                if cid and cid != self.client_id:
                    recipient.addItem(labels.get(cid, cid), f"CLIENT:{cid}")

            priority = ArrowComboBox()
            priority.addItems(["ΟΔΗΓΙΑ", "ΕΠΕΙΓΟΝ", "ΕΝΗΜΕΡΩΣΗ"])
            editor = QPlainTextEdit()
            editor.setPlaceholderText("Γράψτε τι πρέπει να γνωρίζουν ή να κάνουν οι παραλήπτες...")
            editor.setMaximumHeight(130)

            form = QFormLayout()
            form.addRow("Παραλήπτες", recipient)
            form.addRow("Τύπος", priority)
            form.addRow("Μήνυμα / οδηγία", editor)
            lay.addLayout(form)

            send_btn = QPushButton("ΑΠΟΣΤΟΛΗ")
            send_btn.setMinimumHeight(44)
            lay.addWidget(send_btn)

            online_label = QLabel(f"Ενεργοί σταθμοί αυτή τη στιγμή: {len(online)}")
            lay.addWidget(online_label)
            sent_list = QListWidget()
            lay.addWidget(sent_list, 1)

            def refresh_sent():
                sent_list.clear()
                for msg in self.store.recent_sent_messages(20):
                    targets = [_clean(x) for x in msg.get("targets", []) if _clean(x)]
                    acks = msg.get("acknowledgements", [])
                    created = _clean(msg.get("created_at")).replace("T", " ")
                    target_label = _clean(msg.get("target_label")) or "—"
                    body = _clean(msg.get("text"))
                    if len(body) > 90:
                        body = body[:87] + "..."
                    sent_list.addItem(
                        f"{created} | {msg.get('priority','')} | {target_label} | "
                        f"ΕΛΑΒΑΝ {len(acks)}/{len(targets)} | {body}"
                    )

            def do_send():
                target_ids = self._resolve_message_targets(_clean(recipient.currentData()), online)
                try:
                    self.store.send_message(
                        SharedEventStore._person_payload(officer),
                        target_ids,
                        editor.toPlainText(),
                        priority.currentText(),
                        recipient.currentText(),
                    )
                    editor.clear()
                    refresh_sent()
                except Exception as exc:
                    QMessageBox.warning(dlg, "Αποστολή", str(exc))

            send_btn.clicked.connect(do_send)
            refresh_sent()
            dlg.exec()

        def _close_remaining_clients_dialog(self):
            if not self.store:
                return
            clients = [
                c for c in self.store.online_clients()
                if _clean(c.get("client_id")) and _clean(c.get("client_id")) != self.client_id
            ]
            if not clients:
                QMessageBox.information(self, "Τερματικά", "Δεν υπάρχουν άλλα ενεργά τερματικά.")
                return

            labels = self._client_worker_labels(clients)
            dlg = QDialog(self)
            dlg.setWindowTitle("ΕΝΕΡΓΑ ΤΕΡΜΑΤΙΚΑ ΜΕΤΑ ΤΗ ΛΗΞΗ ΒΑΡΔΙΑΣ")
            dlg.resize(820, 520)
            lay = QVBoxLayout(dlg)
            title = QLabel(
                "Οι παρακάτω εφαρμογές παραμένουν ανοικτές.\n"
                "Επιλέξτε ποιες θέλετε να κλείσετε ή κλείστε τες όλες."
            )
            title.setWordWrap(True)
            lay.addWidget(title)
            listing = QListWidget()
            for client in clients:
                cid = _clean(client.get("client_id"))
                item = QListWidgetItem(labels.get(cid, cid))
                item.setData(Qt.UserRole, cid)
                item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
                item.setCheckState(Qt.Unchecked)
                listing.addItem(item)
            lay.addWidget(listing, 1)

            buttons = QHBoxLayout()
            selected_btn = QPushButton("ΚΛΕΙΣΕ ΕΠΙΛΕΓΜΕΝΟΥΣ")
            all_btn = QPushButton("ΚΛΕΙΣΕ ΟΛΟΥΣ")
            leave_btn = QPushButton("ΑΣΤΑ ΑΝΟΙΧΤΑ")
            buttons.addWidget(selected_btn)
            buttons.addWidget(all_btn)
            buttons.addStretch()
            buttons.addWidget(leave_btn)
            lay.addLayout(buttons)

            def send_to(ids: list[str]):
                ids = [x for x in ids if x]
                if not ids:
                    QMessageBox.warning(dlg, "Τερματικά", "Δεν έχει επιλεγεί κανένα τερματικό.")
                    return
                sender = self._sender_payload()
                for cid in ids:
                    self.store.send_close_command(sender, cid, "Λήξη βάρδιας από Αξιωματικό Υπηρεσίας")
                QMessageBox.information(
                    dlg,
                    "Τερματικά",
                    f"Εστάλη εντολή κλεισίματος σε {len(ids)} τερματικά. "
                    "Θα κλείσουν μόλις παραλάβουν την εντολή."
                )
                dlg.accept()

            def close_selected():
                ids = []
                for i in range(listing.count()):
                    item = listing.item(i)
                    if item.checkState() == Qt.Checked:
                        ids.append(_clean(item.data(Qt.UserRole)))
                send_to(ids)

            selected_btn.clicked.connect(close_selected)
            all_btn.clicked.connect(lambda: send_to([_clean(x.get("client_id")) for x in clients]))
            leave_btn.clicked.connect(dlg.reject)
            dlg.exec()

        def _toggle_theme(self):
            self.config.theme = "light" if _clean(self.config.theme).lower() == "dark" else "dark"
            self.config.save()
            self._apply_theme()

        def _apply_theme(self):
            dark = _clean(self.config.theme).lower() == "dark"
            t = THEME_DARK if dark else THEME_LIGHT

            with coherent_widget_update(self):
                palette = self.palette()
                palette.setColor(QPalette.Window, QColor(t["window"]))
                self.setPalette(palette)
                self.setAutoFillBackground(True)

                root_palette = self.root.palette()
                root_palette.setColor(QPalette.Window, QColor(t["window"]))
                self.root.setPalette(root_palette)
                self.root.setAutoFillBackground(True)

                self.setStyleSheet(_theme_stylesheet(self.config.theme))
                self.theme_btn.setText("☀ ΦΩΤΕΙΝΟ" if dark else "☾ ΣΚΟΤΕΙΝΟ")
                if hasattr(self, "greece_map"):
                    self.greece_map.set_dark(dark)
                self._agms_changed()
                self._officer_changed()

        def _find_personnel_file(self) -> Path | None:
            candidates = []
            # Cached roster first, so a powered-off shared host never blocks
            # the operator window during Windows account login.
            cache = _local_root() / "roster"
            for name in ("ΠΡΟΣΩΠΙΚΟ.xlsx", "ΠΡΟΣΩΠΙΚΟ.xlsm", "PROSOPIKO.xlsx", "PROSOPIKO.xlsm"):
                candidates.append(cache / name)
            legacy = Path(self.config.personnel_file) if _clean(self.config.personnel_file) else None
            if legacy:
                candidates.append(legacy)
            for base in (_app_dir() / "CONFIG", _app_dir(), Path(self.config.central_root)):
                for name in ("ΠΡΟΣΩΠΙΚΟ.xlsx", "ΠΡΟΣΩΠΙΚΟ.xlsm", "PROSOPIKO.xlsx", "PROSOPIKO.xlsm"):
                    candidates.append(base / name)
            for path in candidates:
                if path.is_file():
                    return path
            return None

        def _apply_config(self):
            self.position.clear(); self.position.addItems(self.config.positions or [])
            self._position_changed(self.position.currentText())
            if not self.config.central_root:
                self.config.central_root = str(_app_dir() / ".tde_runtime")
            # Never touch a possibly disconnected UNC path on the GUI thread.
            self.store = None
            try:
                self.offline.auto_close_previous_windows_users(
                    socket.gethostname(), os.environ.get("USERNAME") or os.environ.get("USER") or "")
            except Exception as exc:
                self._handoff_initial_error = str(exc)
            if getattr(sys, "frozen", False) and not self.config.shared_exe_path:
                self.config.shared_exe_path = str(Path(sys.executable).resolve())
            personnel = self._find_personnel_file()
            if personnel:
                try:
                    self._load_people(str(personnel))
                    if not self.config.personnel_file:
                        self.config.personnel_file = str(personnel)
                except Exception as exc:
                    self.info.setText(f"Πρόβλημα προσωπικού: {exc}")
            self.config.save()
            self._refresh_info(); self._agms_changed(); self._refresh_shift_status()

        def _refresh_info(self):
            if self._handoff_initial_error:
                self.info.setText(
                    "ΕΚΚΡΕΜΕΙ Η ΑΛΛΑΓΗ ΧΡΗΣΤΗ WINDOWS — " + self._handoff_initial_error
                )
                return
            if not self.store:
                self.info.setText("ΧΩΡΙΣ ΚΕΝΤΡΙΚΟ — τοπικές εκκρεμείς εγγραφές, θα γίνει συγχρονισμός μόλις επιστρέψει η σύνδεση.")
            elif self.offline.conflicts():
                self.info.setText("ΑΣΥΜΦΩΝΙΑ ΣΥΓΧΡΟΝΙΣΜΟΥ — απαιτείται έλεγχος ΑΞΥΠ.")
            elif self.offline.pending_count():
                self.info.setText(f"ΑΝΑΜΟΝΗ ΣΥΓΧΡΟΝΙΣΜΟΥ: {self.offline.pending_count()} καταχωρίσεις.")
            elif self.store and self.people:
                self.info.setText(f"Σύστημα έτοιμο — Λίστα προσωπικού: {len(self.people)} άτομα")
            else:
                self.info.setText("Δεν έχει επιλεγεί λίστα προσωπικού.")

        def _save_config(self):
            self.config.positions = [self.position.itemText(i) for i in range(self.position.count()) if _clean(self.position.itemText(i))]
            self.config.save(); self._refresh_info()

        def _load_people(self, path: str):
            self.people = PersonnelRepository.load(path)
            self._people_by_agms = {_clean(p.agms): p for p in self.people}
            try:
                cached = _local_root() / "roster" / Path(path).name
                cached.parent.mkdir(parents=True, exist_ok=True)
                if Path(path).resolve() != cached.resolve():
                    shutil.copy2(path, cached)
            except OSError:
                pass
            self._agms_changed(); self._officer_changed()

        def _choose_personnel_list(self):
            start_dir = str(Path(self.config.personnel_file).parent) if _clean(self.config.personnel_file) else str(_app_dir())
            path, _ = QFileDialog.getOpenFileName(self, "Επιλογή λίστας προσωπικού", start_dir, "Excel (*.xlsx *.xlsm)")
            if not path:
                return
            try:
                self._load_people(path)
                self.config.personnel_file = path
                self.config.save()
                if self.store:
                    self.store.log_admin_action("PERSONNEL_LIST_SELECTED", {"file": path, "people": len(self.people)})
                self._refresh_info()
                QMessageBox.information(self, "Λίστα προσωπικού", f"Επιλέχθηκε λίστα με {len(self.people)} άτομα.\nΘα χρησιμοποιείται αυτόματα στις επόμενες εκκινήσεις.")
            except Exception as exc:
                QMessageBox.critical(self, "Λίστα προσωπικού", f"Δεν ήταν δυνατή η ανάγνωση του αρχείου.\n{exc}")

        def _selected(self) -> PersonnelRecord | None:
            return self._people_by_agms.get(_clean(self.agms_input.text()))

        def _officer_selected(self) -> PersonnelRecord | None:
            return self._people_by_agms.get(_clean(self.officer_agms.text()))

        def _officer_changed(self, *_):
            p = self._officer_selected()
            self.officer_name.setText(_clean(f"{p.rank} {p.surname} {p.given_name}") if p else "—")
            self._refresh_shift_status()

        def _agms_changed(self, *_):
            self._agms_query_serial = getattr(self, "_agms_query_serial", 0) + 1
            serial = self._agms_query_serial
            person = self._selected()
            self.start_btn.setEnabled(False)
            self.end_btn.setEnabled(False)
            if not person:
                self.person_name.setText("—")
                self.active_status.setText(
                    "Ο Α.Γ.Μ.Σ. δεν βρέθηκε στο προσωπικό"
                    if _clean(self.agms_input.text()) else "Πληκτρολογήστε Α.Γ.Μ.Σ."
                )
                return
            self.person_name.setText(_clean(f"{person.rank} {person.surname} {person.given_name}"))
            if self._position_operation_future is not None:
                self.active_status.setText("ΚΑΤΑΧΩΡΙΣΗ ΣΕ ΕΞΕΛΙΞΗ — μην την επαναλάβετε.")
                return
            local = self.offline.local_active(person.agms)
            if self.offline.conflicts():
                self.active_status.setText("ΑΣΥΜΦΩΝΙΑ ΣΥΓΧΡΟΝΙΣΜΟΥ — ΕΝΗΜΕΡΩΣΤΕ ΤΟΝ ΑΞΥΠ")
                return
            if not self.store or self.offline.pending_count():
                self.active_status.setText(
                    (f"ΤΟΠΙΚΗ ΘΕΣΗ {local.get('position','')} — ΕΚΚΡΕΜΕΙ ΚΕΝΤΡΙΚΗ ΕΠΙΒΕΒΑΙΩΣΗ"
                     if local else "ΕΚΤΟΣ ΣΥΝΔΕΣΗΣ — νέα έναρξη θα παραμείνει ΠΡΟΣΩΡΙΝΗ")
                )
                self.end_btn.setEnabled(bool(local))
                self.start_btn.setEnabled(not bool(local))
                return
            self.active_status.setText("ΕΛΕΓΧΟΣ ΕΝΕΡΓΗΣ ΘΕΣΗΣ — αναμονή κεντρικού.")
            def load_remote():
                record = self.store.get_active(person.agms)
                self.offline.remember(person.agms, record)
                return record
            future = self._state_pool.submit(load_remote)
            query_started = time.monotonic()

            def poll():
                if serial != self._agms_query_serial:
                    return
                if not future.done():
                    if time.monotonic() - query_started > 3.0:
                        local = self.offline.local_active(person.agms)
                        self.active_status.setText("ΔΙΚΤΥΟ ΧΩΡΙΣ ΑΠΟΚΡΙΣΗ — τοπική λειτουργία με εκκρεμή επιβεβαίωση")
                        self.start_btn.setEnabled(not bool(local))
                        self.end_btn.setEnabled(bool(local))
                        return
                    QTimer.singleShot(200, poll)
                    return
                try:
                    active = future.result()
                except Exception:
                    local = self.offline.local_active(person.agms)
                    self.active_status.setText("ΧΩΡΙΣ ΣΥΝΔΕΣΗ — τοπικές εγγραφές ΕΚΚΡΕΜΟΥΝ")
                    self.start_btn.setEnabled(not bool(local))
                    self.end_btn.setEnabled(bool(local))
                    return
                if self._position_operation_future is not None:
                    return
                if active:
                    ts = _clean(active.get("timestamp", "")).replace("T", " ")
                    self.active_status.setText(
                        f"ΕΝΕΡΓΗ: ΘΕΣΗ {active.get('position','')} — "
                        f"{active.get('direction','')} — από {ts}"
                    )
                    self.end_btn.setEnabled(True)
                else:
                    self.active_status.setText("Δεν υπάρχει ανοικτή καταχώριση")
                    self.start_btn.setEnabled(True)

            QTimer.singleShot(100, poll)

        @staticmethod
        def _suggested_direction(position: str) -> str:
            pos = _clean(position)
            if pos.isdigit():
                n = int(pos)
                if 1 <= n <= 7: return "ΕΙΣΟΔΟΣ"
                if 8 <= n <= 14: return "ΕΞΟΔΟΣ"
            return "ΕΙΣΟΔΟΣ"

        def _position_changed(self, position: str):
            suggested = self._suggested_direction(position)
            idx = self.direction.findText(suggested)
            if idx >= 0: self.direction.setCurrentIndex(idx)

        def _offline_conflicts_dialog(self):
            officer = self._officer_selected()
            if not officer:
                QMessageBox.warning(self, "ΑΞΥΠ", "Συμπληρώστε πρώτα έγκυρο ΑΓΜΣ ΑΞΥΠ.")
                return
            conflicts = self.offline.conflicts()
            if not conflicts:
                QMessageBox.information(self, "Offline", "Δεν υπάρχουν εκκρεμείς ασυμφωνίες.")
                return
            dlg = QDialog(self)
            dlg.setWindowTitle("ΑΞΥΠ — ΑΣΥΜΦΩΝΙΕΣ ΣΥΓΧΡΟΝΙΣΜΟΥ")
            dlg.resize(820, 520)
            lay = QVBoxLayout(dlg)
            label = QLabel("Η παρακάτω τοπική εγγραφή ΔΕΝ καταχωρίστηκε κεντρικά. "
                           "Ελέγξτε την πραγματική εργασία πριν απορρίψετε την εγγραφή.")
            label.setWordWrap(True)
            lay.addWidget(label)
            listing = QListWidget()
            for rec in conflicts:
                listing.addItem(f"{rec['created_at']} | ΑΓΜΣ {rec['agms']} | "
                                f"{rec['kind']} | ΘΕΣΗ {rec['position']} | {rec['error']}")
                listing.item(listing.count()-1).setData(Qt.UserRole, rec['id'])
            lay.addWidget(listing, 1)
            row = QHBoxLayout()
            reject = QPushButton("ΑΠΟΡΡΙΨΗ ΜΕ ΑΙΤΙΟΛΟΓΙΑ")
            close = QPushButton("ΚΛΕΙΣΙΜΟ")
            row.addWidget(reject); row.addWidget(close); lay.addLayout(row)
            def resolve():
                item = listing.currentItem()
                if not item:
                    return
                reason, ok = QInputDialog.getText(dlg, "Αιτιολογία", "Γιατί απορρίπτεται η τοπική εγγραφή;")
                if not ok: return
                if QMessageBox.question(dlg, "ΕΠΙΒΕΒΑΙΩΣΗ",
                    "Η καταχώριση δεν θα μεταφερθεί κεντρικά. Διατηρείται στο ιστορικό. Συνέχεια;",
                    QMessageBox.Yes | QMessageBox.No, QMessageBox.No) != QMessageBox.Yes:
                    return
                try:
                    self.offline.resolve_conflict(_clean(item.data(Qt.UserRole)), officer.agms, reason)
                    dlg.accept()
                    self._refresh_info()
                    self._agms_changed()
                except Exception as exc:
                    QMessageBox.critical(dlg, "Ασυμφωνία", str(exc))
            reject.clicked.connect(resolve)
            close.clicked.connect(dlg.reject)
            dlg.exec()

        def _offline_tick(self):
            if self._offline_future is not None:
                if not self._offline_future.done():
                    return
                future = self._offline_future
                self._offline_future = None
                try:
                    connected, result = future.result()
                    self.store = connected
                    self._handoff_initial_error = ""
                    if result.get("conflicts"):
                        self.info.setText("ΑΣΥΜΦΩΝΙΑ ΑΠΟΣΥΝΔΕΔΕΜΕΝΩΝ ΚΑΤΑΧΩΡΙΣΕΩΝ — έλεγχος ΑΞΥΠ")
                    self._refresh_info()
                except Exception as exc:
                    self.store = None
                    self._refresh_info()
                    self.info.setText(
                        "ΔΕΝ ΟΛΟΚΛΗΡΩΘΗΚΕ Ο ΣΥΓΧΡΟΝΙΣΜΟΣ / HANDOFF: "
                        + str(exc)[:220]
                        + " — οι εκκρεμείς εγγραφές παραμένουν αποθηκευμένες."
                    )
                self._agms_changed()
            if self._offline_future is not None:
                return
            if not self.config.central_root:
                return
            def work():
                if not Path(self.config.central_root).is_dir():
                    raise BusyError("Ο κεντρικός φάκελος δεν είναι διαθέσιμος.")
                store = SharedEventStore(self.config.central_root)
                report = reconcile_handoff_before_replay(
                    store, self.offline, PersonnelRecord, socket.gethostname(),
                    os.environ.get("USERNAME") or os.environ.get("USER") or "",
                )
                if self.config.shared_exe_path and time.monotonic() - self._last_bootstrap_check > 300:
                    self._last_bootstrap_check = time.monotonic()
                    try:
                        cache_windows_client(_local_root(), self.config.shared_exe_path)
                    except Exception:
                        pass  # cache update must not block shifts
                return store, report
            self._offline_future = self._network_pool.submit(work)

        def _position_action(self, kind: str, person: PersonnelRecord, position: str = "", direction: str = ""):
            if self._handoff_initial_error:
                QMessageBox.warning(
                    self, "ΑΝΑΜΟΝΗ ΣΥΜΦΩΝΙΑΣ WINDOWS",
                    "Δεν επιβεβαιώθηκε η λήξη του προηγούμενου χρήστη. "
                    + self._handoff_initial_error,
                )
                return
            if self._position_operation_future is not None:
                QMessageBox.warning(self, "Καταχώριση", "Υπάρχει ήδη εντολή σε εξέλιξη.")
                return
            self.start_btn.setEnabled(False)
            self.end_btn.setEnabled(False)
            self.active_status.setText("ΤΟΠΙΚΗ ΑΠΟΘΗΚΕΥΣΗ ΚΑΤΑΧΩΡΙΣΗΣ…")
            def work():
                # Never wait on a sleeping/disconnected SMB host while a booth
                # action is being recorded. Store the operation locally first;
                # a separate network worker performs idempotent central replay.
                result = self.offline.enqueue(kind, person, position, direction, self.client_id)
                result["pending_sync"] = True
                result["conflict"] = bool(self.offline.conflicts())
                return result
            future = self._position_pool.submit(work)
            self._position_operation_future = future
            def poll():
                if not future.done():
                    QTimer.singleShot(200, poll)
                    return
                self._position_operation_future = None
                try:
                    outcome = future.result()
                except Exception as exc:
                    QMessageBox.critical(self, "Δεν καταχωρίστηκε", str(exc))
                else:
                    if outcome.get("conflict"):
                        QMessageBox.warning(self, "ΑΣΥΜΦΩΝΙΑ", "Απαιτείται έλεγχος ΑΞΥΠ. Η τοπική εγγραφή δεν διαγράφηκε.")
                    elif outcome["pending_sync"]:
                        QMessageBox.warning(self, "ΕΚΚΡΕΜΕΙ ΣΥΓΧΡΟΝΙΣΜΟΣ", "Η καταχώριση φυλάχθηκε ΤΟΠΙΚΑ. Δεν έχει ακόμη επιβεβαιωθεί κεντρικά.")
                    else:
                        QMessageBox.information(self, "Καταχώριση", f"{kind}: Η καταχώριση επιβεβαιώθηκε κεντρικά.")
                self._refresh_info()
                self._agms_changed()
                # Launch central reconciliation asynchronously; the button
                # remains responsive regardless of central SMB availability.
                self._offline_tick()
            QTimer.singleShot(100, poll)

        def _start(self):
            person = self._selected()
            if not person:
                QMessageBox.warning(self, "Έλεγχος", "Επιλέξτε λίστα προσωπικού.")
                return
            position = _clean(self.position.currentText())
            direction = _clean(self.direction.currentText()).upper()
            if position not in {str(i) for i in range(1, 15)}:
                QMessageBox.warning(self, "Θέση", "Επιλέξτε θέση 1 έως 14.")
                return
            self._position_action("START", person, position, direction)

        def _end(self):
            person = self._selected()
            if not person:
                QMessageBox.warning(self, "Έλεγχος", "Επιλέξτε λίστα προσωπικού.")
                return
            self._position_action("END", person)

        def _preview_list_dialog(self, title: str, intro: str, agms_list: list[str], allow_exclude: bool) -> tuple[bool,list[str],list[str]]:
            dlg=QDialog(self); dlg.setWindowTitle(title); dlg.resize(620,480)
            lay=QVBoxLayout(dlg); lay.setContentsMargins(20,18,20,18); lay.setSpacing(10)
            heading=QLabel(title); heading.setStyleSheet("font-size:18px;font-weight:900")
            text_label=QLabel(intro); text_label.setWordWrap(True)
            lay.addWidget(heading); lay.addWidget(text_label)
            listing=QListWidget()
            for agms in agms_list:
                p=self._people_by_agms.get(_clean(agms))
                name=_clean(f"{p.rank} {p.surname} {p.given_name}") if p else f"ΑΓΜΣ {agms}"
                item=QListWidgetItem(f"{agms} — {name}")
                item.setData(Qt.UserRole,_clean(agms))
                if allow_exclude:
                    item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
                    item.setCheckState(Qt.Checked)
                listing.addItem(item)
            lay.addWidget(listing,1)
            buttons=QHBoxLayout(); cancel=QPushButton("ΑΚΥΡΩΣΗ"); confirm=QPushButton("ΕΠΙΒΕΒΑΙΩΣΗ")
            cancel.clicked.connect(dlg.reject); confirm.clicked.connect(dlg.accept)
            buttons.addStretch(); buttons.addWidget(cancel); buttons.addWidget(confirm); lay.addLayout(buttons)
            if dlg.exec()!=QDialog.Accepted:
                return False,[],[]
            approved=[]; excluded=[]
            for i in range(listing.count()):
                item=listing.item(i); agms=_clean(item.data(Qt.UserRole))
                if allow_exclude and item.checkState()!=Qt.Checked: excluded.append(agms)
                else: approved.append(agms)
            return True,approved,excluded

        def _start_officer_shift_ui(self):
            if not self.store:
                QMessageBox.warning(self,"Έλεγχος","Δεν είναι διαθέσιμος ο φάκελος δεδομένων."); return
            p=self._officer_selected()
            if not p:
                QMessageBox.warning(self,"Έλεγχος","Πληκτρολογήστε έγκυρο Α.Γ.Μ.Σ. Αξιωματικού Υπηρεσίας."); return
            if self.store.current_shift():
                QMessageBox.warning(self,"Βάρδια","Υπάρχει ήδη ενεργή βάρδια."); return
            active=self.store.active_all()
            agms=[]
            for item in active:
                a=_clean(item.get("agms"))
                if a and a not in agms: agms.append(a)
            ok,approved,excluded=self._preview_list_dialog(
                "ΠΡΟΕΠΙΣΚΟΠΗΣΗ ΕΝΑΡΞΗΣ ΒΑΡΔΙΑΣ",
                "Ελέγξτε το προσωπικό που βρίσκεται ήδη σε ενεργή θέση. Αποεπιλέξτε όποιον δεν πρέπει να συμπεριληφθεί στη νέα βάρδια.",
                agms, True
            )
            if not ok: return
            try:
                shift=self.store.start_officer_shift(p,initial_approved=approved,initial_excluded=excluded)
                self._last_shift_signature=""
                self._refresh_shift_status()
                QMessageBox.information(self,"Βάρδια",f"Η βάρδια ξεκίνησε στις {shift['start_at'].replace('T',' ')}.")
            except Exception as exc:
                QMessageBox.critical(self,"Βάρδια",str(exc))

        def _resolve_pending_members(self) -> bool:
            if not self.store or self._pending_prompt_active:
                return False
            shift=self.store.sync_active_shift()
            pending=list(shift.get("pending",[])) if shift else []
            if not pending:
                return True
            self._pending_prompt_active=True
            try:
                for agms in pending:
                    p=self._people_by_agms.get(_clean(agms))
                    name=_clean(f"{p.rank} {p.surname} {p.given_name}") if p else f"ΑΓΜΣ {agms}"
                    answer=QMessageBox.question(
                        self,
                        "Νέο άτομο στη βάρδια",
                        f"{name}\n\nΚαταγράφηκε σε θέση μετά την έναρξη της βάρδιας.\nΝα συμπεριληφθεί στη βάρδια;",
                        QMessageBox.Yes | QMessageBox.No,
                        QMessageBox.Yes
                    )
                    if answer == QMessageBox.Yes:
                        self.store.set_shift_member_state(agms, "APPROVED")
                    else:
                        confirm_exclusion = QMessageBox.question(
                            self,
                            "Επιβεβαίωση εξαίρεσης",
                            f"{name}\n\nΕίστε βέβαιοι ότι θέλετε να ΕΞΑΙΡΕΣΕΤΕ το άτομο από τη βάρδια;",
                            QMessageBox.Yes | QMessageBox.No,
                            QMessageBox.No
                        )
                        if confirm_exclusion == QMessageBox.Yes:
                            self.store.set_shift_member_state(agms, "EXCLUDED")
            finally:
                self._pending_prompt_active=False
            self._refresh_shift_status()
            return True

        def _confirm_and_end_officer_shift(self, officer: PersonnelRecord) -> bool:
            """Per-person END and group END inside shift close; live map unchanged."""
            if not self.store:
                return False
            previous = getattr(self, "_shift_closure_future", None)
            if previous is not None and not previous.done():
                QMessageBox.warning(
                    self, "Εκκρεμής καταχώριση",
                    "Υπάρχει προηγούμενη εντολή λήξης σε εξέλιξη. "
                    "Ελέγξτε ξανά την κατάσταση αφού ολοκληρωθεί.",
                )
                return False

            dlg = QDialog(self)
            dlg.setWindowTitle("ΛΗΞΗ ΒΑΡΔΙΑΣ — ΕΚΚΡΕΜΕΙΣ ΘΕΣΕΙΣ")
            dlg.resize(920, 670)
            lay = QVBoxLayout(dlg)
            lay.setContentsMargins(20, 18, 20, 18)
            lay.setSpacing(12)
            heading = QLabel("ΑΣΤΥΝΟΜΙΚΟΙ ΠΟΥ ΔΕΝ ΕΧΟΥΝ ΚΑΝΕΙ ΛΗΞΗ")
            heading.setStyleSheet("font-size:17px;font-weight:800")
            lay.addWidget(heading)
            status = QLabel("Έλεγχος ενεργών θέσεων…")
            status.setWordWrap(True)
            lay.addWidget(status)
            pending_note = QLabel("")
            pending_note.setWordWrap(True)
            lay.addWidget(pending_note)
            positions_list = QListWidget()
            positions_list.setAlternatingRowColors(True)
            lay.addWidget(positions_list, 1)
            excluded_note = QLabel("")
            excluded_note.setWordWrap(True)
            lay.addWidget(excluded_note)
            footer = QHBoxLayout()
            refresh_btn = QPushButton("ΑΝΑΝΕΩΣΗ")
            group_btn = QPushButton("ΟΜΑΔΙΚΗ ΛΗΞΗ")
            finalize_btn = QPushButton("ΟΛΟΚΛΗΡΩΣΗ ΒΑΡΔΙΑΣ")
            cancel_btn = QPushButton("ΚΛΕΙΣΙΜΟ")
            for button in (refresh_btn, group_btn, finalize_btn, cancel_btn):
                button.setMinimumHeight(42)
                footer.addWidget(button)
            lay.addLayout(footer)

            worker = ThreadPoolExecutor(max_workers=1, thread_name_prefix="TDE-shift-close")
            state = {"future": None, "kind": "", "shift": None,
                     "active": [], "excluded": [], "row_buttons": [],
                     "finalized": False, "pending": False}

            def full_name(item: dict) -> str:
                name = _clean(
                    f"{item.get('rank', '')} {item.get('surname', '')} {item.get('given_name', '')}"
                )
                return name or _clean(item.get("full_name")) or f"ΑΓΜΣ {_clean(item.get('agms'))}"

            def position_label(item: dict) -> str:
                return (
                    f"{full_name(item)}  |  ΘΕΣΗ {_clean(item.get('position'))}"
                    f" — {_clean(item.get('direction'))}  |  "
                    f"ΕΝΑΡΞΗ {_clean(item.get('timestamp')).replace('T', ' ')}"
                )

            def set_busy(message: str):
                status.setText(message)
                refresh_btn.setEnabled(False)
                group_btn.setEnabled(False)
                finalize_btn.setEnabled(False)
                for button in state["row_buttons"]:
                    button.setEnabled(False)

            def submit(kind: str, task):
                if state["future"] is not None:
                    return
                set_busy("Επικοινωνία με τον κεντρικό φάκελο — αναμονή επιβεβαίωσης…")
                state["kind"] = kind
                future = worker.submit(task)
                state["future"] = future
                if kind in {"single", "finalize"}:
                    self._shift_closure_future = future

            def load_data():
                shift = self.store.sync_active_shift()
                active = self.store.active_all() if shift else []
                return shift, active

            def reload():
                submit("load", load_data)

            def selected_member_end(item: dict):
                if state["future"] is not None:
                    return
                name = full_name(item)
                confirm = QMessageBox.question(
                    dlg, "ΜΕΜΟΝΩΜΕΝΗ ΛΗΞΗ",
                    f"{position_label(item)}\n\n"
                    f"Να καταχωριστεί ΛΗΞΗ μόνο για τον/την {name};\n"
                    "Η βάρδια θα παραμείνει ανοικτή.",
                    QMessageBox.Yes | QMessageBox.No, QMessageBox.No,
                )
                if confirm != QMessageBox.Yes:
                    return
                shift_id = _clean(state["shift"].get("shift_id"))
                agms = _clean(item.get("agms"))
                start_id = _clean(item.get("event_id"))
                submit(
                    "single",
                    lambda: self.store.end_officer_shift_member(
                        officer, agms, shift_id, start_id
                    ),
                )

            def finalize_shift(group: bool):
                if state["future"] is not None or not state["shift"]:
                    return
                pending = bool(state["shift"].get("pending"))
                if pending:
                    QMessageBox.warning(
                        dlg, "Άτομα σε αναμονή",
                        "Υπάρχουν άτομα σε ΑΝΑΜΟΝΗ. "
                        "Ολοκληρώστε πρώτα την ΑΠΟΔΟΧΗ ή ΕΞΑΙΡΕΣΗ από τον πίνακα ΑΞΥΠ.",
                    )
                    return
                items = list(state["active"])
                if not group and items:
                    return
                if group and not items:
                    return
                if group:
                    names = "\n".join(f"• {position_label(x)}" for x in items)
                    question = (
                        "Να καταχωριστεί ΟΜΑΔΙΚΗ ΛΗΞΗ στους παρακάτω αστυνομικούς;\n\n"
                        f"{names}\n\n"
                        "Θα ολοκληρωθεί και η βάρδια. Οι ΕΞΑΙΡΕΜΕΝΟΙ δεν θα επηρεαστούν."
                    )
                else:
                    question = (
                        "Δεν υπάρχουν πλέον ανοικτές θέσεις στους ΑΠΟΔΕΚΤΟΥΣ.\n\n"
                        "Να ολοκληρωθεί η βάρδια;"
                    )
                confirm = QMessageBox.question(
                    dlg, "ΕΠΙΒΕΒΑΙΩΣΗ ΛΗΞΗΣ ΒΑΡΔΙΑΣ", question,
                    QMessageBox.Yes | QMessageBox.No, QMessageBox.No,
                )
                if confirm != QMessageBox.Yes:
                    return
                shift_id = _clean(state["shift"].get("shift_id"))
                snapshot = {
                    _clean(x.get("agms")): _clean(x.get("event_id"))
                    for x in items
                }
                submit(
                    "finalize",
                    lambda: self.store.end_officer_shift(
                        officer, expected_active_start_ids=snapshot,
                        expected_shift_id=shift_id,
                    ),
                )

            def render(shift: dict | None, active: list[dict]):
                positions_list.clear()
                state["row_buttons"] = []
                state["shift"] = shift
                state["active"] = []
                state["excluded"] = []
                if not shift:
                    status.setText("Δεν υπάρχει ενεργή βάρδια.")
                    pending_note.setText("")
                    excluded_note.setText("")
                    group_btn.setEnabled(False)
                    finalize_btn.setEnabled(False)
                    refresh_btn.setEnabled(True)
                    return
                approved = {_clean(x) for x in shift.get("approved", [])}
                state["active"] = [x for x in active if _clean(x.get("agms")) in approved]
                state["excluded"] = [x for x in active if _clean(x.get("agms")) not in approved]
                state["pending"] = bool(shift.get("pending"))
                if state["pending"]:
                    pending_note.setText(
                        "ΠΡΟΣΟΧΗ: Υπάρχουν αστυνομικοί σε ΑΝΑΜΟΝΗ. "
                        "Αποφασίστε ΑΠΟΔΟΧΗ ή ΕΞΑΙΡΕΣΗ στον πίνακα ΑΞΥΠ "
                        "πριν ολοκληρώσετε τη βάρδια."
                    )
                else:
                    pending_note.setText("")
                for item in state["active"]:
                    row_widget = QWidget()
                    row_lay = QHBoxLayout(row_widget)
                    row_lay.setContentsMargins(8, 4, 8, 4)
                    label = QLabel(position_label(item))
                    label.setWordWrap(True)
                    row_lay.addWidget(label, 1)
                    button = QPushButton("ΛΗΞΗ")
                    button.setMinimumWidth(100)
                    button.setMinimumHeight(38)
                    button.clicked.connect(
                        lambda checked=False, rec=dict(item): selected_member_end(rec)
                    )
                    row_lay.addWidget(button)
                    entry = QListWidgetItem()
                    entry.setSizeHint(row_widget.sizeHint())
                    positions_list.addItem(entry)
                    positions_list.setItemWidget(entry, row_widget)
                    state["row_buttons"].append(button)
                count = len(state["active"])
                status.setText(
                    f"Εκκρεμούν {count} λήξεις για την τρέχουσα βάρδια."
                    if count else "Όλοι οι ΑΠΟΔΕΚΤΟΙ έχουν κάνει ΛΗΞΗ."
                )
                if state["excluded"]:
                    excluded_note.setText(
                        "ΕΞΑΙΡΕΜΕΝΟΙ / ΕΚΤΟΣ ΒΑΡΔΙΑΣ — ΔΕΝ ΘΑ ΚΛΕΙΣΟΥΝ:\n"
                        + "\n".join(position_label(x) for x in state["excluded"])
                    )
                else:
                    excluded_note.setText("")
                enabled = not state["pending"]
                refresh_btn.setEnabled(True)
                group_btn.setEnabled(enabled and count > 0)
                finalize_btn.setEnabled(enabled and count == 0)
                for button in state["row_buttons"]:
                    button.setEnabled(enabled)

            def poll():
                future = state["future"]
                if future is None:
                    return
                if not future.done():
                    return
                kind = state["kind"]
                state["future"] = None
                state["kind"] = ""
                try:
                    output = future.result()
                except Exception as exc:
                    QMessageBox.critical(
                        dlg, "Δεν επιβεβαιώθηκε η καταχώριση",
                        f"{exc}\n\nΕλέγξτε ξανά την κατάσταση πριν επαναλάβετε.",
                    )
                    if kind != "load":
                        reload()
                    else:
                        status.setText("Δεν είναι διαθέσιμα τα κεντρικά δεδομένα.")
                        refresh_btn.setEnabled(True)
                    return
                if kind == "load":
                    render(*output)
                elif kind == "single":
                    reload()
                elif kind == "finalize":
                    state["finalized"] = True
                    self._refresh_shift_status()
                    QMessageBox.information(
                        dlg, "Η ΒΑΡΔΙΑ ΟΛΟΚΛΗΡΩΘΗΚΕ",
                        "Η βάρδια οριστικοποιήθηκε και αποθηκεύτηκε.\n"
                        f"Αυτόματες λήξεις με την ολοκλήρωση: "
                        f"{len(output.get('closed_active_agms', []))}.",
                    )
                    dlg.accept()

            refresh_btn.clicked.connect(reload)
            group_btn.clicked.connect(lambda: finalize_shift(True))
            finalize_btn.clicked.connect(lambda: finalize_shift(False))
            cancel_btn.clicked.connect(dlg.reject)
            timer = QTimer(dlg)
            timer.timeout.connect(poll)
            timer.start(180)
            reload()
            try:
                dlg.exec()
            finally:
                timer.stop()
                worker.shutdown(wait=False, cancel_futures=True)
            return bool(state["finalized"])

        def _end_officer_shift_ui(self):
            if not self.store:
                return
            officer = self._officer_selected()
            if not officer:
                QMessageBox.warning(
                    self, "Έλεγχος",
                    "Πληκτρολογήστε έγκυρο Α.Γ.Μ.Σ. Αξιωματικού Υπηρεσίας."
                )
                return
            try:
                self._confirm_and_end_officer_shift(officer)
            except Exception as exc:
                QMessageBox.critical(self, "Λήξη βάρδιας", str(exc))

        def _refresh_shift_status(self):
            if not hasattr(self, "shift_status"):
                return
            if not self.store:
                self.shift_status.setText("ΑΝΑΜΟΝΗ ΣΥΝΔΕΣΗΣ ΜΕ ΤΑ ΚΕΝΤΡΙΚΑ ΔΕΔΟΜΕΝΑ")
                self.start_shift_btn.setEnabled(False)
                self.end_shift_btn.setEnabled(False)
                return
            previous = getattr(self, "_shift_status_future", None)
            if previous is not None and not previous.done():
                return
            self.start_shift_btn.setEnabled(False)
            self.end_shift_btn.setEnabled(False)

            def read_status():
                shift = self.store.current_shift()
                final = None if shift else self.store.current_report_shift()
                return shift, final

            future = self._state_pool.submit(read_status)
            self._shift_status_future = future

            def poll():
                if self._shift_status_future is not future:
                    return
                if not future.done():
                    QTimer.singleShot(200, poll)
                    return
                self._shift_status_future = None
                try:
                    shift, report_shift = future.result()
                except Exception:
                    self.shift_status.setText(
                        "ΑΝΑΜΟΝΗ ΣΥΝΔΕΣΗΣ ΜΕ ΤΑ ΚΕΝΤΡΙΚΑ ΔΕΔΟΜΕΝΑ"
                    )
                    self.start_shift_btn.setEnabled(False)
                    self.end_shift_btn.setEnabled(False)
                    return
                officer = self._officer_selected()
                if not shift:
                    if report_shift and _clean(report_shift.get("status")).upper() == "FINALIZED":
                        ended = _clean(report_shift.get("end_at"))
                        try:
                            ended_dt = datetime.fromisoformat(ended)
                            ended_text = f"{ended_dt:%d/%m/%Y %H:%M}"
                        except Exception:
                            ended_text = ended.replace("T", " ") if ended else "—"
                        self.shift_status.setText(
                            f"Η προηγούμενη βάρδια έληξε {ended_text}.\\n"
                            "Η αναφορά της παραμένει ΤΡΕΧΟΥΣΑ μέχρι την έναρξη της επόμενης."
                        )
                    else:
                        self.shift_status.setText("Δεν υπάρχει ενεργή βάρδια.")
                    self.start_shift_btn.setEnabled(bool(officer))
                    self.end_shift_btn.setEnabled(False)
                    return
                started = datetime.fromisoformat(shift["start_at"])
                starter = shift.get("start_officer") or {}
                name = _clean(f"{starter.get('rank','')} {starter.get('full_name','')}")
                pending = len(shift.get("pending", []))
                msg = f"Ενεργή βάρδια από {started:%d/%m/%Y %H:%M}\\n{name}"
                if pending:
                    msg += f"\\nΑναμονή επιβεβαίωσης: {pending}"
                self.shift_status.setText(msg)
                self.start_shift_btn.setEnabled(False)
                self.end_shift_btn.setEnabled(bool(officer))

            QTimer.singleShot(100, poll)

        def _check_shift_updates(self):
            if not self.store:
                return
            future = self._shift_future
            if future is None:
                self._shift_future = self._network_pool.submit(self.store.sync_active_shift)
                return
            if not future.done():
                return
            self._shift_future = None
            try:
                shift = future.result()
            except Exception:
                if hasattr(self, "shift_status"):
                    self.shift_status.setText("ΑΝΑΜΟΝΗ ΣΥΝΔΕΣΗΣ ΜΕ ΚΕΝΤΡΙΚΑ ΔΕΔΟΜΕΝΑ")
                return
            if not shift:
                if self._last_shift_signature != "NO_SHIFT":
                    self._last_shift_signature = "NO_SHIFT"
                    self._refresh_shift_status()
                return
            signature = json.dumps({
                "approved": shift.get("approved", []),
                "excluded": shift.get("excluded", []),
                "pending": shift.get("pending", [])
            }, ensure_ascii=False, sort_keys=True)
            if signature != self._last_shift_signature:
                self._last_shift_signature = signature
                self._refresh_shift_status()
            if shift.get("pending") and not self._pending_prompt_active:
                self._resolve_pending_members()

        def _live_map(self):
            if not self.store:
                QMessageBox.warning(self,"Έλεγχος","Δεν είναι διαθέσιμος ο φάκελος δεδομένων."); return
            LiveMapDialog(self.store, self._people_by_agms, self._officer_selected, self).exec()

        def _shift_admin(self):
            pass

        def _backup(self):
            if not self.store:
                QMessageBox.warning(self, "Backup", "Ορίστε πρώτα κεντρικό φάκελο δεδομένων.")
                return
            try:
                path = self.store.backup_data()
                QMessageBox.information(self, "Backup", f"Το αντίγραφο ασφαλείας δημιουργήθηκε:\n{path}")
            except Exception as exc:
                QMessageBox.critical(self, "Backup", str(exc))

        def _audit_view(self):
            if not self.store:
                QMessageBox.warning(self, "Audit log", "Ορίστε πρώτα κεντρικό φάκελο δεδομένων.")
                return
            path = self.store.audit / "admin_audit.jsonl"
            if not path.exists():
                QMessageBox.information(self, "Ιστορικό ενεργειών", "Δεν υπάρχουν ακόμη καταγραφές ενεργειών.")
                return
            dlg = QDialog(self); dlg.setWindowTitle("ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ - Ιστορικό ενεργειών"); dlg.resize(900, 600)
            lay = QVBoxLayout(dlg); browser = QTextBrowser(); browser.setStyleSheet("font-family:Consolas,monospace")
            lines = path.read_text(encoding="utf-8", errors="replace").splitlines()[-300:]
            pretty=[]
            for line in lines:
                try:
                    r=json.loads(line); pretty.append(f"{r.get('timestamp','')} | {r.get('user','')}@{r.get('computer','')} | {r.get('action','')} | {json.dumps(r.get('details',{}),ensure_ascii=False)}")
                except Exception:
                    pretty.append(line)
            browser.setPlainText("\n".join(pretty)); lay.addWidget(browser); dlg.exec()

        def _admin(self):
            if not self.store:
                QMessageBox.warning(self, "Έλεγχος", "Δεν είναι διαθέσιμο το .tde_runtime."); return
            try:
                self.store.log_admin_action("OFFICER_DUTY_PANEL_OPENED")
            except Exception:
                pass
            dlg = QDialog(self); dlg.setWindowTitle("ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ"); dlg.resize(720, 620)
            t = THEME_DARK if _clean(self.config.theme).lower() == "dark" else THEME_LIGHT
            dlg.setStyleSheet(_theme_stylesheet(self.config.theme))
            lay = QVBoxLayout(dlg); lay.setContentsMargins(24,22,24,22); lay.setSpacing(10)
            header = QLabel("ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ"); header.setStyleSheet(f"font-size:20px;font-weight:900;color:{t['text']}")
            lay.addWidget(header)
            officer_agms = QLineEdit(); officer_agms.setPlaceholderText("Α.Γ.Μ.Σ. Αξιωματικού Υπηρεσίας")
            officer_name = QLabel("—"); officer_name.setStyleSheet("font-weight:800;font-size:15px")
            form = QFormLayout(); form.addRow("Α.Γ.Μ.Σ.", officer_agms); form.addRow("Αξιωματικός", officer_name); lay.addLayout(form)
            status = QLabel(""); status.setWordWrap(True); status.setStyleSheet(
                f"background:{t['card2']};color:{t['text']};border:1px solid {t['border']};border-radius:10px;padding:9px;font-weight:700"
            )
            lay.addWidget(status)
            members = QListWidget(); lay.addWidget(members, 1)
            row = QHBoxLayout()
            accept_btn = QPushButton("ΑΠΟΔΟΧΗ"); exclude_btn = QPushButton("ΕΞΑΙΡΕΣΗ"); refresh_btn = QPushButton("ΑΝΑΝΕΩΣΗ")
            row.addWidget(accept_btn); row.addWidget(exclude_btn); row.addWidget(refresh_btn); lay.addLayout(row)
            row2 = QHBoxLayout()
            start_shift = QPushButton("ΕΝΑΡΞΗ ΒΑΡΔΙΑΣ"); end_shift = QPushButton("ΛΗΞΗ ΒΑΡΔΙΑΣ"); live = QPushButton("ΖΩΝΤΑΝΗ ΕΙΚΟΝΑ")
            start_shift.setMinimumHeight(48); end_shift.setMinimumHeight(48)
            row2.addWidget(start_shift); row2.addWidget(end_shift); row2.addWidget(live); lay.addLayout(row2)

            def selected_officer():
                return self._people_by_agms.get(_clean(officer_agms.text()))

            def officer_changed(*_):
                p = selected_officer()
                officer_name.setText(_clean(f"{p.rank} {p.full_name}") if p else "—")

            def refresh_members():
                shift = self.store.sync_active_shift() if self.store else None
                members.clear()
                if not shift:
                    status.setText("Δεν υπάρχει ενεργή βάρδια.")
                    accept_btn.setEnabled(False); exclude_btn.setEnabled(False); end_shift.setEnabled(False); start_shift.setEnabled(True)
                    return
                start_shift.setEnabled(False); end_shift.setEnabled(True); accept_btn.setEnabled(True); exclude_btn.setEnabled(True)
                st = datetime.fromisoformat(shift["start_at"])
                off = shift.get("start_officer") or {}
                status.setText(f"Ενεργή βάρδια από {st:%d/%m/%Y %H:%M} — Αξιωματικός: {_clean(f"{off.get('rank','')} {off.get('full_name','')}")}")
                for state_key, label in (("approved", "ΑΠΟΔΕΚΤΟΣ"), ("pending", "ΑΝΑΜΟΝΗ"), ("excluded", "ΕΞΑΙΡΕΣΗ")):
                    for agms in shift.get(state_key, []):
                        person = self._people_by_agms.get(_clean(agms))
                        name = _clean(f"{person.rank} {person.full_name}") if person else f"ΑΓΜΣ {agms}"
                        members.addItem(f"{label} | {agms} | {name}")
                        item = members.item(members.count()-1); item.setData(Qt.UserRole, _clean(agms))

            def change_state(state):
                item = members.currentItem()
                if not item:
                    QMessageBox.warning(dlg, "Έλεγχος", "Επιλέξτε πρώτα άτομο."); return
                if state == "EXCLUDED":
                    agms = _clean(item.data(Qt.UserRole))
                    person = self._people_by_agms.get(agms)
                    name = _clean(f"{person.rank} {person.full_name}") if person else f"ΑΓΜΣ {agms}"
                    confirm = QMessageBox.question(
                        dlg,
                        "Επιβεβαίωση εξαίρεσης",
                        f"{name}\n\nΕίστε βέβαιοι ότι θέλετε να ΕΞΑΙΡΕΣΕΤΕ το άτομο από τη βάρδια;",
                        QMessageBox.Yes | QMessageBox.No,
                        QMessageBox.No
                    )
                    if confirm != QMessageBox.Yes:
                        return
                try:
                    self.store.set_shift_member_state(item.data(Qt.UserRole), state); refresh_members()
                except Exception as exc:
                    QMessageBox.critical(dlg, "Βάρδια", str(exc))

            def do_start():
                p = selected_officer()
                if not p:
                    QMessageBox.warning(dlg, "Έλεγχος", "Πληκτρολογήστε έγκυρο Α.Γ.Μ.Σ. Αξιωματικού Υπηρεσίας."); return
                try:
                    shift = self.store.start_officer_shift(p)
                    QMessageBox.information(dlg, "Βάρδια", f"Έναρξη βάρδιας {shift['start_at'].replace('T',' ')}.\nΕλέγξτε τους ΑΠΟΔΕΚΤΟΥΣ και εξαιρέστε όποιον ανήκει στην προηγούμενη βάρδια.")
                    refresh_members()
                except Exception as exc:
                    QMessageBox.critical(dlg, "Βάρδια", str(exc))

            def do_end():
                p = selected_officer()
                if not p:
                    QMessageBox.warning(dlg, "Έλεγχος", "Πληκτρολογήστε έγκυρο Α.Γ.Μ.Σ. Αξιωματικού Υπηρεσίας."); return
                try:
                    if self._confirm_and_end_officer_shift(p):
                        refresh_members()
                except Exception as exc:
                    QMessageBox.critical(dlg, "Λήξη βάρδιας", str(exc))

            officer_agms.textChanged.connect(officer_changed)
            accept_btn.clicked.connect(lambda: change_state("APPROVED")); exclude_btn.clicked.connect(lambda: change_state("EXCLUDED")); refresh_btn.clicked.connect(refresh_members)
            start_shift.clicked.connect(do_start); end_shift.clicked.connect(do_end); live.clicked.connect(self._live_map)
            timer = QTimer(dlg); timer.timeout.connect(refresh_members); timer.start(2500)
            refresh_members(); dlg.exec()

        def _live_map(self):
            if not self.store:
                QMessageBox.warning(self, "Έλεγχος", "Ορίστε πρώτα κεντρικό φάκελο δεδομένων.")
                return
            LiveMapDialog(self.store, self._people_by_agms, self._officer_selected, self).exec()

        def _shift_admin(self):
            self._admin()

        def _add_position(self):
            value, ok = QInputDialog.getText(self, "Νέα θέση", "Ονομασία / αριθμός θέσης")
            value = _clean(value)
            if ok and value:
                existing = [self.position.itemText(i) for i in range(self.position.count())]
                if value not in existing: self.position.addItem(value)
                self.position.setCurrentText(value); self._save_config()

        def _report(self):
            if not self.store:
                QMessageBox.warning(self, "Έλεγχος", "Δεν είναι διαθέσιμο το .tde_runtime."); return
            shift = self.store.sync_active_shift()
            if not shift:
                shift = self.store.current_report_shift()
            if not shift:
                QMessageBox.information(self, "Τρέχον έντυπο", "Δεν υπάρχει διαθέσιμη τρέχουσα αναφορά."); return
            try:
                path = create_shift_pdf(self.store, shift, self._people_by_agms)
                _open_path(path)
            except Exception as exc:
                QMessageBox.critical(self, "PDF", str(exc))

        def _open_files(self):
            if not self.store:
                QMessageBox.warning(self, "Έλεγχος", "Δεν είναι διαθέσιμο το .tde_runtime."); return
            try:
                _open_path(self.store.archive)
            except Exception as exc:
                QMessageBox.critical(self, "Αρχεία", str(exc))

        def _report_archive(self):
            if not self.store:
                QMessageBox.warning(self, "Έλεγχος", "Δεν είναι διαθέσιμο το .tde_runtime.")
                return

            dlg = QDialog(self)
            dlg.setWindowTitle("ΑΡΧΕΙΟ ΑΝΑΦΟΡΩΝ ΒΑΡΔΙΑΣ")
            dlg.resize(900, 650)
            lay = QVBoxLayout(dlg)
            lay.setContentsMargins(20, 18, 20, 18)
            lay.setSpacing(10)

            title = QLabel("ΑΡΧΕΙΟ ΑΝΑΦΟΡΩΝ ΒΑΡΔΙΑΣ")
            title.setStyleSheet("font-size:18px;font-weight:900")
            lay.addWidget(title)

            hint = QLabel(
                "Εμφανίζονται αρχικά οι 20 πιο πρόσφατες βάρδιες. "
                "Χρησιμοποιήστε ημερομηνίες ή/και Αξιωματικό Υπηρεσίας για παλαιότερη αναζήτηση."
            )
            hint.setWordWrap(True)
            lay.addWidget(hint)

            filters = QFormLayout()
            date_filter = QCheckBox("Χρήση ημερομηνιών")
            from_date = QDateEdit()
            to_date = QDateEdit()
            from_date.setCalendarPopup(True)
            to_date.setCalendarPopup(True)
            from_date.setDisplayFormat("dd/MM/yyyy")
            to_date.setDisplayFormat("dd/MM/yyyy")
            from_date.setDate(QDate.currentDate().addMonths(-1))
            to_date.setDate(QDate.currentDate())
            from_date.setEnabled(False)
            to_date.setEnabled(False)
            officer_edit = QLineEdit()
            officer_edit.setPlaceholderText("Επώνυμο, όνομα ή Α.Γ.Μ.Σ. Αξιωματικού")
            filters.addRow("", date_filter)
            filters.addRow("Από", from_date)
            filters.addRow("Έως", to_date)
            filters.addRow("Αξιωματικός", officer_edit)
            lay.addLayout(filters)

            actions = QHBoxLayout()
            search_btn = QPushButton("ΑΝΑΖΗΤΗΣΗ")
            clear_btn = QPushButton("ΚΑΘΑΡΙΣΜΟΣ")
            actions.addWidget(search_btn)
            actions.addWidget(clear_btn)
            actions.addStretch()
            lay.addLayout(actions)

            listing = QListWidget()
            lay.addWidget(listing, 1)
            info = QLabel("")
            lay.addWidget(info)

            buttons = QHBoxLayout()
            open_btn = QPushButton("ΑΝΟΙΓΜΑ ΑΝΑΦΟΡΑΣ")
            close_btn = QPushButton("ΚΛΕΙΣΙΜΟ")
            buttons.addWidget(open_btn)
            buttons.addStretch()
            buttons.addWidget(close_btn)
            lay.addLayout(buttons)

            current_results: list[dict] = []

            def shift_label(shift: dict) -> str:
                try:
                    start_at = datetime.fromisoformat(_clean(shift.get("start_at")))
                    end_at = datetime.fromisoformat(_clean(shift.get("end_at")))
                    period = f"{start_at:%d/%m/%Y %H:%M} – {end_at:%d/%m/%Y %H:%M}"
                except Exception:
                    period = _clean(shift.get("start_at")).replace("T", " ")
                officer = shift.get("start_officer") or shift.get("end_officer") or {}
                officer_name = _clean(
                    f"{officer.get('rank','')} {officer.get('surname','')} {officer.get('given_name','')}"
                ) or _clean(officer.get("full_name")) or "ΧΩΡΙΣ ΟΝΟΜΑ"
                return f"{period}  |  {officer_name}"

            def fill_results(shifts: list[dict], default_view: bool = False):
                nonlocal current_results
                current_results = shifts
                listing.clear()
                for shift in shifts:
                    item = QListWidgetItem(shift_label(shift))
                    item.setData(Qt.UserRole, _clean(shift.get("shift_id")))
                    listing.addItem(item)
                if default_view:
                    info.setText(f"Τελευταίες ολοκληρωμένες βάρδιες: {len(shifts)}")
                else:
                    info.setText(f"Βρέθηκαν {len(shifts)} βάρδιες.")
                if shifts:
                    listing.setCurrentRow(0)

            def show_recent():
                fill_results(self.store.search_finalized_shifts(limit=20), default_view=True)

            def do_search():
                use_dates = date_filter.isChecked()
                d_from = from_date.date().toPython() if use_dates else None
                d_to = to_date.date().toPython() if use_dates else None
                if d_from and d_to and d_from > d_to:
                    QMessageBox.warning(dlg, "Αναφορές", "Η ημερομηνία «Από» είναι μετά την ημερομηνία «Έως».")
                    return
                officer_query = _clean(officer_edit.text())
                if not use_dates and not officer_query:
                    show_recent()
                    return
                shifts = self.store.search_finalized_shifts(
                    date_from=d_from,
                    date_to=d_to,
                    officer_query=officer_query,
                )
                fill_results(shifts)

            def clear_filters():
                date_filter.setChecked(False)
                from_date.setDate(QDate.currentDate().addMonths(-1))
                to_date.setDate(QDate.currentDate())
                officer_edit.clear()
                show_recent()

            def selected_shift():
                item = listing.currentItem()
                if not item:
                    return None
                sid = _clean(item.data(Qt.UserRole))
                return next((s for s in current_results if _clean(s.get("shift_id")) == sid), None)

            def open_report():
                shift = selected_shift()
                if not shift:
                    QMessageBox.warning(dlg, "Αναφορές", "Επιλέξτε πρώτα μία βάρδια.")
                    return
                try:
                    path = create_shift_pdf(self.store, shift, self._people_by_agms)
                    _open_path(path)
                except Exception as exc:
                    QMessageBox.critical(dlg, "Αναφορές", str(exc))

            date_filter.toggled.connect(from_date.setEnabled)
            date_filter.toggled.connect(to_date.setEnabled)
            search_btn.clicked.connect(do_search)
            clear_btn.clicked.connect(clear_filters)
            open_btn.clicked.connect(open_report)
            listing.itemDoubleClicked.connect(lambda *_: open_report())
            officer_edit.returnPressed.connect(do_search)
            close_btn.clicked.connect(dlg.accept)

            show_recent()
            dlg.exec()

        def _personnel_history(self):
            if not self.store:
                QMessageBox.warning(self, "Έλεγχος", "Δεν είναι διαθέσιμο το .tde_runtime.")
                return
            pin, ok = QInputDialog.getText(self, "Ιστορικό προσωπικού", "Κωδικός πρόσβασης", QLineEdit.Password)
            if not ok:
                return
            if hashlib.sha256(_clean(pin).encode("utf-8")).hexdigest() != STATISTICS_PIN_SHA256:
                QMessageBox.warning(self, "Ιστορικό προσωπικού", "Λανθασμένος κωδικός.")
                return

            dlg = QDialog(self)
            dlg.setWindowTitle("ΙΣΤΟΡΙΚΟ ΠΡΟΣΩΠΙΚΟΥ")
            dlg.resize(1180, 780)
            lay = QVBoxLayout(dlg)
            lay.setContentsMargins(20, 18, 20, 18)
            lay.setSpacing(9)

            title = QLabel("ΙΣΤΟΡΙΚΟ ΠΡΟΣΩΠΙΚΟΥ")
            title.setStyleSheet("font-size:18px;font-weight:900")
            lay.addWidget(title)
            hint = QLabel(
                "Βασικές αναζητήσεις: «ΠΟΥ/ΠΟΤΕ ΔΟΥΛΕΨΕ Ο ΤΑΔΕ;» ή «ΠΟΙΟΣ ΗΤΑΝ ΕΚΕΙ ΤΟΤΕ;». "
                "Τα υπόλοιπα φίλτρα είναι προαιρετικά και συνδυάζονται μεταξύ τους."
            )
            hint.setWordWrap(True)
            lay.addWidget(hint)

            workers = self.store.list_personnel_workers()
            person_combo = ArrowComboBox()
            person_combo.setEditable(True)
            person_combo.addItem("ΟΛΟΙ ΟΣΟΙ ΕΧΟΥΝ ΔΟΥΛΕΨΕΙ", "")
            for worker in workers:
                label = _clean(
                    f"{worker.get('rank','')} {worker.get('surname','')} {worker.get('given_name','')} "
                    f"(ΑΓΜΣ {worker.get('agms','')})"
                )
                person_combo.addItem(label, _clean(worker.get("agms")))

            use_dates = QCheckBox("Φίλτρο ημερομηνιών")
            from_date = QDateEdit()
            to_date = QDateEdit()
            from_date.setCalendarPopup(True)
            to_date.setCalendarPopup(True)
            from_date.setDisplayFormat("dd/MM/yyyy")
            to_date.setDisplayFormat("dd/MM/yyyy")
            from_date.setDate(QDate.currentDate().addMonths(-1))
            to_date.setDate(QDate.currentDate())
            from_date.setEnabled(False)
            to_date.setEnabled(False)

            positions = self.store.list_session_positions()
            position_combo = ArrowComboBox()
            position_combo.addItem("ΟΛΕΣ")
            for pos in positions:
                position_combo.addItem(pos)
            direction_combo = ArrowComboBox()
            direction_combo.addItems(["ΟΛΑ", "ΕΙΣΟΔΟΣ", "ΕΞΟΔΟΣ"])

            use_clock = QCheckBox("Φίλτρο ώρας μέσα στην ημέρα")
            time_from = QTimeEdit()
            time_to = QTimeEdit()
            time_from.setDisplayFormat("HH:mm")
            time_to.setDisplayFormat("HH:mm")
            time_from.setTime(QTime(0, 0))
            time_to.setTime(QTime(23, 59))
            time_from.setEnabled(False)
            time_to.setEnabled(False)

            exact_moment = QCheckBox("Στιγμιαία αναζήτηση — ποιος ήταν εκεί τότε")
            moment_date = QDateEdit()
            moment_time = QTimeEdit()
            moment_date.setCalendarPopup(True)
            moment_date.setDisplayFormat("dd/MM/yyyy")
            moment_date.setDate(QDate.currentDate())
            moment_time.setDisplayFormat("HH:mm")
            moment_time.setTime(QTime.currentTime())
            moment_date.setEnabled(False)
            moment_time.setEnabled(False)

            min_duration = QSpinBox()
            max_duration = QSpinBox()
            for box in (min_duration, max_duration):
                box.setRange(0, 24 * 60)
                box.setSuffix(" λεπτά")
                box.setSpecialValueText("Χωρίς όριο")
            min_duration.setToolTip("Ελάχιστη διάρκεια μίας συνεχόμενης παραμονής στη θέση")
            max_duration.setToolTip("Μέγιστη διάρκεια μίας συνεχόμενης παραμονής στη θέση")

            form = QFormLayout()
            form.addRow("Άτομο", person_combo)
            form.addRow("", use_dates)
            form.addRow("Από ημερομηνία", from_date)
            form.addRow("Έως ημερομηνία", to_date)
            form.addRow("Θέση", position_combo)
            form.addRow("Κατεύθυνση", direction_combo)
            form.addRow("", use_clock)
            form.addRow("Ώρα από", time_from)
            form.addRow("Ώρα έως", time_to)
            form.addRow("", exact_moment)
            form.addRow("Ημερομηνία στιγμής", moment_date)
            form.addRow("Ώρα στιγμής", moment_time)
            form.addRow("Διάρκεια από", min_duration)
            form.addRow("Διάρκεια έως", max_duration)
            lay.addLayout(form)

            quick = QHBoxLayout()
            person_mode_btn = QPushButton("ΠΟΥ / ΠΟΤΕ ΔΟΥΛΕΨΕ Ο ΤΑΔΕ;")
            moment_mode_btn = QPushButton("ΠΟΙΟΣ ΗΤΑΝ ΕΚΕΙ ΤΟΤΕ;")
            quick.addWidget(person_mode_btn)
            quick.addWidget(moment_mode_btn)
            quick.addStretch()
            lay.addLayout(quick)

            actions = QHBoxLayout()
            search_btn = QPushButton("ΑΝΑΖΗΤΗΣΗ")
            clear_btn = QPushButton("ΚΑΘΑΡΙΣΜΟΣ")
            pdf_btn = QPushButton("PDF ΑΠΟΤΕΛΕΣΜΑΤΩΝ")
            actions.addWidget(search_btn)
            actions.addWidget(clear_btn)
            actions.addWidget(pdf_btn)
            actions.addStretch()
            lay.addLayout(actions)

            summary_box = QLabel("")
            summary_box.setWordWrap(True)
            summary_box.setObjectName("infoBox")
            lay.addWidget(summary_box)

            result_label = QLabel("")
            lay.addWidget(result_label)

            table = QTableWidget(0, 8)
            table.setHorizontalHeaderLabels(
                ["ΗΜΕΡΟΜΗΝΙΑ", "ΑΓΜΣ", "ΟΝΟΜΑ", "ΘΕΣΗ", "ΚΑΤΕΥΘΥΝΣΗ", "ΕΝΑΡΞΗ", "ΛΗΞΗ", "ΔΙΑΡΚΕΙΑ"]
            )
            table.setAlternatingRowColors(True)
            table.setEditTriggers(QAbstractItemView.NoEditTriggers)
            table.setSelectionBehavior(QAbstractItemView.SelectRows)
            table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
            table.horizontalHeader().setStretchLastSection(True)
            lay.addWidget(table, 1)

            current_rows: list[dict] = []
            current_criteria = ""

            def qtime_to_time(value: QTime) -> dtime:
                return dtime(value.hour(), value.minute())

            def selected_person_query() -> str:
                text_value = _clean(person_combo.currentText())
                if _norm_search(text_value) == _norm_search("ΟΛΟΙ ΟΣΟΙ ΕΧΟΥΝ ΔΟΥΛΕΨΕΙ"):
                    return ""
                idx = person_combo.currentIndex()
                if idx >= 0 and text_value == _clean(person_combo.itemText(idx)):
                    data = person_combo.itemData(idx)
                    if _clean(data):
                        return _clean(data)
                return text_value

            def summary_text(rows: list[dict]) -> str:
                summary = summarize_personnel_history(rows)
                line1 = (
                    f"Καταχωρίσεις: {len(rows)}   |   Άτομα: {summary['people']}   |   "
                    f"Συνολικός χρόνος: {_format_seconds(summary['total_seconds'])}"
                )
                parts = [line1]
                for direction in ("ΕΙΣΟΔΟΣ", "ΕΞΟΔΟΣ"):
                    sec = int(summary["by_direction"].get(direction, 0))
                    if not sec:
                        continue
                    pos_parts = []
                    for (d, pos), psec in sorted(summary["by_direction_position"].items(), key=lambda x: (x[0][0], x[0][1])):
                        if d == direction and psec:
                            pos_parts.append(f"Θέση {pos}: {_format_seconds(psec)}")
                    detail = " · ".join(pos_parts)
                    parts.append(f"{direction}: {_format_seconds(sec)}" + (f"   |   {detail}" if detail else ""))
                return "\n".join(parts)

            def criteria_text() -> str:
                bits = []
                pq = selected_person_query()
                if pq:
                    bits.append(f"Άτομο: {person_combo.currentText()}")
                else:
                    bits.append("Άτομο: ΟΛΟΙ")
                if exact_moment.isChecked():
                    bits.append(
                        f"Στιγμή: {moment_date.date().toString('dd/MM/yyyy')} {moment_time.time().toString('HH:mm')}"
                    )
                elif use_dates.isChecked():
                    bits.append(
                        f"Ημερομηνίες: {from_date.date().toString('dd/MM/yyyy')}–{to_date.date().toString('dd/MM/yyyy')}"
                    )
                pos = _clean(position_combo.currentText())
                direction = _clean(direction_combo.currentText())
                if pos != "ΟΛΕΣ":
                    bits.append(f"Θέση: {pos}")
                if direction != "ΟΛΑ":
                    bits.append(f"Κατεύθυνση: {direction}")
                if use_clock.isChecked() and not exact_moment.isChecked():
                    bits.append(f"Ώρες: {time_from.time().toString('HH:mm')}–{time_to.time().toString('HH:mm')}")
                if min_duration.value():
                    bits.append(f"Διάρκεια ≥ {min_duration.value()} λεπτά")
                if max_duration.value():
                    bits.append(f"Διάρκεια ≤ {max_duration.value()} λεπτά")
                return " | ".join(bits)

            def render_rows(rows: list[dict], label_prefix: str = "Βρέθηκαν"):
                nonlocal current_rows, current_criteria
                current_rows = rows
                current_criteria = criteria_text()
                table.setRowCount(len(rows))
                for r, row in enumerate(rows):
                    try:
                        sdt = datetime.fromisoformat(row["start_at"])
                        edt = datetime.fromisoformat(row["end_at"])
                        date_text = sdt.strftime("%d/%m/%Y")
                        start_text = sdt.strftime("%H:%M:%S")
                        end_text = edt.strftime("%H:%M:%S")
                    except Exception:
                        date_text = _clean(row.get("start_at"))[:10]
                        start_text = _clean(row.get("start_at"))[11:19]
                        end_text = _clean(row.get("end_at"))[11:19]
                    name = _clean(f"{row.get('rank','')} {row.get('surname','')} {row.get('given_name','')}") or _clean(row.get("full_name"))
                    values = [
                        date_text,
                        _clean(row.get("agms")),
                        name,
                        _clean(row.get("position")),
                        _clean(row.get("direction")),
                        start_text,
                        end_text,
                        _format_seconds(int(row.get("duration_seconds") or 0)),
                    ]
                    for col, value in enumerate(values):
                        table.setItem(r, col, QTableWidgetItem(str(value)))
                result_label.setText(f"{label_prefix} {len(rows)} καταχωρίσεις.")
                summary_box.setText(summary_text(rows) if rows else "Δεν βρέθηκαν καταχωρίσεις με αυτά τα κριτήρια.")

            def do_search():
                d_from = from_date.date().toPython() if use_dates.isChecked() and not exact_moment.isChecked() else None
                d_to = to_date.date().toPython() if use_dates.isChecked() and not exact_moment.isChecked() else None
                if d_from and d_to and d_from > d_to:
                    QMessageBox.warning(dlg, "Ιστορικό", "Η ημερομηνία «Από» είναι μετά την ημερομηνία «Έως».")
                    return
                moment = None
                if exact_moment.isChecked():
                    qd = moment_date.date()
                    qt = moment_time.time()
                    moment = datetime(qd.year(), qd.month(), qd.day(), qt.hour(), qt.minute())
                tf = qtime_to_time(time_from.time()) if use_clock.isChecked() and not exact_moment.isChecked() else None
                tt = qtime_to_time(time_to.time()) if use_clock.isChecked() and not exact_moment.isChecked() else None
                if max_duration.value() and min_duration.value() and max_duration.value() < min_duration.value():
                    QMessageBox.warning(dlg, "Ιστορικό", "Η μέγιστη διάρκεια είναι μικρότερη από την ελάχιστη.")
                    return

                rows = self.store.search_personnel_sessions(
                    person_query=selected_person_query(),
                    date_from=d_from,
                    date_to=d_to,
                    position=_clean(position_combo.currentText()),
                    direction=_clean(direction_combo.currentText()),
                    time_from=tf,
                    time_to=tt,
                    moment=moment,
                    min_duration_minutes=min_duration.value(),
                    max_duration_minutes=max_duration.value(),
                )
                render_rows(rows)
                try:
                    self.store.log_admin_action("PERSONNEL_HISTORY_SEARCH", {
                        "criteria": criteria_text(), "results": len(rows)
                    })
                except Exception:
                    pass

            def clear_search():
                person_combo.setCurrentIndex(0)
                use_dates.setChecked(False)
                from_date.setDate(QDate.currentDate().addMonths(-1))
                to_date.setDate(QDate.currentDate())
                position_combo.setCurrentIndex(0)
                direction_combo.setCurrentIndex(0)
                use_clock.setChecked(False)
                time_from.setTime(QTime(0, 0))
                time_to.setTime(QTime(23, 59))
                exact_moment.setChecked(False)
                moment_date.setDate(QDate.currentDate())
                moment_time.setTime(QTime.currentTime())
                min_duration.setValue(0)
                max_duration.setValue(0)
                rows = self.store.search_personnel_sessions(limit=100)
                render_rows(rows, "Τελευταίες")

            def export_pdf():
                if not current_rows:
                    QMessageBox.warning(dlg, "PDF", "Δεν υπάρχουν αποτελέσματα για εξαγωγή.")
                    return
                try:
                    path = create_personnel_history_pdf(
                        current_rows,
                        current_criteria,
                        summarize_personnel_history(current_rows),
                    )
                    _open_path(path)
                except Exception as exc:
                    QMessageBox.critical(dlg, "PDF", str(exc))

            def update_filter_states():
                moment_on = exact_moment.isChecked()
                date_on = use_dates.isChecked() and not moment_on
                clock_on = use_clock.isChecked() and not moment_on
                from_date.setEnabled(date_on)
                to_date.setEnabled(date_on)
                time_from.setEnabled(clock_on)
                time_to.setEnabled(clock_on)
                moment_date.setEnabled(moment_on)
                moment_time.setEnabled(moment_on)
                use_dates.setEnabled(not moment_on)
                use_clock.setEnabled(not moment_on)

            def choose_person_mode():
                exact_moment.setChecked(False)
                person_combo.setFocus()
                if person_combo.lineEdit():
                    person_combo.lineEdit().selectAll()

            def choose_moment_mode():
                person_combo.setCurrentIndex(0)
                exact_moment.setChecked(True)
                position_combo.setCurrentIndex(0)
                direction_combo.setCurrentIndex(0)
                moment_date.setFocus()

            use_dates.toggled.connect(update_filter_states)
            use_clock.toggled.connect(update_filter_states)
            exact_moment.toggled.connect(update_filter_states)
            search_btn.clicked.connect(do_search)
            clear_btn.clicked.connect(clear_search)
            pdf_btn.clicked.connect(export_pdf)
            person_mode_btn.clicked.connect(choose_person_mode)
            moment_mode_btn.clicked.connect(choose_moment_mode)
            person_combo.lineEdit().returnPressed.connect(do_search)

            update_filter_states()
            rows = self.store.search_personnel_sessions(limit=100)
            render_rows(rows, "Τελευταίες")
            dlg.exec()

        def _snapshot(self):
            if not self.store:
                QMessageBox.warning(self, "Έλεγχος", "Ορίστε κεντρικό φάκελο."); return
            active = self.store.active_all()
            if not active:
                QMessageBox.information(self,"Στιγμιαία εικόνα","Δεν υπάρχουν ανοικτές θέσεις."); return
            lines = [f"{x.get('rank','')} {x.get('full_name','')} — Θ. {x.get('position','')} {x.get('direction','')} — από {x.get('timestamp','').replace('T',' ')}" for x in active]
            QMessageBox.information(self,"Στιγμιαία εικόνα","\n".join(lines))


    app = QApplication(list(args or sys.argv))
    app.setApplicationName("TDE Position Tracking")
    try:
        window = MainWindow()
    except RuntimeError as exc:
        # Fail closed on an inaccessible shared queue or old pending records.
        QMessageBox.critical(None, "ΚΟΙΝΗ ΟΥΡΑ ΘΕΣΕΩΝ — ΑΣΦΑΛΕΙΑ", str(exc))
        return 2
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(run())
